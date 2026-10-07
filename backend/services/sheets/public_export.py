"""Read a Google Sheet through its public "Anyone with the link can view" export.

No Google credentials are involved: the whole spreadsheet is downloaded as
one .xlsx (every tab, full cell text, cell colours and hyperlinks) with a
single request, which keeps the poller well inside any read limits.

If the sheet stops being link-shared, Google answers with a sign-in page
instead of a workbook, and SheetAccessError says exactly that.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Optional

import httpx
from openpyxl import load_workbook

EXPORT_URL = "https://docs.google.com/spreadsheets/d/{id}/export?format=xlsx"
OPEN_URL = "https://docs.google.com/spreadsheets/d/{id}/edit"
XLSX_MAGIC = b"PK"
TIMEOUT_SECONDS = 60
MAX_BYTES = 50_000_000

# Also matches multi-account links such as /spreadsheets/u/1/d/<id>/edit.
_ID_IN_URL = re.compile(r"/spreadsheets/(?:u/\d+/)?d/([A-Za-z0-9_-]{20,})")
_BARE_ID = re.compile(r"^[A-Za-z0-9_-]{25,}$")


class SheetAccessError(Exception):
    """The spreadsheet could not be read; the message is safe to show users."""


def parse_spreadsheet_id(url_or_id: str) -> str:
    text = (url_or_id or "").strip()
    m = _ID_IN_URL.search(text)
    if m:
        return m.group(1)
    if _BARE_ID.match(text):
        return text
    raise SheetAccessError("That doesn't look like a Google Sheets link. Paste the full URL from the browser.")


def open_url(spreadsheet_id: str) -> str:
    return OPEN_URL.format(id=spreadsheet_id)


def fetch_workbook_bytes(spreadsheet_id: str, *, client: Optional[httpx.Client] = None) -> bytes:
    url = EXPORT_URL.format(id=spreadsheet_id)
    own = client is None
    client = client or httpx.Client(timeout=TIMEOUT_SECONDS, follow_redirects=True)
    try:
        resp = client.get(url)
    except httpx.HTTPError as exc:
        raise SheetAccessError(f"Could not reach Google Sheets ({exc.__class__.__name__}). Try again shortly.")
    finally:
        if own:
            client.close()
    if resp.status_code == 404:
        raise SheetAccessError("Google says this spreadsheet doesn't exist. Check the link.")
    if resp.status_code in (401, 403) or (resp.status_code == 200 and not resp.content.startswith(XLSX_MAGIC)):
        raise SheetAccessError(
            "This sheet isn't readable by link. In Google Sheets choose Share → General access → "
            "\"Anyone with the link\" → Viewer, then try again."
        )
    if resp.status_code == 429:
        raise SheetAccessError("Google is rate-limiting reads of this sheet. The next poll will retry.")
    if resp.status_code != 200:
        raise SheetAccessError(f"Google Sheets answered {resp.status_code}. Try again shortly.")
    if len(resp.content) > MAX_BYTES:
        raise SheetAccessError("This spreadsheet is too large to track (over 50 MB exported).")
    return resp.content


@dataclass
class RowData:
    row_number: int
    values: dict                 # header -> text ("" when blank)
    links: dict = field(default_factory=dict)   # header -> hyperlink target, when the cell is a link
    is_example: bool = False


@dataclass
class TabData:
    name: str
    headers: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    hidden: bool = False


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes") if (value.hour or value.minute) else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _is_yellow(cell) -> bool:
    """Pale-yellow fill (the template's example-row colour)."""
    fill = getattr(cell, "fill", None)
    if not fill or not fill.fill_type:
        return False
    rgb = getattr(fill.fgColor, "rgb", None)
    if not isinstance(rgb, str) or len(rgb) < 6:
        return False
    try:
        r, g, b = (int(rgb[-6:][i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return False
    # e.g. FFF2CC / FFFF99 / FCE8B2 are yellow; white, F2F2F2 grey and navy are not.
    return r >= 200 and g >= 190 and b <= 215 and (r - b) >= 35 and (g - b) >= 25


def parse_workbook(data: bytes, header_row: int = 1) -> dict:
    """{tab name: TabData}. A hyperlinked cell's target is kept in RowData.links,
    so a "Final Doc" cell showing "Open doc" still yields the URL."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    tabs = {}
    for ws in wb.worksheets:
        tab = TabData(name=ws.title, hidden=ws.sheet_state != "visible")
        header_cells = [c for c in ws[header_row]] if ws.max_row >= header_row else []
        headers = [cell_text(c.value) for c in header_cells]
        while headers and not headers[-1]:
            headers.pop()
        tab.headers = headers
        if not headers:
            tabs[ws.title] = tab
            continue
        width = len(headers)
        for row in ws.iter_rows(min_row=header_row + 1, max_col=width):
            values, links, yellow, filled = {}, {}, 0, 0
            for header, c in zip(headers, row):
                if not header:
                    continue
                text = cell_text(c.value)
                link = getattr(getattr(c, "hyperlink", None), "target", None)
                if link:
                    links[header] = link
                values[header] = text
                if text:
                    filled += 1
                if _is_yellow(c):
                    yellow += 1
            tab.rows.append(RowData(row_number=row[0].row, values=values, links=links,
                                    is_example=yellow >= max(1, width // 2) and filled > 0))
        tabs[ws.title] = tab
    return tabs


def read_sheet(spreadsheet_id: str, header_row: int = 1) -> dict:
    return parse_workbook(fetch_workbook_bytes(spreadsheet_id), header_row=header_row)
