"""Smoke test for the Buffer tab's Purpose Sheet (BufferSheetPurpose /
BufferSheetRow + routers/buffer_router.py's "Purpose Sheet" endpoints):
confirm create/edit is gated behind the buffer_sheet_edit grant (or admin),
edit/clear row fields, confirm delete is admin-only, delete a row, delete a
purpose (checking its rows cascade), export to .xlsx, and the self-upload
download gate's auto-sync into the sheet (record_self_upload_download) -
the only way a row is ever created, since there is no manual "add row"
endpoint.

Run: backend/venv/Scripts/python.exe tests/buffer_sheet_smoke.py
"""

import os
import sys
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool


BACKEND_DIR = Path(__file__).resolve().parents[1]
os.environ.setdefault("DATABASE_URL", "postgresql://placeholder:placeholder@localhost:5432/placeholder")
os.environ.setdefault("ARCHIVE_DATABASE_URL", os.environ["DATABASE_URL"])
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from models_new import (  # noqa: E402
    Base,
    BufferSelfUpload,
    BufferSelfUploadDownload,
    BufferSheetPurpose,
    BufferSheetRow,
    User,
    UserFeatureAccess,
    UserRole,
)
from routers.buffer_router import (  # noqa: E402
    BufferSelfUploadDownloadCreate,
    BufferSheetPurposeCreate,
    BufferSheetPurposeUpdate,
    BufferSheetRowUpdate,
    create_sheet_purpose,
    delete_sheet_purpose,
    delete_sheet_row,
    export_sheet_xlsx,
    get_sheet,
    record_self_upload_download,
    update_sheet_row,
)
from services.feature_access_service import FEATURE_BUFFER_SHEET_EDIT, set_feature_access  # noqa: E402


engine = create_engine(
    "sqlite://",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)


@event.listens_for(engine, "connect")
def _enable_sqlite_foreign_keys(dbapi_connection, _connection_record):
    # SQLite silently ignores ON DELETE CASCADE unless a connection opts in
    # with this pragma - Postgres (this app's real database) enforces it
    # unconditionally, so without this the test would pass even if the
    # cascade were broken.
    dbapi_connection.execute("PRAGMA foreign_keys=ON")


SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base.metadata.create_all(
    bind=engine,
    tables=[
        User.__table__,
        UserRole.__table__,
        UserFeatureAccess.__table__,
        BufferSheetPurpose.__table__,
        BufferSheetRow.__table__,
        BufferSelfUpload.__table__,
        BufferSelfUploadDownload.__table__,
    ],
)


def _assert(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def _preload_permission_relationships(db, user: User) -> None:
    """has_any_role() lazy-loads role_assignments and has_feature_access()
    lazy-loads feature_grants - force both loads now, while the session is
    still open, so the detached object the caller returns doesn't blow up
    with a DetachedInstanceError the first time some later call (in a
    different session) needs a role/grant check."""
    _ = user.role_assignments
    _ = user.feature_grants


def _create_user(email: str, *, is_admin: bool = False) -> User:
    with SessionLocal() as db:
        user = User(
            email=email,
            name=email.split("@", 1)[0],
            hashed_password="hashed-password",
            is_active=True,
            is_deleted=False,
            is_admin=is_admin,
        )
        db.add(user)
        db.commit()
        db.refresh(user)
        _preload_permission_relationships(db, user)
        db.expunge(user)
        return user


def _grant_sheet_edit(user: User) -> User:
    """Grants FEATURE_BUFFER_SHEET_EDIT to an already-created user and
    returns a fresh, detached User reflecting it - same admin-grants-access
    flow as Admin Queue -> Section Access (routers/admin_router.py's
    set_user_feature_access), just called directly instead of through HTTP."""
    with SessionLocal() as db:
        fresh = db.query(User).filter(User.id == user.id).first()
        set_feature_access(db, fresh, FEATURE_BUFFER_SHEET_EDIT, True, granted_by=None)
        db.commit()
        db.refresh(fresh)
        _preload_permission_relationships(db, fresh)
        db.expunge(fresh)
        return fresh


def test_create_purpose_and_rows_then_export() -> None:
    viewer = _create_user("viewer@example.com")  # Buffer access only, no edit grant
    editor = _grant_sheet_edit(_create_user("editor@example.com"))
    admin = _create_user("admin@example.com", is_admin=True)

    with SessionLocal() as db:
        # A user with plain Buffer access but no buffer_sheet_edit grant
        # must be refused - viewing the sheet is open to anyone with Buffer
        # access, but creating/editing is a narrower, admin-granted layer on
        # top (see buffer_router.py's _require_sheet_edit_access).
        create_denied = False
        try:
            create_sheet_purpose(BufferSheetPurposeCreate(name="Janmashtami"), current_user=viewer, db=db)
        except HTTPException as exc:
            create_denied = exc.status_code == 403
        _assert(create_denied, "creating a purpose without the edit grant should get 403")

        purpose_resp = create_sheet_purpose(BufferSheetPurposeCreate(name="Janmashtami"), current_user=editor, db=db)
        purpose_id = purpose_resp["purpose"]["id"]
        _assert(purpose_resp["purpose"]["color"], "a newly created purpose should get a default color")

        # No create_sheet_row endpoint exists (see this module's docstring) -
        # rows only ever come from _sync_purpose_sheet, so test fixtures are
        # built the same way that helper builds them: directly on the table.
        row1 = BufferSheetRow(
            purpose_id=purpose_id, audio_name="Janmashtami v1",
            client_name="Mansha", editor_name="Kavinder", sort_order=0,
            created_by_user_id=editor.id, updated_by_user_id=editor.id,
        )
        row2 = BufferSheetRow(
            purpose_id=purpose_id, audio_name="Janmashtami v2", sort_order=1,
            created_by_user_id=editor.id, updated_by_user_id=editor.id,
        )
        db.add_all([row1, row2])
        db.commit()
        db.refresh(row1)
        db.refresh(row2)
        row1, row2 = {"id": row1.id}, {"id": row2.id}

        sheet = get_sheet(current_user=viewer, db=db)
        _assert(len(sheet["purposes"]) == 1, "one purpose should be listed")
        _assert(len(sheet["purposes"][0]["rows"]) == 2, "both rows should be listed under it")
        _assert(sheet["clientSuggestions"] == ["Mansha"], "only the one filled-in client should be suggested")
        _assert(sheet["editorSuggestions"] == ["Kavinder"], "only the one filled-in editor should be suggested")

        # A viewer without the edit grant must not be able to edit an
        # existing row's fields either.
        edit_denied = False
        try:
            update_sheet_row(row2["id"], BufferSheetRowUpdate(clientName="Northwind"), current_user=viewer, db=db)
        except HTTPException as exc:
            edit_denied = exc.status_code == 403
        _assert(edit_denied, "editing a row without the edit grant should get 403")

        # Fill in row2's client/editor, then clear row1's client with an
        # explicit empty string (the "clear" convention, same as
        # update_self_upload).
        update_sheet_row(row2["id"], BufferSheetRowUpdate(clientName="Northwind", editorName="Rahul"), current_user=editor, db=db)
        updated_row1 = update_sheet_row(row1["id"], BufferSheetRowUpdate(clientName=""), current_user=editor, db=db)["row"]
        _assert(updated_row1["clientName"] is None, "an explicit empty string should clear clientName")
        _assert(updated_row1["editorName"] == "Kavinder", "editorName should be untouched by a clientName-only update")
        # Fully unassign row1 (clear its remaining editor too) so the
        # onlyFilled export check below actually has an unfilled row to drop.
        update_sheet_row(row1["id"], BufferSheetRowUpdate(editorName=""), current_user=editor, db=db)

        raised = False
        try:
            update_sheet_row(row1["id"], BufferSheetRowUpdate(audioName="   "), current_user=editor, db=db)
        except HTTPException as exc:
            raised = exc.status_code == 400
        _assert(raised, "clearing the required audioName to blank should 400, not silently null it out")

        export = export_sheet_xlsx(purposeIds=None, columns=None, onlyFilled=False, current_user=viewer, db=db)
        _assert(export.status_code == 200, "export should succeed")
        _assert("Buffer-Sheet_" in export.headers["content-disposition"], "export filename should be present")

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(export.body))
        ws = wb.active
        cell_values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        _assert("JANMASHTAMI" in cell_values, "the purpose title row should render uppercased")
        _assert("Janmashtami v1" in cell_values, "row 1's audio name should appear in the export")
        _assert("Northwind" in cell_values, "row 2's client should appear in the export")

        # onlyFilled should drop the still-unassigned row.
        only_filled_export = export_sheet_xlsx(purposeIds=None, columns=None, onlyFilled=True, current_user=viewer, db=db)
        wb2 = load_workbook(BytesIO(only_filled_export.body))
        values2 = [c.value for row in wb2.active.iter_rows() for c in row if c.value is not None]
        _assert("Janmashtami v1" not in values2, "row1 has no client/editor left after clearing and should be filtered out")
        _assert("Janmashtami v2" in values2, "row2 is assigned and should remain")

        # Deletion is admin-only - stricter than the edit grant. `editor` can
        # freely create/edit rows and purposes but must still be refused a
        # delete, proving the edit grant does not implicitly unlock it.
        row_delete_denied = False
        try:
            delete_sheet_row(row2["id"], current_user=editor, db=db)
        except HTTPException as exc:
            row_delete_denied = exc.status_code == 403
        _assert(row_delete_denied, "an editor (non-admin) deleting a row should get 403")

        purpose_delete_denied = False
        try:
            delete_sheet_purpose(purpose_id, current_user=editor, db=db)
        except HTTPException as exc:
            purpose_delete_denied = exc.status_code == 403
        _assert(purpose_delete_denied, "an editor (non-admin) deleting a purpose should get 403")

        delete_sheet_row(row2["id"], current_user=admin, db=db)
        remaining = get_sheet(current_user=viewer, db=db)["purposes"][0]["rows"]
        _assert(len(remaining) == 1 and remaining[0]["id"] == row1["id"], "row2 should be gone after an admin deletes it")

        delete_sheet_purpose(purpose_id, current_user=admin, db=db)
        after_delete = get_sheet(current_user=viewer, db=db)
        _assert(after_delete["purposes"] == [], "purpose should be gone")
        orphaned_rows = db.query(BufferSheetRow).filter(BufferSheetRow.purpose_id == purpose_id).count()
        _assert(orphaned_rows == 0, "deleting a purpose should cascade-delete its rows, not orphan them")


def test_self_upload_download_syncs_to_purpose_sheet() -> None:
    owner = _create_user("owner@example.com")
    downloader = _create_user("downloader@example.com")

    with SessionLocal() as db:
        upload = BufferSelfUpload(
            owner_user_id=owner.id,
            title="Grenade Explosion",
            media_type="audio",
            asset_url="https://example.com/grenade.mp3",
        )
        db.add(upload)
        db.commit()
        db.refresh(upload)

        # First download: "Dussehra" doesn't exist yet as a purpose - it
        # should be created, with a row for this client/editor.
        resp1 = record_self_upload_download(
            upload.id,
            BufferSelfUploadDownloadCreate(clientName="Northwind", purpose="Dussehra"),
            current_user=downloader,
            db=db,
        )
        _assert(resp1["sheetPurpose"]["name"] == "Dussehra", "a new purpose should be created from the typed value")
        _assert(resp1["sheetRow"]["audioName"] == "Grenade Explosion", "the row's audio name should come from the upload's title")
        _assert(resp1["sheetRow"]["clientName"] == "Northwind", "the row's client should come from the download form")
        _assert(resp1["sheetRow"]["editorName"] == downloader.name, "the row's editor should be whoever downloaded it")

        # Second download of the SAME file, for the SAME client, under a
        # purpose typed with different case ("dussehra") - should land in
        # the SAME purpose group and update the SAME row, not create dupes.
        second_editor = _create_user("second-editor@example.com")
        resp2 = record_self_upload_download(
            upload.id,
            BufferSelfUploadDownloadCreate(clientName="Northwind", purpose="dussehra"),
            current_user=second_editor,
            db=db,
        )
        _assert(resp2["sheetPurpose"]["id"] == resp1["sheetPurpose"]["id"], "a case-different repeat should reuse the same purpose")
        _assert(resp2["sheetRow"]["id"] == resp1["sheetRow"]["id"], "same file + same client should update the existing row, not duplicate it")
        _assert(resp2["sheetRow"]["editorName"] == second_editor.name, "editor should update to whoever downloaded it most recently")

        # A download of the same file for a DIFFERENT client under the same
        # purpose should add a second, distinct row.
        resp3 = record_self_upload_download(
            upload.id,
            BufferSelfUploadDownloadCreate(clientName="Sikka", purpose="Dussehra"),
            current_user=downloader,
            db=db,
        )
        _assert(resp3["sheetRow"]["id"] != resp1["sheetRow"]["id"], "a different client should get its own row")

        sheet = get_sheet(current_user=downloader, db=db)
        dussehra = next(p for p in sheet["purposes"] if p["name"] == "Dussehra")
        _assert(len(dussehra["rows"]) == 2, f"expected 2 rows under Dussehra (one per client), got {len(dussehra['rows'])}")

        downloads_logged = db.query(BufferSelfUploadDownload).filter(BufferSelfUploadDownload.upload_id == upload.id).count()
        _assert(downloads_logged == 3, "every download attempt should still get its own log row regardless of sheet dedup")


def test_color_palette_cycles() -> None:
    user = _grant_sheet_edit(_create_user("palette@example.com"))
    with SessionLocal() as db:
        colors = []
        for i in range(3):
            resp = create_sheet_purpose(BufferSheetPurposeCreate(name=f"Event {i}"), current_user=user, db=db)
            colors.append(resp["purpose"]["color"])
        _assert(len(set(colors)) == 3, f"the first three purposes should each get a distinct color, got {colors}")


def run() -> None:
    tests = [
        test_create_purpose_and_rows_then_export,
        test_self_upload_download_syncs_to_purpose_sheet,
        test_color_palette_cycles,
    ]
    for test in tests:
        test()
        print(f"  ok  {test.__name__}")
    print(f"\nbuffer sheet smoke: {len(tests)} passed")


if __name__ == "__main__":
    run()
