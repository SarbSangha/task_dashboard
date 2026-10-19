"""Usage & Spend covers every tool the company pays for.

ChatGPT and Claude (one row per prompt), Grammarly (one row per document) and
Splice (one row per sample download, pay-per-use) are reported through the
same facts as every generation tool, and a paid tool with no usage capture
(Semrush) still counts in total spend as "Paid, not tracked". All figures are
worked out by hand below.
"""

import os
import sys
from datetime import date, datetime
from pathlib import Path

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

BACKEND_DIR = Path(__file__).resolve().parents[1]
os.environ.setdefault("DATABASE_URL", "postgresql://placeholder:placeholder@localhost:5432/placeholder")
os.environ.setdefault("ARCHIVE_DATABASE_URL", os.environ["DATABASE_URL"])
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from models_new import (  # noqa: E402
    Base, GenerationClient, GenerationRecord, ITPortalTool, ITPortalToolCredential, ITPortalToolUsageEvent,
    ToolCreditRate, User, UserFeatureAccess, UserRole,
)
from providers.chatgpt.models import ConversationPrompt, ConversationRecord, ConversationResponse  # noqa: E402
from providers.elevenlabs.models import ElevenlabsGeneration  # noqa: E402
from providers.envato.models import EnvatoGeneration  # noqa: E402
from providers.epidemicsound.models import EpidemicAdaptation  # noqa: E402
from providers.flow.models import FlowGeneration  # noqa: E402
from providers.freepik.models import FreepikGeneration  # noqa: E402
from providers.grammarly_docs.models import GrammarlyDocSession  # noqa: E402
from providers.heygen.models import HeygenGeneration  # noqa: E402
from providers.higgsfield.models import HiggsfieldGeneration  # noqa: E402
from providers.splice.models import SpliceDownload  # noqa: E402
from providers.suno.models import SunoGeneration  # noqa: E402
from utils.credit_report.facts import resolve_output_url  # noqa: E402
from utils.credit_report.output_media import chat_transcript  # noqa: E402
from utils.credit_report.live_sheet import build_live_sheet  # noqa: E402

CLAUDE_ID = "3f0c2a1e-1111-4222-8333-944455556666"
# What the Claude capture reports: the API request it intercepted, not the page.
CLAUDE_API_URL = f"https://claude.ai/api/organizations/0a0b0c0d-0000-4000-8000-000000000000/chat_conversations/{CLAUDE_ID}?tree=True"

engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
TABLES = [User, UserRole, UserFeatureAccess, GenerationClient, ITPortalTool, ITPortalToolCredential,
          ITPortalToolUsageEvent, ToolCreditRate, GenerationRecord, FreepikGeneration, EnvatoGeneration,
          HeygenGeneration, HiggsfieldGeneration, ElevenlabsGeneration, FlowGeneration, SunoGeneration,
          EpidemicAdaptation, SpliceDownload, GrammarlyDocSession, ConversationRecord, ConversationPrompt,
          ConversationResponse]
Base.metadata.create_all(bind=engine, tables=[m.__table__ for m in TABLES])

# A whole month, so each monthly fee is counted in full.
START, END = "2026-10-01", "2026-10-31"
FEES = {"ChatGPT": 3000.0, "Claude": 1800.0, "Grammarly": 900.0, "Semrush": 6000.0}
SPLICE_RATE = 10.0          # rupees per Splice credit; a sample download is 1 credit
# ChatGPT: 3 prompts share 3000 -> 1000 each (Asha 2, Bob 1). Claude: Bob's 1 prompt takes 1800.
# Grammarly: Asha's 1 document takes 900. Splice: Bob's 2 downloads x 1 credit x 10 = 20.
ITEM_SPEND = 3000 + 1800 + 900 + 2 * SPLICE_RATE
TOTAL_SPEND = ITEM_SPEND + FEES["Semrush"]
IDS: dict = {}


def check(condition, message):
    if not condition:
        raise AssertionError(message)


def close(a, b, message, tol=0.01):
    check(abs(float(a) - float(b)) <= tol, f"{message}: {a} != {b}")


def at(day, hour=6):
    return datetime(2026, 10, day, hour, 0, 0)    # UTC, well inside the IST day


def seed():
    with SessionLocal() as db:
        for name in ("Asha", "Bob"):
            u = User(email=f"{name.lower()}@example.com", name=name, hashed_password="x", department="DESIGN",
                     is_active=True, is_deleted=False, is_admin=False, employee_id=f"E-{name}")
            db.add(u)
            db.flush()
            IDS[name] = u.id

        def tool_with_account(name, slug, fee):
            t = ITPortalTool(name=name, slug=slug, website_url=f"https://{slug}.example")
            db.add(t)
            db.flush()
            c = ITPortalToolCredential(tool_id=t.id, scope="company", tool_cost=fee, renewal_type="MONTHLY",
                                       is_active=True, auto_renew=False)
            db.add(c)
            db.flush()
            return c.id

        for name in ("ChatGPT", "Claude", "Grammarly", "Semrush"):
            IDS[name] = tool_with_account(name, name.lower(), FEES[name])
        IDS["Splice"] = tool_with_account("Splice", "splice", None)
        db.add(ToolCreditRate(credential_id=IDS["Splice"], rate_per_credit=SPLICE_RATE,
                              effective_from=date(2026, 9, 1), currency="INR"))

        def conversation(provider, owner, prompts, key, url=None):
            record = ConversationRecord(provider=provider, provider_conversation_id=key, owner_user_id=owner,
                                        title=f"{provider} chat", conversation_url=url, created_at=at(2))
            db.add(record)
            db.flush()
            for i, text in enumerate(prompts):
                db.add(ConversationPrompt(conversation_id=record.id, sequence_index=i, prompt_text=text,
                                          prompt_timestamp=at(3 + i), created_at=at(3 + i)))

        conversation("chatgpt", IDS["Asha"], ["Write a tagline", "Shorter please"], "c1")
        conversation("chatgpt", IDS["Bob"], ["Summarise this brief"], "c2")
        conversation("claude", IDS["Bob"], ["Review this script"], CLAUDE_ID, url=CLAUDE_API_URL)
        db.add(GrammarlyDocSession(owner_user_id=IDS["Asha"], credential_id=IDS["Grammarly"], session_key="g1",
                                   doc_title="Pitch deck copy", started_at=at(5), status="closed", created_at=at(5)))
        for i in range(2):
            db.add(SpliceDownload(owner_user_id=IDS["Bob"], credential_id=IDS["Splice"], asset_title=f"Kick {i}",
                                  downloaded_at=at(6), created_at=at(6)))
        db.commit()


def sheet(db, **filters):
    return build_live_sheet(db, start=START, end=END, **filters)


def test_all_tools_counted():
    with SessionLocal() as db:
        data = sheet(db)
    kpis, tools = data["kpis"], {t["tool"]: t for t in data["sheets"]["tools"]}
    close(kpis["totalSpendRupees"], TOTAL_SPEND, "total spend = every fee + pay-per-use")
    close(kpis["itemSpendRupees"], ITEM_SPEND, "spend on items")
    check(kpis["totalGenerations"] == 3 + 1 + 1 + 2, f"items counted: {kpis['totalGenerations']}")
    check(tools["ChatGPT"]["generations"] == 3, "Claude prompts must not be counted as ChatGPT")
    check(tools["Claude"]["generations"] == 1, "Claude prompts counted")
    close(tools["Splice"]["totalSpendRupees"], 2 * SPLICE_RATE, "Splice: credits x rate")
    close(tools["Grammarly"]["totalSpendRupees"], FEES["Grammarly"], "Grammarly fee shared by its documents")
    semrush = tools.get("Semrush")
    check(semrush is not None and semrush["tracked"] is False, "a paid, untracked tool is listed")
    close(semrush["totalSpendRupees"], FEES["Semrush"], "Semrush fee")
    close(sum(t["totalSpendRupees"] for t in tools.values()), TOTAL_SPEND, "Tools sheet adds up to total spend")
    print("ok  ChatGPT, Claude, Grammarly, Splice and an untracked paid tool all reach total spend")


def test_people_and_filters():
    with SessionLocal() as db:
        data = sheet(db)
        users = {u["userName"]: u for u in data["sheets"]["users"]}
        close(users["Asha"]["spendRupees"], 2 * 1000 + FEES["Grammarly"], "Asha: 2 ChatGPT prompts + 1 document")
        close(users["Bob"]["spendRupees"], 1000 + FEES["Claude"] + 2 * SPLICE_RATE, "Bob: ChatGPT + Claude + Splice")
        options = data["options"]["tools"]
        check({"ChatGPT", "Claude", "Grammarly", "Splice"} <= set(options), f"tool filter lists new tools: {options}")
        check("Semrush" not in options, "an untracked tool has nothing to filter on")
        claude_only = sheet(db, tool="Claude")
        check(claude_only["kpis"]["totalGenerations"] == 1, "Claude filter shows only Claude prompts")
        check({r["tool"] for r in claude_only["sheets"]["log"]} == {"Claude"}, "Claude filter log")
        asha_only = sheet(db, user_id=IDS["Asha"])
        check(asha_only["kpis"]["totalGenerations"] == 3, "a user filter applies to conversation tools too")
    print("ok  per-person spend, tool filter list and filters cover the new tools")


def test_workbooks_export():
    import tempfile

    from openpyxl import load_workbook

    from utils.credit_report import ReportFilters, generate_report

    with SessionLocal() as db, tempfile.TemporaryDirectory() as tmp:
        for version in ("full", "management"):
            path = os.path.join(tmp, f"{version}.xlsx")
            generate_report(db, ReportFilters.build(start=START, end=END), path, version=version)
            text = " ".join(str(c.value) for ws in load_workbook(path).worksheets
                            for row in ws.iter_rows() for c in row if c.value is not None)
            for name in ("ChatGPT", "Claude", "Grammarly", "Splice"):
                check(name in text, f"{version} workbook mentions {name}")
    print("ok  full and management Excel exports include the new tools")


def test_chat_links():
    with SessionLocal() as db:
        rows = {r["tool"]: r for r in sheet(db)["sheets"]["log"] if r["tool"] in ("Claude", "ChatGPT")}
        claude, chatgpt = rows["Claude"], rows["ChatGPT"]
        check(claude["hasOutput"] and chatgpt["hasOutput"], "Claude and ChatGPT prompts link to their chat")
        check(resolve_output_url(db, "Claude", int(claude["recordId"])) == f"https://claude.ai/chat/{CLAUDE_ID}",
              "a captured Claude API URL becomes the chat page")
        check(resolve_output_url(db, "ChatGPT", int(chatgpt["recordId"])).startswith("https://chatgpt.com/c/"),
              "a ChatGPT chat with no stored URL links by its id")
        t = chat_transcript(db, "Claude", int(claude["recordId"]))
        check(t and [m["role"] for m in t["messages"]] == ["user"] and t["messages"][0]["target"],
              "the transcript returns the conversation with the clicked prompt marked")
    print("ok  Claude / ChatGPT prompts open their real conversation")


def main():
    seed()
    test_all_tools_counted()
    test_people_and_filters()
    test_workbooks_export()
    test_chat_links()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
