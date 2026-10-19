"""
Credit Consumption Report - management view checks (rupees, Summary, Team &
Clients, the management version). Run: python tests/credit_report_management_smoke.py

Seeds an in-memory SQLite database whose rupee answers are worked out by
hand below, then checks:

* costing: pay-per-use = credits x the account's rate (a rate-less account
  borrows its tool's rate, flagged); subscriptions = the fee pro-rated by
  days, shared by size; unused fees are "Paid, not used";
* Kling jobs recorded twice count once (starter kept; history-only jobs go
  to the account's only starter, flagged), rows with no job ID are listed;
* Freepik uploads are not generations; ElevenLabs free preview models are ₹0;
* ₹ totals agree: Summary = Team & Clients = detail sheets = Generation Log,
  overall and per tool;
* the previous period, recomputed here from the raw rows, matches the report;
* every key point / needs-attention sentence appears exactly when its
  condition holds (fixtures where it must and must not);
* the management version holds no prompts, outputs or test accounts;
* the Summary is set to print on one A4 landscape page (and, with
  LibreOffice installed, its PDF has exactly one page);
* the full version still passes the workbook's link audit.
"""

from __future__ import annotations

import os
import re
import sys
import tempfile
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
os.environ.setdefault("DATABASE_URL", "postgresql://placeholder:placeholder@localhost:5432/placeholder")
os.environ.setdefault("ARCHIVE_DATABASE_URL", os.environ["DATABASE_URL"])
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from openpyxl import Workbook, load_workbook  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

from models_new import (Base, GenerationClient, GenerationRecord, ITPortalTool, ITPortalToolCredential,  # noqa: E402
                        ITPortalToolUsageEvent, ToolCreditRate, User, UserFeatureAccess, UserRole)
from providers.elevenlabs.models import ElevenlabsGeneration  # noqa: E402
from providers.envato.models import EnvatoGeneration  # noqa: E402
from providers.epidemicsound.models import EpidemicAdaptation  # noqa: E402
from providers.flow.models import FlowGeneration  # noqa: E402
from providers.freepik.models import FreepikGeneration  # noqa: E402
from providers.heygen.models import HeygenGeneration  # noqa: E402
from providers.higgsfield.models import HiggsfieldGeneration  # noqa: E402
from providers.suno.models import SunoGeneration  # noqa: E402
from utils.credit_report import (CHARGED, NO_CLIENT, ReportFilters, generate_report, load_groups,  # noqa: E402
                                 soffice_path, to_pdf)
from utils.credit_report import insights as I  # noqa: E402
from utils.credit_report import management as M  # noqa: E402
from utils.credit_report import workbook as W  # noqa: E402
from utils.credit_report import workbook_summary as WS  # noqa: E402
from utils.credit_report.costing import period_fee  # noqa: E402
from utils.credit_report.display import Budgets, DisplayConfig  # noqa: E402
from utils.credit_report.facts import DayGroup, FactGroup  # noqa: E402

engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
TABLES = [User, UserRole, UserFeatureAccess, GenerationClient, ITPortalTool, ITPortalToolCredential,
          ITPortalToolUsageEvent, ToolCreditRate, GenerationRecord, FreepikGeneration, EnvatoGeneration,
          HeygenGeneration, HiggsfieldGeneration, ElevenlabsGeneration, FlowGeneration, SunoGeneration,
          EpidemicAdaptation]
Base.metadata.create_all(bind=engine, tables=[m.__table__ for m in TABLES])

START, END = "2026-10-01", "2026-10-14"
SECRET_PROMPT = "SECRET-PROMPT-XYZ"
SECRET_OUTPUT = "https://freepik.example/secret-output.png"
IDS: dict = {}

# --- the answers, by hand ----------------------------------------------------
# Fees (monthly) pro-rated to 1-14 Oct = 14/31 of October.
OCT = 14 / 31
SEP = 14 / 30                     # previous period: 17-30 Sep
FEE = {"F1": 3100.0, "F2": None, "E": 3100.0, "S": 62.0, "H": 3100.0, "V": 3100.0}
OTHER_FEE = 3100.0               # ChatGPT: an AI subscription whose usage is not tracked
LAPSED_FEE = 3100.0              # Higgsfield P: unused, renewal date passed -> to confirm, never counted
RATE = 0.5                        # Kling account KA; KB has none and borrows it
# Kling (pay-per-use): J1 100 cr (Asha) + J2 200 cr (Bob, rate assumed) + J3 60 cr (credited to Asha)
# + no-job-ID 40 cr (Cara); the second copies of J1 and J3 are removed.
KLING = {"Asha": (100 + 60) * RATE, "Bob": 200 * RATE, "Cara": 40 * RATE}
# Freepik F1: Asha 30 charged + 10 unlimited-plan estimate share the fee (Tester's item takes none).
F1 = FEE["F1"] * OCT
# ElevenLabs E: Bob 100 counted + 300 characters (counter stuck) share it; the preview model gets ₹0.
E = FEE["E"] * OCT
S = FEE["S"] * OCT                # Suno: Cara's 2 songs share it
H = FEE["H"] * OCT                # Higgsfield: Cara's 10 items share it
V = FEE["V"] * OCT                # Envato: no use -> paid, not used
PAY_PER_USE = sum(KLING.values())
ALLOCATED = PAY_PER_USE + F1 + E + S + H
NOT_USED = V
OTHER = OTHER_FEE * OCT
TO_CONFIRM = LAPSED_FEE * OCT
TOTAL = ALLOCATED + NOT_USED + OTHER
ESTIMATED = 200 * RATE + E * 300 / 400 + S + H
PER_TOOL = {"Kling": PAY_PER_USE, "Freepik": F1, "ElevenLabs": E, "Suno": S, "Higgsfield": H}
ITEMS = 4 + 3 + 3 + 2 + 10        # Kling, Freepik (2 Asha + 1 Bob), ElevenLabs, Suno, Higgsfield


def _assert(condition, message):
    if not condition:
        raise AssertionError(message)


def _close(a, b, message, tol=0.01):
    _assert(abs(float(a) - float(b)) <= tol, f"{message}: {a} != {b}")


def seed():
    with SessionLocal() as db:
        def user(name, dept, admin=False):
            u = User(email=f"{name.lower()}@example.com", name=name, hashed_password="x", department=dept,
                     is_active=True, is_deleted=False, is_admin=admin, employee_id=f"E-{name}")
            db.add(u)
            db.flush()
            IDS[name] = u.id
            return u.id

        asha, bob, cara = user("asha", "DESIGN"), user("Bob", "VIDEO"), user("Cara", "VIDEO")
        admin = user("Admin", "ADMIN", admin=True)
        tester = user("Tester", "APPPLICANT")           # a test account by department (config)

        def tool(name, slug):
            t = ITPortalTool(name=name, slug=slug, website_url=f"https://{slug}.example")
            db.add(t)
            db.flush()
            return t.id

        def cred(key, tool_id, fee, renewal=None, active=True):
            c = ITPortalToolCredential(tool_id=tool_id, scope="company", tool_cost=fee, renewal_type="MONTHLY",
                                       renewal_date=renewal, is_active=active, auto_renew=False)
            db.add(c)
            db.flush()
            IDS[key] = c.id
            return c.id

        kling, freepik, eleven = tool("Kling", "kling"), tool("freepik", "freepik"), tool("ElevenLabs", "elevenlabs")
        suno, higgs, envato = tool("Suno", "suno"), tool("Higgsfield", "higgsfield"), tool("Envato", "envato")
        ka, kb = cred("KA", kling, 12600), cred("KB", kling, None)
        f1, f2 = cred("F1", freepik, FEE["F1"]), cred("F2", freepik, FEE["F2"])
        e = cred("E", eleven, FEE["E"])
        s = cred("S", suno, FEE["S"])
        h = cred("H", higgs, FEE["H"], renewal=date(2026, 10, 10))
        v = cred("V", envato, FEE["V"])
        cred("P", higgs, LAPSED_FEE, renewal=date(2026, 9, 1))
        cred("G", tool("CHAT GPT", "chat-gpt"), OTHER_FEE, renewal=date(2026, 11, 5))
        db.add_all([
            ToolCreditRate(credential_id=ka, package_credits=1000, package_rupees=500, rate_per_credit=RATE,
                           effective_from=date(2026, 9, 1), currency="INR"),
            # The seeded global placeholder must never cost a Kling account.
            ToolCreditRate(credential_id=None, rate_per_credit=1.0, effective_from=date(2026, 7, 1), currency="INR"),
        ])

        def ev(uid, credits, day, hour, account, job, source, client=None, prompt=None):
            row = ITPortalToolUsageEvent(tool_id=kling, user_id=uid, event_type="generation", event_date=day,
                                         credits_burned=credits, status="settled", credential_id=account,
                                         generation_id=job, source=source, linked_client_name=client,
                                         prompt_text=prompt, created_at=datetime(day.year, day.month, day.day, hour))
            db.add(row)
            db.flush()
            return row.id

        IDS["J1"] = ev(asha, 100, date(2026, 10, 3), 6, ka, "J1", "xhr_response", "RMW", "j1 prompt")
        IDS["J1copy"] = ev(bob, 100, date(2026, 10, 5), 6, ka, "J1", "asset_history_discovered")
        IDS["J2"] = ev(bob, 200, date(2026, 10, 4), 6, kb, "J2", "generate_intent", "Acme")
        IDS["J3"] = ev(bob, 60, date(2026, 10, 6), 6, ka, "J3", "asset_history_discovered")
        IDS["J3copy"] = ev(cara, 60, date(2026, 10, 6), 7, ka, "J3", "asset_history_discovered")
        # Found by a history scan (not a Generate click), so KA still has one starter: Asha.
        IDS["noid"] = ev(cara, 40, date(2026, 10, 7), 6, ka, None, "asset_history_discovered", "Applicant")
        ev(admin, 70, date(2026, 10, 3), 8, ka, "JA", "xhr_response")
        ev(asha, 50, date(2026, 9, 20), 6, ka, "J0", "xhr_response")                 # previous period

        at = datetime(2026, 10, 2, 9)
        db.add_all([
            FreepikGeneration(creation_id="f1", owner_user_id=asha, credential_id=f1, credits_charged=30,
                              linked_client_name="Acme", provider_created_at=at, prompt=SECRET_PROMPT,
                              raw_url=SECRET_OUTPUT, status="completed", tool_name="text-to-image"),
            FreepikGeneration(creation_id="f2", owner_user_id=asha, credential_id=f1, credits_charged=0,
                              credits_estimated=10, unlimited_credits=10, linked_client_name="Acme",
                              provider_created_at=at, status="completed", tool_name="text-to-image"),
            FreepikGeneration(creation_id="f3", owner_user_id=asha, credential_id=f1, credits_charged=0,
                              provider_created_at=at, status="completed", tool="upload-reference"),
            FreepikGeneration(creation_id="f4", owner_user_id=tester, credential_id=f1, credits_charged=25,
                              provider_created_at=at, status="completed", prompt="tester prompt"),
            FreepikGeneration(creation_id="f5", owner_user_id=bob, credential_id=f2, credits_charged=20,
                              provider_created_at=datetime(2026, 10, 3, 9), status="completed"),
            FreepikGeneration(creation_id="f0", owner_user_id=asha, credential_id=f1, credits_charged=15,
                              provider_created_at=datetime(2026, 9, 21, 9), status="completed"),
            ElevenlabsGeneration(provider_creation_id="e1", owner_user_id=bob, credential_id=e, credits_used=100,
                                 prompt_length=100, provider_created_at=datetime(2026, 10, 8, 9),
                                 metadata_json={"model_id": "eleven_v3"}),
            ElevenlabsGeneration(provider_creation_id="e2", owner_user_id=bob, credential_id=e, credits_used=0,
                                 prompt_length=300, provider_created_at=datetime(2026, 10, 8, 9),
                                 metadata_json={"model_id": "eleven_v3"}),
            ElevenlabsGeneration(provider_creation_id="e3", owner_user_id=bob, credential_id=e, credits_used=0,
                                 prompt_length=500, provider_created_at=datetime(2026, 10, 8, 9),
                                 metadata_json={"model_id": "eleven_v4_exp"}),
            SunoGeneration(provider_creation_id="s1", owner_user_id=cara, credential_id=s, status="complete",
                           provider_created_at=datetime(2026, 10, 9, 9)),
            SunoGeneration(provider_creation_id="s2", owner_user_id=cara, credential_id=s, status="complete",
                           provider_created_at=datetime(2026, 10, 9, 9)),
        ])
        db.add_all([HiggsfieldGeneration(generation_id=f"h{i}", owner_user_id=cara, credential_id=h,
                                         status="completed", provider_created_at=datetime(2026, 10, 10, 9))
                    for i in range(10)])
        db.add(EnvatoGeneration(item_uuid="v0", owner_user_id=bob, credential_id=v, provider_created_at=datetime(2026, 8, 1)))
        db.commit()


def _filters(db, start=START, end=END, **kw):
    return ReportFilters.build(start=start, end=end, exclude_test_accounts=True, **kw)


def _charged(groups):
    return [g for g in groups if g.charge_status == CHARGED]


# --------------------------------------------------------------------------- #
# Facts: rupees, Kling double recording, uploads, previews
# --------------------------------------------------------------------------- #
def test_costing_and_dedupe():
    with SessionLocal() as db:
        f = _filters(db)
        groups = _charged(load_groups(db, f))
        by_tool, by_user = {}, {}
        for g in groups:
            by_tool[g.tool] = by_tool.get(g.tool, 0.0) + g.rupees
            by_user[g.user_name] = by_user.get(g.user_name, 0.0) + g.rupees
        for tool, expected in PER_TOOL.items():
            _close(by_tool.get(tool, 0.0), expected, f"{tool} ₹")
        _close(sum(g.rupees for g in groups), ALLOCATED, "₹ on items")
        _close(sum(g.rupees_est for g in groups), ESTIMATED, "estimated ₹")
        _assert(sum(g.generations for g in groups) == ITEMS, f"items {sum(g.generations for g in groups)}")
        _assert("Tester" not in by_user and "Admin" not in by_user, "test & admin accounts left out")
        _assert("Asha" in by_user and "asha" not in by_user, "names cleaned ('asha' -> 'Asha')")
        _assert({g.department for g in groups} == {"Design", "Video"}, "departments cleaned (DESIGN -> Design)")
        _close(sum(g.rupees for g in groups if g.user_name == "Asha" and g.tool == "Kling"), KLING["Asha"],
               "Asha's Kling: J1 + J3 (history-only job credited to the account's only starter)")
        clients = {g.client for g in groups}
        _assert("Internal (RMW)" in clients and "RMW" not in clients and "Applicant" not in clients,
                f"client display names {clients}")
        k = f.kling
        _assert(len(k.jobs) == 2 and k.removed_rows == 2 and k.removed_credits == 160, "two jobs recorded twice")
        _assert({r.id for j in k.jobs for r in j.removed} == {IDS["J1copy"], IDS["J3copy"]}, "the copies are removed")
        j3 = next(j for j in k.jobs if j.job == "J3")
        _assert(j3.owner_uncertain and j3.owner_id == IDS["asha"], "J3: owner uncertain, credited to Asha")
        j1 = next(j for j in k.jobs if j.job == "J1")
        _assert(not j1.owner_uncertain and j1.kept.id == IDS["J1"], "J1: the starter's (API) row is kept")
        _assert([r.id for r in k.no_job_id] == [IDS["noid"]], "the row with no job ID is listed")
        # Rate fallback: KB has no rate of its own -> KA's (the tool's) rate, flagged; never the ₹1 placeholder.
        rate, assumed = f.cost.rate_on("Kling", IDS["KB"], date(2026, 10, 4))
        _assert(rate == RATE and assumed, f"rate assumed for KB: {rate} {assumed}")
        _assert(f.cost.rate_on("Kling", IDS["KA"], date(2026, 10, 4)) == (RATE, False), "KA's own rate")
    _close(period_fee(3100, date(2026, 10, 1), date(2026, 10, 14)), 1400, "fee pro-rated by days")
    _close(period_fee(3100, date(2026, 9, 17), date(2026, 10, 14)), 3100 * 14 / 30 + 1400, "fee across two months")
    print("ok  costing: pay-per-use, subscription shares, Kling recorded twice, uploads, preview models")


# --------------------------------------------------------------------------- #
# Previous period, recomputed independently from the raw rows
# --------------------------------------------------------------------------- #
def _independent_previous(db) -> tuple:
    """17-30 Sep from the tables, by the written rules - not by the report's code."""
    start, end = date(2026, 9, 17), date(2026, 9, 30)
    utc_s, utc_e = datetime(2026, 9, 16, 18, 30), datetime(2026, 9, 30, 18, 30)
    kling = sum(float(e.credits_burned) * RATE for e in db.query(ITPortalToolUsageEvent)
                .filter(ITPortalToolUsageEvent.event_date >= start, ITPortalToolUsageEvent.event_date <= end,
                        ITPortalToolUsageEvent.user_id != IDS["Admin"]))
    fees = {k: (v or 0.0) * SEP for k, v in FEE.items()}
    other = OTHER_FEE * SEP              # counted; Higgsfield P (renewal passed, unused) is not
    used = {"F1": db.query(FreepikGeneration).filter(FreepikGeneration.provider_created_at >= utc_s,
                                                     FreepikGeneration.provider_created_at < utc_e).count()}
    allocated = kling + (fees["F1"] if used["F1"] else 0.0)
    not_used = sum(v for k, v in fees.items() if not (k == "F1" and used["F1"]))
    return kling + sum(fees.values()) + other, allocated, not_used, kling


def test_previous_period():
    # September 4,787 items vs August 4,765 is +0.46%: a real change, never "no change".
    for cur, prev, want in ((4787, 4765, "▲ <1%"), (4765, 4787, "▼ <1%"), (4787, 4787, "— no change"),
                            (112, 100, "▲ 12%"), (5, 0, "new"), (0, 0, "—")):
        _assert(M.change_text(cur, prev) == want, f"change_text({cur}, {prev}) = {M.change_text(cur, prev)}")
    with SessionLocal() as db:
        f = _filters(db)
        view = M.build_management(db, f, load_groups(db, f))
        total, allocated, not_used, ppu = _independent_previous(db)
        p = view.prev
        _assert((p.start, p.end) == (date(2026, 9, 17), date(2026, 9, 30)), f"previous period {p.label}")
        _close(p.total_spend, total, "previous total spend")
        _close(p.total.rupees, allocated, "previous spend on items")
        _close(p.not_used, not_used, "previous paid, not used")
        _close(p.pay_per_use, ppu, "previous pay-per-use")
        _close(view.cur.total_spend, TOTAL, "current total spend")
        _close(view.cur.not_used, NOT_USED, "current paid, not used")
        _close(view.cur.other_total, OTHER, "other AI subscriptions (usage not tracked)")
        _close(view.cur.to_confirm, TO_CONFIRM, "renewal passed and unused: to confirm, not counted")
        _assert([s.tool for s in view.cur.other_subscriptions] == ["ChatGPT"], "ChatGPT matched from CHAT GPT")
        _assert("ChatGPT" not in view.cur.tools and view.cur.cost_per_item("ChatGPT") is None,
                "untracked tools stay out of tools and cost per item")
        _assert(M.previous_period(date(2026, 9, 1), date(2026, 9, 30)) == (date(2026, 8, 1), date(2026, 8, 31)),
                "a whole month compares with the month before")
        _assert(M.last_full_month(date(2026, 10, 14)) == (date(2026, 9, 1), date(2026, 9, 30)), "default period")
        subs = {s.account: s for s in view.cur.subscriptions}
        _assert(subs[IDS["P"]].status == M.RENEWAL_PASSED and subs[IDS["P"]].not_used == 0
                and abs(subs[IDS["P"]].to_confirm - TO_CONFIRM) < 0.01, "lapsed account held back")
        _assert(subs[IDS["V"]].status == M.NOT_USED and subs[IDS["F2"]].status == M.NO_FEE
                and subs[IDS["H"]].status == M.RENEWAL_PASSED and subs[IDS["S"]].status == M.IN_USE
                and subs[IDS["E"]].status == M.LOW_USE, {k: s.status for k, s in subs.items()})
        _close(subs[IDS["S"]].per_item, S / 2, "Suno ₹ per item")
    print("ok  previous period recomputed from the raw rows matches; subscription statuses")


# --------------------------------------------------------------------------- #
# Sentences: each appears exactly when its condition holds
# --------------------------------------------------------------------------- #
def _view(groups, *, prev_groups=(), days=None, subs=(), issues=None, filtered=False, duplicates=0,
          start=date(2026, 10, 1), end=date(2026, 10, 7)):
    display = DisplayConfig.load()
    cur = M.fill_period(M.PeriodView(start=start, end=end, filtered=filtered), list(groups),
                        [DayGroup(1, "Kling", NO_CLIENT, d, CHARGED, 1, 0.0, r) for d, r in (days or {}).items()])
    cur.subscriptions = list(subs)
    cur.issues = issues or {"rate_assumed": (0, 0.0), "estimated": (0, 0.0), "free_preview": (0, 0.0)}
    prev = M.fill_period(M.PeriodView(start=start - (end - start) - timedelta(days=1), end=start - timedelta(days=1)),
                         list(prev_groups))
    return M.ManagementView(cur=cur, prev=prev, budgets=Budgets(), display=display,
                            generated_at=datetime(2026, 10, 15), duplicates=duplicates, duplicate_rupees=50.0)


def _g(uid, name, tool, rupees, *, dept="Video", client="Acme", items=1):
    return FactGroup(user_id=uid, user_name=name, employee_id="", user_deleted=False, department=dept, tool=tool,
                     client=client, charge_status=CHARGED, month="2026-10", generations=items, credits=rupees,
                     rupees=rupees, rupees_est=0.0)


def _sub(label, fee, items, status, allocated=None):
    return M.SubRow(tool=label.split(" #")[0], account=1, label=label, monthly_fee=fee, period_fee=fee, items=items,
                    allocated=fee if allocated is None else allocated, not_used=0.0 if items else fee,
                    last_used=None, renewal_date=None, status=status)


def test_sentences():
    def has(rule, view, text=None):
        out = rule(view)
        return out is not None and (text is None or text in (out.sentence if isinstance(out, M.Attention) else out))

    # Dominant tool: 60% -> yes; 40% -> no.
    _assert(has(I.dominant_tool, _view([_g(1, "A", "Kling", 60), _g(2, "B", "Suno", 40)]), "Kling is 60%"), "tool 60%")
    _assert(not has(I.dominant_tool, _view([_g(1, "A", "Kling", 40), _g(2, "B", "Suno", 30),
                                            _g(3, "C", "Freepik", 30)])), "tool 40%: no sentence")
    # People: 4 people, top 2 = 70% -> yes; 3 people -> no; top 2 = 30% -> no.
    four = [_g(1, "A", "Kling", 40), _g(2, "B", "Kling", 30), _g(3, "C", "Kling", 15), _g(4, "D", "Kling", 15)]
    _assert(has(I.people_concentration, _view(four), "2 people (A and B) account for 70%"), "people 70%")
    _assert(not has(I.people_concentration, _view(four[:3])), "3 people: no sentence")
    even = [_g(i, f"P{i}", "Kling", 10) for i in range(1, 8)]
    _assert(not has(I.people_concentration, _view(even)), "even spread: no sentence")
    # Change vs previous: +25% -> yes; +10% -> no; nothing before -> no.
    _assert(has(I.change_vs_previous, _view([_g(1, "A", "Kling", 125)], prev_groups=[_g(1, "A", "Kling", 100)]),
                "rose 25%"), "change +25%")
    _assert(not has(I.change_vs_previous, _view([_g(1, "A", "Kling", 110)], prev_groups=[_g(1, "A", "Kling", 100)])),
            "change +10%: no sentence")
    _assert(not has(I.change_vs_previous, _view([_g(1, "A", "Kling", 110)])), "no previous spend: no sentence")
    # Unused subscriptions: 50% of total -> yes; 5% -> no; filtered report -> no.
    idle = _sub("Envato #9", 100.0, 0, M.NOT_USED, allocated=0.0)
    _assert(has(I.unused_subscriptions, _view([_g(1, "A", "Kling", 100)], subs=[idle]), "Envato #9"), "idle 50%")
    small = _sub("Envato #9", 5.0, 0, M.NOT_USED, allocated=0.0)
    _assert(not has(I.unused_subscriptions, _view([_g(1, "A", "Kling", 100)], subs=[small])), "idle 5%: no")
    _assert(not has(I.unused_subscriptions, _view([_g(1, "A", "Kling", 100)], subs=[idle], filtered=True)),
            "filtered: no")
    # Clients: 3 clients, top 2 = 80% -> yes; only 2 clients -> no.
    clients = [_g(1, "A", "Kling", 50, client="MPF"), _g(1, "A", "Kling", 30, client="Vedvan"),
               _g(1, "A", "Kling", 20, client="Onyx")]
    _assert(has(I.client_concentration, _view(clients), "MPF and Vedvan together take 80%"), "clients 80%")
    _assert(not has(I.client_concentration, _view(clients[:2])), "2 clients: no sentence")
    # Internal share: 30% -> yes; 10% -> no.
    _assert(has(I.internal_share, _view([_g(1, "A", "Kling", 30, client="Internal (RMW)"),
                                         _g(1, "A", "Kling", 70)]), "Internal work is 30%"), "internal 30%")
    _assert(not has(I.internal_share, _view([_g(1, "A", "Kling", 10, client="Internal (RMW)"),
                                             _g(1, "A", "Kling", 90)])), "internal 10%: no")
    # Department on one tool: 90% -> yes; 70% -> no.
    _assert(has(I.department_dependence, _view([_g(1, "A", "Kling", 90, dept="Gen AI"),
                                                _g(1, "A", "Suno", 10, dept="Gen AI")]), "Gen AI 90% Kling"),
            "dependence 90%")
    _assert(not has(I.department_dependence, _view([_g(1, "A", "Kling", 70, dept="Gen AI"),
                                                    _g(1, "A", "Suno", 30, dept="Gen AI")])), "dependence 70%: no")
    # Spike day: one day 40% of 7 -> yes; even days -> no; 3 active days -> no.
    days = {f"2026-10-0{d}": (40.0 if d == 1 else 10.0) for d in range(1, 8)}
    _assert(has(I.spike_day, _view([_g(1, "A", "Kling", 100)], days=days), "Spend on 1 Oct was 40%"), "spike day")
    _assert(not has(I.spike_day, _view([_g(1, "A", "Kling", 70)], days={f"2026-10-0{d}": 10.0 for d in range(1, 8)})),
            "even days: no")
    _assert(not has(I.spike_day, _view([_g(1, "A", "Kling", 70)], days={"2026-10-01": 50.0, "2026-10-02": 10.0,
                                                                         "2026-10-03": 10.0})), "3 days: no")
    # Needs attention.
    costly = _sub("Higgsfield #5", 500.0, 5, M.LOW_USE)
    v = _view([_g(1, "A", "Kling", 100, client=NO_CLIENT)], subs=[idle, costly],
              issues={"rate_assumed": (3, 60.0), "estimated": (0, 0.0), "free_preview": (0, 0.0)}, duplicates=2)
    _assert(has(I.idle_subscriptions, v, "Envato #9") and has(I.costly_per_item, v, "Higgsfield #5 ₹100 per item"),
            "idle and costly subscriptions")
    _assert(has(I.assumed_rates, v, "₹60") and has(I.untagged_client, v, "no client") and has(I.double_charges, v),
            "assumed rates, no client, double charges")
    items = I.attention_items(v)
    _assert(len(items) == I.MAX_ATTENTION and items[0].text.startswith("₹100 paid for subscriptions not used")
            and items[1].text.startswith("Subscriptions costing over"), [a.text for a in items])
    quiet = _view([_g(1, "A", "Kling", 100)])
    _assert(I.attention_items(quiet) == [] and not has(I.untagged_client, quiet) and not has(I.double_charges, quiet),
            "nothing to raise: no attention items")
    _assert(not has(I.costly_per_item, _view([_g(1, "A", "Kling", 100)], subs=[_sub("Suno #1", 20.0, 2, M.IN_USE)])),
            "₹10 per item: not costly")
    print("ok  sentences: every key point and attention item appears only when its condition holds")


# --------------------------------------------------------------------------- #
# Workbooks
# --------------------------------------------------------------------------- #
def _sheet_text(path: str) -> str:
    with zipfile.ZipFile(path) as z:
        return " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))


def _table(ws, name):
    t = ws.tables[name]
    first, last = t.ref.split(":")
    hr, lr = int(re.sub(r"[A-Z]", "", first)), int(re.sub(r"[A-Z]", "", last))
    headers = [c.value for c in ws[hr]][: len(t.tableColumns)]
    return headers, list(ws.iter_rows(min_row=hr + 1, max_row=lr, max_col=len(headers), values_only=True))


def _find_row(ws, label, col=1):
    for row in ws.iter_rows():
        if row[col - 1].value == label:
            return row
    raise AssertionError(f"{ws.title}: no row '{label}'")


def test_full_version_totals_and_links():
    import credit_report_smoke as base     # the workbook's own link audit

    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        with SessionLocal() as db:
            model = generate_report(db, _filters(db), path, dashboard_url="https://dash.example.com")
        wb = load_workbook(path)
        _assert(wb.sheetnames[:4] == [WS.SUMMARY, WS.TEAM, WS.GLOSSARY, WS.DETAILS], f"order {wb.sheetnames[:5]}")
        # Summary card = Team & Clients total = detail sheets + paid-not-used = log + paid-not-used.
        summary, team = wb[WS.SUMMARY], wb[WS.TEAM]
        _close(summary["A7"].value, round(TOTAL), "Summary total spend card", tol=0.5)
        tracked_row = _find_row(team, "Total (tracked)")
        _close(tracked_row[2].value, round(TOTAL - OTHER), "Team & Clients total, tracked tools", tol=0.5)
        lines = [c.value for row in team.iter_rows() for c in row if isinstance(c.value, str)]
        _assert(any(t.startswith("Total spend: ") and M.inr(TOTAL) in t for t in lines), "Team & Clients total spend line")
        _assert(any(t.startswith(WS.OTHER_AI_LINE) and M.inr(OTHER) in t for t in lines), "Other AI line")
        _assert("Other AI, usage not tracked" not in lines, "Other AI is not a department row")
        _close(_find_row(team, "To confirm, not counted")[3].value, round(TO_CONFIRM), "to confirm row", tol=0.5)
        _assert("To confirm" in (summary["A15"].value or ""), "Summary To confirm line")
        not_used_row = _find_row(team, "Subscriptions paid, not used")
        _close(not_used_row[2].value, round(NOT_USED), "Team & Clients paid, not used", tol=0.5)
        dept_h, dept_rows = _table(wb[W.DEPT], "tblDepartments")
        detail = {
            "By Department": sum(r[dept_h.index(W.RUPEES)] for r in dept_rows),
            "By User": sum(r[_table(wb[W.USER], "tblUsers")[0].index(W.RUPEES)] for r in _table(wb[W.USER], "tblUsers")[1]),
            "By Tool": sum(r[_table(wb[W.TOOL], "tblTools")[0].index(W.RUPEES)] for r in _table(wb[W.TOOL], "tblTools")[1]),
            "By Client": sum(r[_table(wb[W.CLIENT], "tblClients")[0].index(W.RUPEES)]
                             for r in _table(wb[W.CLIENT], "tblClients")[1]),
            "Generation Log": sum(model.log_stats.rupees_by_tool.values()),
        }
        lh, lrows = _table(wb[W.LOG], "tblGenerationLog")
        detail["Generation Log cells"] = sum(r[lh.index(W.RUPEES)] or 0 for r in lrows
                                             if r[lh.index("Charge status")] == CHARGED)
        for name, value in detail.items():
            _close(value, ALLOCATED, f"{name} ₹ = spend on items", tol=0.05)
        _close(round(ALLOCATED) + round(NOT_USED) + round(OTHER), round(TOTAL),
               "items + paid, not used + other AI = total", tol=1.5)
        # The same money on By Date, Monthly Trend, Month Drill-down (and Period Explorer, recalculated).
        by_date_total = _find_row(wb[W.BY_DATE], "Total")
        _close(by_date_total[2].value, ALLOCATED, "By Date ₹ total", tol=0.05)
        trh, trows = _table(wb[W.TREND], "tblTrendToolRupees")
        _close(sum(r[trh.index("Total")] for r in trows), ALLOCATED, "Monthly Trend ₹ by tool", tol=0.05)
        sh, srows = _table(wb[W.TREND], "tblTrendSummary")
        _close(sum(r[sh.index(W.RUPEES)] for r in srows), ALLOCATED, "Monthly Trend month summary ₹", tol=0.05)
        drill = [row for row in wb[W.DRILL].iter_rows(values_only=True) if row and len(row) > 10 and row[9] == W.RUPEES]
        _close(sum(row[10] for row in drill), ALLOCATED, "Month Drill-down ₹", tol=0.05)
        if soffice_path():
            recalced = base._recalc(path)[W.PERIOD]
            _close(recalced["B23"].value, ALLOCATED, "Period Explorer ₹ (recalculated)", tol=0.05)
        # Per tool: log = By Tool = Team & Clients (tool spend less its unused fees).
        th, trows = _table(wb[W.TOOL], "tblTools")
        by_tool = {W.N.strip_decoration(base._display(r[0])): r[th.index(W.RUPEES)] for r in trows}
        tools_row = _find_row(team, "Tools")
        team_tools = {}
        for row in team.iter_rows(min_row=tools_row[0].row + 2):
            if row[0].value in (None, ""):
                break
            team_tools[row[0].value] = row[2].value
        for tool, expected in PER_TOOL.items():
            _close(model.log_stats.rupees_by_tool.get(tool, 0.0), expected, f"log ₹ {tool}", tol=0.05)
            _close(by_tool[tool], expected, f"By Tool ₹ {tool}", tol=0.05)
            _close(team_tools[tool], round(expected), f"Team & Clients ₹ {tool}", tol=0.5)
        _close(team_tools["Envato"], round(V), "Envato: its whole fee, not used", tol=0.5)
        # Data Quality lists the removed Kling rows with both owners, and the row with no job ID.
        dq = wb[W.QUALITY]
        jh, jrows = _table(dq, "tblKlingRecordedTwice")
        _assert(len(jrows) == 2 and {r[jh.index("Removed rows recorded by")] for r in jrows} == {"Bob", "Cara"},
                f"recorded-twice list {jrows}")
        _assert({r[jh.index("Credited to")] for r in jrows} == {"Asha"}, "both kept rows credited to Asha")
        _assert(len(_table(dq, "tblKlingNoJobId")[1]) == 1, "no-job-ID list")
        # The full version still passes the workbook's own link checks.
        base._validate_links(wb)
        base._validate_navigation(wb, expect_links=False)
        base._audit_landings(wb)
        base._audit_graph(wb)
        missing = base._reachability(wb)
        _assert(not missing, f"rows not reachable within 4 clicks: { {k: v[:8] for k, v in missing.items()} }")
        # Management sheets: whole percentages, no raw decimals.
        for ws in (summary, team):
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, float) and c.number_format in ("General", None):
                        raise AssertionError(f"{ws.title}!{c.coordinate}: raw decimal {c.value}")
                    if c.number_format == "0%":
                        _assert(abs(c.value * 100 - round(c.value * 100)) < 1e-9, f"{ws.title}!{c.coordinate} rounded %")
    finally:
        os.remove(path)
    print("ok  full version: ₹ totals agree (Summary = Team & Clients = detail sheets = log), link audit passes")


def test_management_version_and_print():
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        with SessionLocal() as db:
            generate_report(db, _filters(db), path, version="management")
        wb = load_workbook(path)
        visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
        _assert(visible == list(WS.MANAGEMENT_SHEETS), f"management sheets {wb.sheetnames}")
        # Excel applies a cell's number format (₹, dates) only when its style says so.
        with zipfile.ZipFile(path) as z:
            styles = z.read("xl/styles.xml").decode("utf-8")
        xfs = re.findall(r"<xf\b[^>]*>", re.search(r"<cellXfs\b[^>]*>(.*?)</cellXfs>", styles, re.S).group(1))
        unflagged = [x for x in xfs if 'numFmtId="0"' not in x and 'applyNumberFormat="1"' not in x]
        _assert(xfs and not unflagged, f"management version: {len(unflagged)} styles without applyNumberFormat")
        text = _sheet_text(path)
        for secret in (SECRET_PROMPT, SECRET_OUTPUT, "j1 prompt", "tester prompt", "Tester", "Admin"):
            _assert(secret not in text, f"management version leaks '{secret}'")
        _assert("Asha" in text and "₹" in text, "names and rupees are there")
        summary = wb[WS.SUMMARY]
        ps = summary.page_setup
        _assert(ps.orientation == "landscape" and str(ps.paperSize) == "9" and ps.fitToWidth == 1 and ps.fitToHeight == 1
                and summary.sheet_properties.pageSetUpPr.fitToPage, "Summary: A4 landscape, fit to 1 page")
        _assert(summary.print_area and "$A$3" in str(summary.print_area), f"print area {summary.print_area}")
        _assert(summary.sheet_view.showGridLines is False and not summary.print_options.gridLines, "no gridlines")
        _assert(summary["A3"].value.startswith("AI Tools — Spend & Usage · ") and summary["A3"].font.sz >= 20, "headline ≥ 20pt")
        body = [c.font.sz for row in summary.iter_rows(min_row=4) for c in row
                if isinstance(c.value, str) and c.value and not c.value.startswith("=")]
        _assert(body and min(body) >= 11, f"body text ≥ 11pt (smallest {min(body)})")
        team = wb[WS.TEAM]
        _assert(team.page_setup.fitToHeight == 2 and team.page_setup.orientation == "landscape", "Team & Clients ≤ 2 pages")
        soffice = soffice_path()
        if soffice:
            # Print the Summary on its own: exactly one page.
            only = _summary_only(path)
            pdf = only[:-5] + ".pdf"
            to_pdf(only, pdf)
            with open(pdf, "rb") as fh:
                pages = len(re.findall(rb"/Type\s*/Page(?!s)", fh.read()))
            _assert(pages == 1, f"Summary prints on {pages} pages")
            print("    LibreOffice: the Summary PDF has exactly one page")
        else:
            print("    SKIP one-page PDF check: LibreOffice not installed (page setup checked instead)")
    finally:
        os.remove(path)
    print("ok  management version: no prompts, outputs or test accounts; Summary set to one A4 landscape page")


def _summary_only(path: str) -> str:
    """A workbook with just the Summary (and its chart data), built the same way."""
    with SessionLocal() as db:
        f = _filters(db)
        view = M.build_management(db, f, load_groups(db, f))
    wb = Workbook(write_only=True)
    ws, data = wb.create_sheet(WS.SUMMARY), wb.create_sheet(WS.CHART_DATA)
    data.sheet_state = "hidden"
    WS.write_summary(ws, data, view, with_details=False)
    out = path[:-5] + "_summary.xlsx"
    wb.save(out)
    return out


def main() -> int:
    seed()
    test_costing_and_dedupe()
    test_previous_period()
    test_sentences()
    test_full_version_totals_and_links()
    test_management_version_and_print()
    print("ALL MANAGEMENT REPORT CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    raise SystemExit(main())
