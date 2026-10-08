"""Credit Consumption Report: aggregation, reconciliation and link checks.

Run from backend/:  python tests/credit_report_smoke.py

Seeds every tool table in an in-memory SQLite database with known numbers,
then checks:
  * status mapping (raw status -> Charged / Pending / Failed / Refunded) in
    Python and SQL
  * the GROUP BY queries: credits, Kling clamp, owner-less rows as
    "Unassigned", IST date boundaries, every filter, the test & admin
    account exclusion
  * the model: ranks, ties, Unassigned never ranked, client-tagging rate
  * duplicate detection (same user/tool/prompt within 5 s, both charged)
  * an end-to-end workbook: Charged totals reconcile across every sheet
    (Unassigned included), Charged + Pending + Failed = every log row, no
    Output link contains "localhost", and every internal hyperlink lands on
    an existing cell whose text matches the link label
  * navigation: breadcrumbs on row 1 of every sheet; every link target has a
    way back (breadcrumb or "⬅" back cell); every incoming link lands on the
    back cell for its own source sheet; one-to-one links round-trip
  * Generation Log readability: one-line rows, 120-character prompt
    preview + full prompt column, headers wide enough
  * By User / By Tool detail blocks: every "Generations (n) →" opens the
    first log row of its (user, tool) or (user, tool, day) run, whose rows
    and credits match the block; each block's tables add up to its list row;
    every log back link returns to the row that links to it
  * edge cases: no data in range, formula-looking prompts stay text, long
    prompts are truncated, mostly-empty Task / Model columns are dropped
  * the API: Section Access gating, options, validation, output lookup
"""
import os
import re
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

BACKEND_DIR = Path(__file__).resolve().parents[1]
os.environ.setdefault("DATABASE_URL", "postgresql://placeholder:placeholder@localhost:5432/placeholder")
os.environ.setdefault("ARCHIVE_DATABASE_URL", os.environ["DATABASE_URL"])
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import column_index_from_string  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

from models_new import (  # noqa: E402
    Base,
    GenerationClient,
    GenerationRecord,
    ITPortalTool,
    ITPortalToolUsageEvent,
    User,
    UserFeatureAccess,
    UserRole,
)
from providers.elevenlabs.models import ElevenlabsGeneration  # noqa: E402
from providers.envato.models import EnvatoGeneration  # noqa: E402
from providers.epidemicsound.models import EpidemicAdaptation  # noqa: E402
from providers.flow.models import FlowGeneration  # noqa: E402
from providers.freepik.models import FreepikGeneration  # noqa: E402
from providers.heygen.models import HeygenGeneration  # noqa: E402
from providers.higgsfield.models import HiggsfieldGeneration  # noqa: E402
from providers.suno.models import SunoGeneration  # noqa: E402
from utils.credit_report import (  # noqa: E402
    CHARGED,
    FAILED,
    NO_CLIENT,
    PENDING,
    UNASSIGNED,
    UNASSIGNED_USER_ID,
    ReportFilters,
    build_model,
    charge_status,
    generate_report,
    load_excluded_summary,
    load_groups,
    resolve_excluded_accounts,
)
from utils.credit_report import navigation as N  # noqa: E402
from utils.credit_report import workbook as W  # noqa: E402
from utils.credit_report.facts import EXCLUDED_ACCOUNTS_ENV  # noqa: E402

engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
TABLES = [
    User, UserRole, UserFeatureAccess, GenerationClient, ITPortalTool, ITPortalToolUsageEvent, GenerationRecord,
    FreepikGeneration, EnvatoGeneration, HeygenGeneration, HiggsfieldGeneration, ElevenlabsGeneration,
    FlowGeneration, SunoGeneration, EpidemicAdaptation,
]
Base.metadata.create_all(bind=engine, tables=[m.__table__ for m in TABLES])

OCT1, OCT7, SEP1 = "2026-10-01", "2026-10-07", "2026-09-01"
DASHBOARD = "https://dash.example.com"
LONG_PROMPT = "x" * 40_000
INJECTION = '=HYPERLINK("http://evil.example","click")'
DUP_PROMPT = "a red fox, cinematic"
IDS: dict = {}

# Expected figures for 1-7 Oct 2026 (IST), admin accounts excluded.
# Tester is a test account but is only excluded once listed in CREDIT_REPORT_EXCLUDED_ACCOUNTS.
EXPECTED_USER_CHARGED = {"Ravi": 230, "Asha": 200, "Bob": 1035, "Zed": 10, "Del": 15, "Tester": 33, UNASSIGNED: 999}
CHARGED_TOTAL = 2522
CHARGED_ROWS = 18
PENDING_TOTAL = 2536
PENDING_ROWS = 8
FAILED_TOTAL = 40
FAILED_ROWS = 1
ALL_ROWS = CHARGED_ROWS + PENDING_ROWS + FAILED_ROWS
ADMIN_CREDITS, TESTER_CREDITS = 77, 33


def _assert(condition, message):
    if not condition:
        raise AssertionError(message)


def _close(a, b, message):
    _assert(abs(float(a) - float(b)) < 0.005, f"{message}: {a} != {b}")


def seed():
    """Rows per user (Oct, IST), C = Charged, P = Pending, F = Failed:

    Ravi  Kling 100 C Acme, 50 C, duplicate pair 10+10 C 3 s apart, 10 C 20 s later (not a dup),
          10 P same prompt (not a dup: pending); HeyGen 50 C Acme           = 230 C, 10 P
    Asha  Kling 5000->0 C, Freepik 20 C Beta, HeyGen 180 C Beta, HeyGen 6 P processing,
          HeyGen 0 P draft                                                  = 200 C, 6 P
    Bob   Freepik 30 C Acme (long prompt), Freepik 40 F, ElevenLabs 5 C (no status), ElevenLabs 1500 P
          generating_music, Suno None C complete, Suno P streaming, Epidemic 1000 C Acme, 1000 P draft,
          Kling 0 P submitted, Kling 20 P reconciling                       = 1035 C, 2520 P, 40 F
    Zed   (blank dept) Envato 10 C, Flow 0 C (30 Sep 20:00 UTC = 1 Oct IST)  = 10 C
    Del   (deleted) Higgsfield 15 C Gamma                                   = 15 C
    (no owner) Freepik 999 C, history-sync import                           = 999 C  -> "Unassigned"
    Admin (is_admin) Freepik 77 C; Tester HeyGen 33 C                        -> excluded by default
    Outside the range: Kling Ravi 70 C on 30 Sep, Flow Zed on 8 Oct IST.
    """
    with SessionLocal() as db:
        def user(name, dept, deleted=False, admin=False, email=None):
            u = User(email=email or f"{name.lower()}@example.com", name=name, hashed_password="x", department=dept,
                     is_active=True, is_deleted=deleted, is_admin=admin, employee_id=f"E-{name}")
            db.add(u)
            db.flush()
            IDS[name] = u.id
            return u.id

        ravi, asha = user("Ravi", "Design"), user("Asha", " Design ")
        bob, zed, dele = user("Bob", "Video"), user("Zed", "  "), user("Del", "Video", deleted=True)
        admin = user("Admin", "Management", admin=True)
        tester = user("Tester", "Video", email="tester@example.com")
        role_admin = user("RoleAdmin", "Management")
        db.add(UserRole(user_id=role_admin, role="admin"))
        user("Plain", "Video")
        db.add_all([GenerationClient(name="Acme"), GenerationClient(name="Beta"), GenerationClient(name="Gamma")])
        kling = ITPortalTool(name="Kling AI", slug="kling", website_url="https://kling.example")
        db.add(kling)
        db.flush()

        def ev(uid, credits, client=None, day=date(2026, 10, 3), prompt=None, status="settled", at=None):
            e = ITPortalToolUsageEvent(tool_id=kling.id, user_id=uid, event_type="generation", event_date=day,
                                       credits_burned=credits, linked_client_name=client, prompt_text=prompt,
                                       status=status, created_at=at or datetime(day.year, day.month, day.day, 6))
            db.add(e)
            db.flush()
            return e

        e1 = ev(ravi, 100, "Acme", prompt="a cat on the moon")
        ev(ravi, 50, "  ")
        dup_day = datetime(2026, 10, 4, 6, 0, 0)
        ev(ravi, 10, prompt=DUP_PROMPT, day=date(2026, 10, 4), at=dup_day)
        ev(ravi, 10, prompt=DUP_PROMPT, day=date(2026, 10, 4), at=dup_day.replace(second=3))
        ev(ravi, 10, prompt=DUP_PROMPT, day=date(2026, 10, 4), at=dup_day.replace(second=2), status="submitted")
        ev(ravi, 10, prompt=DUP_PROMPT, day=date(2026, 10, 4), at=dup_day.replace(second=20))
        ev(asha, 5000)                                          # outside 0-3000 -> 0 credits
        ev(ravi, 70, "Acme", day=date(2026, 9, 30))             # September
        ev(bob, 0, status="submitted")
        ev(bob, 20, status="reconciling")
        db.add(GenerationRecord(provider="kling", provider_task_id="k1", source_usage_event_id=e1.id,
                                canonical_asset_url="https://cdn.kling.example/k1.mp4", owner_user_id=ravi))
        IDS["kling_event"] = e1.id

        at = datetime(2026, 10, 2, 9)
        db.add_all([
            FreepikGeneration(creation_id="f1", owner_user_id=asha, credits_charged=20, linked_client_name="Beta",
                              provider_created_at=at, prompt="beta poster", status="completed"),
            FreepikGeneration(creation_id="f2", owner_user_id=bob, credits_charged=30, linked_client_name="Acme",
                              provider_created_at=at, prompt=LONG_PROMPT, raw_url="https://freepik.example/f2.png",
                              status="completed"),
            FreepikGeneration(creation_id="f3", owner_user_id=None, credits_charged=999, provider_created_at=at,
                              status="completed", ingestion_source="recovered", ownership_status="unknown"),
            FreepikGeneration(creation_id="f4", owner_user_id=bob, credits_charged=40, provider_created_at=at,
                              status="failed"),
            FreepikGeneration(creation_id="f5", owner_user_id=admin, credits_charged=ADMIN_CREDITS,
                              provider_created_at=at, status="completed"),
            HeygenGeneration(video_id="h1", owner_user_id=ravi, credits_used=50, linked_client_name="Acme",
                             provider_created_at=at, status="completed"),
            HeygenGeneration(video_id="h2", owner_user_id=asha, credits_used=180, linked_client_name="Beta",
                             provider_created_at=at, status="completed"),
            HeygenGeneration(video_id="h3", owner_user_id=asha, credits_used=6, provider_created_at=at, status="processing"),
            HeygenGeneration(video_id="h4", owner_user_id=asha, credits_used=0, provider_created_at=at, status="draft"),
            HeygenGeneration(video_id="h5", owner_user_id=tester, credits_used=TESTER_CREDITS, provider_created_at=at,
                             status="completed"),
            # TTS: blank status, the provider's state is in metadata_json.
            ElevenlabsGeneration(provider_creation_id="el1", owner_user_id=bob, credits_used=5, prompt=INJECTION,
                                 provider_created_at=at, metadata_json={"state": "created"}),
            ElevenlabsGeneration(provider_creation_id="el2", owner_user_id=bob, credits_used=1500,
                                 provider_created_at=at, status="generating_music"),
            SunoGeneration(provider_creation_id="s1", owner_user_id=bob, credits_used=None, provider_created_at=at,
                           status="complete", asset_mirror_status="mirrored", mirrored_asset_key="suno/s1.mp3",
                           media_url="https://cdn.suno.example/s1.mp3", model_name="chirp-v4"),
            SunoGeneration(provider_creation_id="s2", owner_user_id=bob, provider_created_at=at, status="streaming"),
            EpidemicAdaptation(version_id="ep1", owner_user_id=bob, credits_used=1000, status="completed",
                               linked_client_name="Acme", created_at=at),
            EpidemicAdaptation(version_id="ep2", owner_user_id=bob, credits_used=1000, status="draft", created_at=at),
            EnvatoGeneration(item_uuid="en1", owner_user_id=zed, credits_badge=10, provider_created_at=at),
            FlowGeneration(provider_creation_id="fl1", owner_user_id=zed, provider_created_at=datetime(2026, 9, 30, 20)),
            FlowGeneration(provider_creation_id="fl2", owner_user_id=zed, provider_created_at=datetime(2026, 10, 7, 19)),
            HiggsfieldGeneration(generation_id="hf1", owner_user_id=dele, credits_used=15, linked_client_name="Gamma",
                                 provider_created_at=at, status="completed"),
        ])
        db.commit()


def _filters(db, **kw):
    f = ReportFilters.build(start=kw.pop("start", OCT1), end=kw.pop("end", OCT7), **kw)
    if f.exclude_test_accounts:
        f.excluded_accounts = resolve_excluded_accounts(db)
    return f


def _charged(groups, key):
    out = {}
    for g in groups:
        if g.charge_status == CHARGED:
            out[getattr(g, key)] = out.get(getattr(g, key), 0.0) + g.credits
    return out


def _sum(groups, status, attr="credits"):
    return sum(getattr(g, attr) for g in groups if g.charge_status == status)


# --------------------------------------------------------------------------- #
# Status mapping, aggregation, filters, exclusion
# --------------------------------------------------------------------------- #
def test_status_mapping():
    charged = ["Settled", "settled", "Completed", "complete", "Captured", "success", "done", "active", "", None, "  "]
    pending = ["Submitted", "Queued", "Running", "Generating_Music", "Reconciling", "Streaming", "processing",
               "rendering", "pending", "draft", "brand-new-status"]
    failed = ["failed", "Error", "cancelled", "canceled", "rejected", "timeout", "Refunded"]
    for raw in charged:
        _assert(charge_status(raw) == CHARGED, f"{raw!r} -> Charged")
    for raw in pending:
        _assert(charge_status(raw) == PENDING, f"{raw!r} -> Pending")
    for raw in failed:
        _assert(charge_status(raw) == FAILED, f"{raw!r} -> Failed / Refunded")
    print("ok  status mapping")


def test_group_queries():
    os.environ.pop(EXCLUDED_ACCOUNTS_ENV, None)
    with SessionLocal() as db:
        groups = load_groups(db, _filters(db))
        _close(_sum(groups, CHARGED), CHARGED_TOTAL, "charged credits")
        _assert(_sum(groups, CHARGED, "generations") == CHARGED_ROWS, "charged generations")
        _close(_sum(groups, PENDING), PENDING_TOTAL, "pending credits (SQL status mapping)")
        _assert(_sum(groups, PENDING, "generations") == PENDING_ROWS, "pending generations")
        _close(_sum(groups, FAILED), FAILED_TOTAL, "failed credits")
        _assert(sum(g.generations for g in groups) == ALL_ROWS, "Charged + Pending + Failed = all rows")
        by_user = _charged(groups, "user_name")
        for name, expected in EXPECTED_USER_CHARGED.items():
            _close(by_user.get(name, 0), expected, f"{name} charged credits")
        unassigned = [g for g in groups if g.user_id == UNASSIGNED_USER_ID]
        _assert(unassigned and all(g.department == UNASSIGNED for g in unassigned), "no owner -> Unassigned/Unassigned")
        _close(_charged(groups, "tool")["Kling"], 180, "Kling: clamp 5000 -> 0, September excluded")
        flow = [g for g in groups if g.tool == "Flow"]
        _assert(sum(g.generations for g in flow) == 1, "Flow: IST boundaries keep 1 of 2 rows")
        _assert({g.month for g in groups} == {"2026-10"}, "IST month")
        _assert("Admin" not in by_user and "RoleAdmin" not in by_user, "admins excluded by default")
        _assert("Tester" in by_user, "unlisted test account still included")
        zero = sum(g.zero_credit for g in groups if g.charge_status == CHARGED)
        _assert(zero == 3, f"zero-credit charged rows (clamped Kling, Suno, Flow): {zero}")
    print("ok  group queries")


def test_filters_and_exclusion():
    with SessionLocal() as db:
        design = load_groups(db, _filters(db, department="Design"))
        _assert({g.user_name for g in design} == {"Ravi", "Asha"}, "department filter")
        un = load_groups(db, _filters(db, department=UNASSIGNED))
        _assert({g.user_name for g in un} == {"Zed", UNASSIGNED}, "Unassigned department = blank dept + no owner")
        owner_less = load_groups(db, _filters(db, user_id=UNASSIGNED_USER_ID))
        _assert({g.user_name for g in owner_less} == {UNASSIGNED}, "user filter 0 = owner-less rows")
        _assert({g.tool for g in load_groups(db, _filters(db, tool="Kling"))} == {"Kling"}, "tool filter")
        nc = load_groups(db, _filters(db, client=NO_CLIENT))
        _assert(nc and {g.client for g in nc} == {NO_CLIENT}, "No client filter")
        _close(_sum(load_groups(db, _filters(db, user_id=IDS["Bob"])), CHARGED), 1035, "user filter")
        _assert(load_groups(db, _filters(db, tool="Nope")) == [], "unknown tool returns nothing")

        # Exclusion: off -> admin + tester counted; env list adds the tester.
        off = load_groups(db, _filters(db, exclude_test_accounts=False))
        _close(_sum(off, CHARGED), CHARGED_TOTAL + ADMIN_CREDITS, "exclusion off includes the admin")
        os.environ[EXCLUDED_ACCOUNTS_ENV] = "Tester@Example.com, 999999"
        try:
            f = _filters(db)
            names = {n for _i, n in f.excluded_accounts}
            _assert(names == {"Admin", "RoleAdmin", "Tester"}, f"flag, role row and env list: {names}")
            on = load_groups(db, f)
            _close(_sum(on, CHARGED), CHARGED_TOTAL - TESTER_CREDITS, "env-listed test account excluded")
            ex = load_excluded_summary(db, f)
            _close(ex["credits"], ADMIN_CREDITS + TESTER_CREDITS, "excluded credits summary")
            _assert(ex["accounts"] == 3 and ex["generations"] == 2, f"excluded summary {ex}")
        finally:
            os.environ.pop(EXCLUDED_ACCOUNTS_ENV, None)
    print("ok  filters and exclusion")


def test_model():
    with SessionLocal() as db:
        model = build_model(load_groups(db, _filters(db)))
    ranks = {u.name: (u.company_rank, u.department_rank) for u in model.users}
    _assert(ranks["Bob"] == (1, 1) and ranks["Ravi"] == (2, 1) and ranks["Asha"] == (3, 2), f"ranks {ranks}")
    _assert(ranks[UNASSIGNED] == ("—", "—"), "Unassigned is not ranked")
    _assert(model.users[-1].name == UNASSIGNED, "Unassigned listed last")
    video = next(d for d in model.departments if d.name == "Video")
    _assert(video.top_user.keys == (IDS["Bob"],), "Video's top user")
    # One top-user rule: everyone counts, Unassigned included (credits, then
    # generations); when Unassigned wins, the best named user is kept beside it.
    unassigned_dept = next(d for d in model.departments if d.name == UNASSIGNED)
    _assert(unassigned_dept.top_user.keys == (UNASSIGNED_USER_ID,), "Unassigned (999 credits) tops its department")
    _assert(unassigned_dept.top_named_user.keys == (IDS["Zed"],), "top named user beside it")
    _assert(model.top_user.keys == (IDS["Bob"],) and model.top_named_user.keys == (IDS["Bob"],),
            "company top user: Bob (1,035) beats Unassigned (999)")
    freepik = next(t for t in model.tools if t.name == "Freepik")
    _assert(freepik.top_user.keys == (UNASSIGNED_USER_ID,) and freepik.top_named_user.keys == (IDS["Bob"],),
            "Freepik: Unassigned on top, Bob the top named user")
    bob = model.user(IDS["Bob"])
    _close(bob.client_rate, 0.5, "Bob: 2 of 4 charged generations have a client")
    _close(bob.totals.credits_per_generation, 1035 / 4, "credits per generation")
    _assert(bob.top_tool.keys == ("Epidemic Sound",) and bob.top_tool.generations == 1, "top tool with gens")
    suno = next(t for t in model.tools if t.name == "Suno")
    _assert(suno.cost_note.startswith("Fixed price per song"), "Suno: priced by an admin setting")
    _assert(model.clients[-1].name == NO_CLIENT, "No client sorts last")
    clients_in_order = [r[2] for r in model.user_client_rows]
    _assert(clients_in_order == sorted(clients_in_order, key=[c.name for c in model.clients].index),
            "User x Client sorted by client")
    print("ok  model")


def test_duplicate_detection():
    def row(sec, prompt="p", tool="Kling", credits=10.0, status=CHARGED, user=1):
        return SimpleNamespace(user_id=user, tool=tool, prompt=prompt, credits=credits, charge_status=status,
                               occurred_at=datetime(2026, 10, 1, 6, 0, sec))
    rows = [row(0), row(3), row(9), row(14, credits=0), row(15, tool="Freepik"), row(16, prompt="q"),
            row(17, status=PENDING), row(18, user=2), row(30), row(35)]
    flags = [e["dup"] for e in W.flag_duplicates(rows)]
    _assert(flags == [True, True, False, False, False, False, False, False, True, True],
            f"dup flags: {flags}")
    print("ok  duplicate detection")


def test_block_rules():
    """"Most used" ties and free tools; > 60 active dates become weeks."""
    from datetime import timedelta
    from utils.credit_report.blocks import Agg, DayRow, date_rows, pick_most_used

    tie = pick_most_used({"A": Agg(rows=5, gens=5, credits=10), "B": Agg(rows=9, gens=9, credits=10)},
                         Agg(rows=14, gens=14, credits=20))
    _assert(tie.name == "B" and "50% of credits" in tie.text(), f"credit tie broken by generations: {tie.text()}")
    free = pick_most_used({"Flow": Agg(rows=3, gens=3), "Suno": Agg(rows=7, gens=7)}, Agg(rows=10, gens=10))
    _assert(free.name == "Suno" and free.by_generations and "no credits recorded" in free.text(),
            f"free tools ranked by generations, and it says so: {free.text()}")
    client = pick_most_used({NO_CLIENT: Agg(gens=9, credits=90), "Acme": Agg(gens=1, credits=10)},
                            Agg(gens=10, credits=100), skip=NO_CLIENT)
    _assert(client.name == "Acme" and round(client.share, 2) == 0.10, "No client ignored unless it is the only one")
    only = pick_most_used({NO_CLIENT: Agg(gens=2, credits=5)}, Agg(gens=2, credits=5), skip=NO_CLIENT)
    _assert(only.name == NO_CLIENT, "No client kept when it is the only one")
    days = {date(2026, 1, 1) + timedelta(days=i): DayRow(date(2026, 1, 1) + timedelta(days=i), total=1.0)
            for i in range(61)}
    rows, weekly = date_rows(days)
    _assert(weekly and all(r.day.weekday() == 0 for r in rows) and sum(r.total for r in rows) == 61,
            "more than 60 dates are grouped by Monday-start week, totals kept")
    _assert(not date_rows(dict(list(days.items())[:60]))[1], "60 dates stay daily")
    print("ok  block rules: most used, mostly works on client, weekly above 60 dates")


def test_log_columns_dropped_when_empty():
    m = SimpleNamespace(task_fill=0.19, model_fill=0.25)
    headers = W.log_headers(m)
    _assert("Task" not in headers and "Model / type" in headers, f"80% empty rule: {headers}")
    print("ok  log column drop rule")


# --------------------------------------------------------------------------- #
# Workbook
# --------------------------------------------------------------------------- #
LINK_RE = re.compile(r'^=HYPERLINK\("(?P<target>(?:[^"]|"")*)","(?P<label>(?:[^"]|"")*)"\)$')
INTERNAL_RE = re.compile(r"^#'(?P<sheet>[^']+)'!(?P<col>[A-Z]+)(?P<row>\d+)$")


def _build(path, **kw):
    with SessionLocal() as db:
        return generate_report(db, _filters(db, **kw), path, dashboard_url=DASHBOARD)


def _display(value, fmt=None):
    """What a cell shows: a link's label, a number at 2 dp, a date in its
    number format ("mmm yyyy" -> "Mar 2026"), or the text."""
    if isinstance(value, str):
        m = LINK_RE.match(value)
        return m.group("label").replace('""', '"') if m else value
    if isinstance(value, datetime):
        if fmt == "mmm yyyy":
            return value.strftime("%b %Y")
        if fmt == "dd mmm yyyy":
            return value.strftime("%d %b %Y")
        return value.isoformat(sep=" ")
    if isinstance(value, (int, float)):
        return f"{value:,.2f}"
    return "" if value is None else str(value)


def _normalize_label(label):
    return N.strip_decoration(label)


def _table(ws, name=None):
    tables = list(ws.tables.values())
    table = next(t for t in tables if t.displayName == name) if name else tables[0]
    first, last = table.ref.split(":")
    header_row = int(re.sub(r"[A-Z]", "", first))
    last_row = int(re.sub(r"[A-Z]", "", last))
    headers = [c.value for c in ws[header_row]][: len(table.tableColumns)]
    rows = list(ws.iter_rows(min_row=header_row + 1, max_row=last_row, max_col=len(headers), values_only=True))
    return headers, rows


def _col_sum(ws, header, table=None, where=None):
    headers, rows = _table(ws, table)
    i = headers.index(header)
    if where:
        wi = headers.index(where[0])
        rows = [r for r in rows if r[wi] == where[1]]
    return sum(float(r[i] or 0) for r in rows)


def _links(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.data_type == "f" and cell.value.startswith("=HYPERLINK("):
                    m = LINK_RE.match(cell.value)
                    _assert(m, f"{ws.title}!{cell.coordinate} malformed link {cell.value[:80]}")
                    yield ws, cell, m.group("target").replace('""', '"'), m.group("label").replace('""', '"')


def _internal(value):
    """(sheet, coordinate) an internal link cell points at, else None."""
    m = LINK_RE.match(value) if isinstance(value, str) else None
    t = INTERNAL_RE.match(m.group("target").replace('""', '"')) if m else None
    return (t.group("sheet"), f"{t.group('col')}{t.group('row')}") if t else None


def _is_back_cell(cell):
    return isinstance(cell.value, str) and cell.data_type == "f" and _display(cell.value).startswith("⬅")


def _validate_links(wb):
    """Every link lands on an existing cell that matches it:
    * a back link ("⬅ ...") lands on its label's text, or on a link that
      leads back to it (a round trip; checked in detail by
      _validate_navigation and _validate_blocks);
    * a link that lands on a back cell, or whose label names an action
      ("Open block →", "Generations (48) →"), is checked by those too;
    * any other link lands on its label's text."""
    internal = external = 0
    for ws, cell, target, label in _links(wb):
        where = f"{ws.title}!{cell.coordinate}"
        if not target.startswith("#"):
            _assert("localhost" not in target and "127.0.0.1" not in target, f"{where}: local URL {target}")
            _assert(target.startswith(f"{DASHBOARD}/open-output?ref="), f"{where}: output link {target}")
            external += 1
            continue
        m = INTERNAL_RE.match(target)
        _assert(m, f"{where}: bad internal target {target}")
        sheet, col, row = m.group("sheet"), m.group("col"), int(m.group("row"))
        _assert(sheet in wb.sheetnames, f"{where} -> missing sheet {sheet}")
        dest = wb[sheet]
        _assert(row <= dest.max_row, f"{where} -> {target} past last row {dest.max_row}")
        target_cell = dest[f"{col}{row}"]
        shown = _display(target_cell.value, target_cell.number_format)
        want = _normalize_label(label)
        internal += 1
        if label.startswith("⬅"):
            if shown == want or (want.endswith("…") and shown.startswith(want[:-1])):
                continue                      # (a back label cut to fit its narrow column ends with "…")
            if label.endswith(N.LIST_SUFFIX) and (sheet, f"{col}{row}") == (sheet, f"A{N.list_row(sheet)}"):
                continue                      # "(list)": back to the sheet's list header row
            back = _internal(target_cell.value)
            _assert(back is not None and back[0] == ws.title,
                    f"{where} '{label}' -> {target} shows '{shown}' and does not link back")
        elif N.is_action_label(label) or _is_back_cell(target_cell):
            continue
        else:
            _assert(shown == want, f"{where} '{label}' -> {target} shows '{shown}'")
    return internal, external


def _validate_navigation(wb, expect_links=True):
    """A (breadcrumbs), B/C (back cells), and the round trip."""
    for name in W.SHEET_ORDER[1:]:
        ws = wb[name]
        crumbs = [_display(c.value) for c in ws[1] if c.value is not None]
        _assert(crumbs == W.crumb_labels(name), f"{name} breadcrumbs {crumbs}")
        parent = W.SHEET_PARENT[name]
        _assert(_internal(ws["A1"].value) == (parent, f"A{N.landing_row(parent)}"), f"{name} A1 goes to {parent}")
        _assert(ws.freeze_panes and int(re.sub(r"[A-Z]", "", ws.freeze_panes)) >= 2, f"{name} row 1 frozen")
    forward = round_trips = shared = 0
    for ws, cell, target, label in _links(wb):
        if not target.startswith("#") or label.startswith("⬅"):
            continue
        sheet, coord = _internal(cell.value)
        if sheet == W.HOME or cell.row == 1:
            continue                                      # Home is the root; row 1 is the breadcrumb itself
        dest = wb[sheet]
        landing = dest[coord]
        where = f"{ws.title}!{cell.coordinate} -> {sheet}!{coord}"
        forward += 1
        if not _is_back_cell(landing):
            # Plain sheet navigation may skip a back cell when it lands on the
            # title (the breadcrumb is the way back) or on a row that links
            # straight back to the source row (a block's "⬅ Back to ... list").
            if landing.row <= N.TITLE_ROW and landing.column == 1 and _is_back_cell(dest["A1"]):
                continue
            returns = [_internal(c.value) for c in dest[landing.row] if _is_back_cell(c)]
            if any(r and r[0] == ws.title and int(re.sub(r"[A-Z]", "", r[1])) == cell.row for r in returns):
                continue
            # The row's back slots are full: it lands on the item's name (column
            # A) and returns through the title area's "⬅ <sheet> (list)" cell.
            title_backs = [_internal(c.value) for r in (1, N.TITLE_ROW) for c in dest[r] if _is_back_cell(c)]
            _assert(landing.column == 1 and any(b and b[0] == ws.title for b in title_backs),
                    f"{where}: dead end (no back cell, no link back to the row, no list return)")
            continue
        back_sheet, back_coord = _internal(landing.value)
        _assert(back_sheet == ws.title, f"{where}: lands on the back cell for {back_sheet}, not {ws.title}")
        if _display(landing.value).endswith(N.LIST_SUFFIX):
            shared += 1
            _assert(back_coord == f"A{N.list_row(ws.title)}", f"{where}: list back link to {back_coord}")
        elif sheet == W.LOG:
            # The log's back cells return to the (user, tool) / (user, tool,
            # day) table row that links in; _validate_blocks checks them.
            round_trips += 1
            _assert(int(re.sub(r"[A-Z]", "", back_coord)) == cell.row, f"{where}: log back link returns to {back_coord}")
        else:
            round_trips += 1
            _assert(int(re.sub(r"[A-Z]", "", back_coord)) == cell.row, f"{where}: back link returns to {back_coord}")
    _assert(not expect_links or (forward > 80 and round_trips > 40 and shared > 5),
            f"navigation checked: {forward} links, {round_trips} round trips, {shared} list returns")
    return forward, round_trips, shared


def _log_rows(log):
    """[(row number, values)] of the log's generation rows (not user header / separator rows)."""
    headers = [c.value for c in log[W.HEADER_ROW]]
    status = headers.index("Charge status")
    return headers, [(i, r) for i, r in enumerate(log.iter_rows(min_row=W.FIRST_DATA_ROW, values_only=True),
                                                  start=W.FIRST_DATA_ROW) if r[status] is not None]


def _validate_log_layout(wb):
    log = wb[W.LOG]
    headers = [c.value for c in log[W.HEADER_ROW]]
    _assert(headers[-1] == W.PROMPT_FULL, f"full prompt is the last column: {headers[-1]}")
    for i, h in enumerate(headers, start=1):
        width = log.column_dimensions[W.get_column_letter(i)].width
        if h.startswith("⬅"):
            _assert(width == N.BACK_COL_WIDTH, f"log back column '{h}' is narrow ({width})")
        else:
            _assert(width >= len(h) + 4, f"log header '{h}' fits ({width})")
    preview_i = headers.index("Prompt") + 1
    when = headers.index("Date / time (IST)")
    for row in log.iter_rows(min_row=W.FIRST_DATA_ROW, max_row=log.max_row):
        if not isinstance(row[when].value, datetime):
            continue                                       # header / separator / total rows
        for c in row:
            _assert(not (c.alignment and c.alignment.wrap_text), f"{W.LOG}!{c.coordinate} wraps")
        preview = row[preview_i - 1].value or ""
        _assert(len(preview) <= W.PROMPT_PREVIEW_CHARS + 1 and "\n" not in preview, f"preview {preview[:30]!r}")
        if len(preview) > W.PROMPT_PREVIEW_CHARS:
            _assert(preview.endswith("…"), "a cut preview ends with …")
    for r, dim in log.row_dimensions.items():
        _assert(not dim.ht or dim.ht <= 15, f"log row {r} is {dim.ht}pt")


GENS_RE = re.compile(r"^(?P<all>All )?[Gg]enerations \((?P<n>[\d,]+)(?: · (?P<p>[\d,]+) pending)?"
                     r"(?: · (?P<f>[\d,]+) failed)?\)$")


def _validate_blocks(wb):
    """Detail blocks, their links into the log and the log's links back.

    * every "Generations (n …) →" lands on the first log row of its run - the
      user header (All generations), the (user, tool) separator, or the first
      row of a (user, tool, day, client) - and the run has exactly n rows,
      with the pending / failed counts the label names; where the block shows
      credits for that exact run, they equal the run's Charged credits;
    * each user block: Tools used = Clients = Date-wise = the list row;
      each tool block: Users = Clients = Date-wise = Date › User › Client = the list row;
    * every log back link (⬅ User, ⬅ Tool, ⬅ Date, ⬅ Tool date, ⬅ Client)
      returns to the row that links into its run (⬅ Date / Tool date /
      Client: into exactly that row);
    * in-block links (Date-wise / Clients -> Tool › Date › Client, Date-wise
      -> Date › User › Client) land on a "⬅" cell that returns to their row;
    * every block has a link in (the list) and out (⬅ Back to ... list).
    """
    log = wb[W.LOG]
    lh = [c.value for c in log[W.HEADER_ROW]]
    li = {h: i for i, h in enumerate(lh)}
    rows = {i: r for i, r in enumerate(log.iter_rows(min_row=W.FIRST_DATA_ROW, values_only=True),
                                       start=W.FIRST_DATA_ROW)}

    def is_data(r):
        return r[li["Charge status"]] is not None

    def key(r, kind):
        k = (r[li["User"]], r[li["Tool"]])
        return k + (r[li["Date"]], r[li["Client"]]) if kind == "run" else (k[0],) if kind == "user" else k

    def run(start, kind):
        i = start if kind == "run" else start + 1
        first = None
        out = []
        while i in rows:
            r = rows[i]
            if not is_data(r):
                if kind == "user" and r[li["Tool"]] is not None and r[li["User"]] == rows[start][li["User"]]:
                    i += 1                                         # the user's next tool separator
                    continue
                break
            first = first or key(r, kind)
            if key(r, kind) != first:
                break
            out.append(r)
            i += 1
        return out

    def counts(rs):
        status = [r[li["Charge status"]] for r in rs]
        return len(rs), status.count(PENDING), status.count(FAILED), \
            sum(float(r[li["Credits"]] or 0) for r in rs if r[li["Charged"]] is True)

    runs = 0
    credit_cols = {(W.USER, "A"): "C", (W.USER, "C"): "E", (W.TOOL, "B"): "D", (W.TOOL, "D"): "E"}
    for sheet_name in (W.USER, W.TOOL, W.USER_CLIENT):
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=W.FIRST_DATA_ROW):
            for cell in row:
                if not (isinstance(cell.value, str) and cell.value.startswith("=HYPERLINK(")):
                    continue
                label = _display(cell.value)
                m = GENS_RE.match(N.strip_decoration(label))
                target = _internal(cell.value)
                if not m or not target or target[0] != W.LOG:
                    continue
                where = f"{ws.title}!{cell.coordinate} '{label}'"
                r0 = int(re.sub(r"[A-Z]", "", target[1]))
                col = re.sub(r"\d", "", target[1])
                if m.group("all"):
                    _assert(not is_data(rows[r0]) and rows[r0][li["Tool"]] is None and col == "A", f"{where}: user header")
                    rs = run(r0, "user")
                elif col in ("A", "B"):
                    _assert(not is_data(rows[r0]) and rows[r0][li["Tool"]], f"{where}: tool separator")
                    rs = run(r0, "pair")
                else:
                    prev = rows.get(r0 - 1)
                    _assert(is_data(rows[r0]) and (prev is None or not is_data(prev) or key(prev, "run") != key(rows[r0], "run")),
                            f"{where}: first row of its (user, tool, day, client) run")
                    rs = run(r0, "run")
                n, pending, failed, charged = counts(rs)
                _assert(n == int(m.group("n").replace(",", "")) and pending == int((m.group("p") or "0").replace(",", ""))
                        and failed == int((m.group("f") or "0").replace(",", "")),
                        f"{where}: run has {n} rows, {pending} pending, {failed} failed")
                back = _internal(log[target[1]].value)
                _assert(back and back[0] == ws.title and int(re.sub(r"[A-Z]", "", back[1])) == cell.row,
                        f"{where}: the log's back link returns to {back}")
                shown = credit_cols.get((sheet_name, col))
                if shown and not m.group("all"):
                    _close(charged, float(ws[f"{shown}{cell.row}"].value or 0), f"{where}: credits = log run")
                runs += 1

    for i, r in rows.items():
        for name in W.LOG_BACK_COLUMNS:
            cell = log.cell(i, li[name] + 1)
            back = _internal(cell.value)
            if not back:
                continue
            fwd = _internal(wb[back[0]][back[1]].value)
            _assert(fwd and fwd[0] == W.LOG, f"{W.LOG}!{cell.coordinate} -> {back}: not a link into the log")
            start = int(re.sub(r"[A-Z]", "", fwd[1]))
            _assert(start <= i and rows[start][li["User"]] == r[li["User"]],
                    f"{W.LOG}!{cell.coordinate}: back link returns to a row for another run")
            if name in (W.LOG_BACK_DATE, W.LOG_BACK_TOOL_DATE, W.LOG_BACK_CLIENT):
                _assert(start == i, f"{W.LOG}!{cell.coordinate}: {name} returns to the row that opens this run")

    # In-block links: Date-wise / Clients rows open a "⬅" cell that returns to them.
    inner = 0
    for sheet_name in (W.USER, W.TOOL):
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=W.FIRST_DATA_ROW):
            for cell in row:
                target = _internal(cell.value) if isinstance(cell.value, str) else None
                if not target or target[0] != sheet_name or _display(cell.value).startswith("⬅"):
                    continue
                landing = ws[target[1]]
                if not _is_back_cell(landing) or landing.column_letter not in ("G", "H"):
                    continue
                back = _internal(landing.value)
                _assert(back[0] == sheet_name and int(re.sub(r"[A-Z]", "", back[1])) == cell.row,
                        f"{sheet_name}!{cell.coordinate}: in-block link returns to {back}")
                inner += 1

    def tables_in(ws, first, last):
        """{table title: (headers, [rows])} between two rows."""
        out, r = {}, first
        while r <= last:
            a = ws.cell(r, 1).value
            hdr = [c.value for c in ws[r + 1]] if r + 1 <= last else []
            if isinstance(a, str) and hdr and hdr[0] in ("Tool", "Client", "Date", "Week", "User") and \
                    ws.cell(r + 1, 1).fill.fgColor.rgb in ("001F3864", "FF1F3864"):
                body, k = [], r + 2
                while k <= last and ws.cell(k, 1).value not in (None, ""):
                    body.append([c.value for c in ws[k]])
                    k += 1
                out[a] = (hdr, body)
                r = k
            r += 1
        return out

    def credits_of(table, header):
        headers, body = table
        i = headers.index(header)
        return sum(float(b[i] or 0) for b in body if i < len(b))

    blocks_seen = 0
    for sheet_name, sections in ((W.USER, {"Tools used": "Credits", "Clients": "Credits", "Date-wise": "Total",
                                           "Tool › Date › Client": "Credits"}),
                                 (W.TOOL, {"Users": "Credits", "Clients": "Credits", "Date-wise": "Credits",
                                           "Date › User › Client": "Credits"})):
        ws = wb[sheet_name]
        lh2, lrows = _table(ws)
        list_credit = {_display(r[0]): float(r[lh2.index("Credits")] or 0) for r in lrows}
        incoming = {}
        for s2, c2, _t2, _l in _links(wb):
            tgt = _internal(c2.value)
            if tgt and tgt[0] == sheet_name:
                incoming.setdefault(int(re.sub(r"[A-Z]", "", tgt[1])), []).append(f"{s2.title}!{c2.coordinate}")
        heads = [r for r in range(W.FIRST_DATA_ROW + len(lrows), ws.max_row + 1)
                 if ws.cell(r, 1).fill.fgColor.rgb in ("00BDD7EE", "FFBDD7EE") and ws.cell(r, 1).value]
        _assert(len(heads) == len(lrows), f"{sheet_name}: one block per list row ({len(heads)} vs {len(lrows)})")
        for k, h in enumerate(heads):
            name = ws.cell(h, 1).value
            end = (heads[k + 1] - 1) if k + 1 < len(heads) else ws.max_row
            tables = tables_in(ws, h + 1, end)
            got = {}
            for title, table in tables.items():
                for sec, header in sections.items():
                    if title.startswith(sec):
                        got[sec] = credits_of(table, header)
            _assert(set(got) == set(sections), f"{sheet_name} {name}: sections {sorted(tables)}")
            for sec, value in got.items():
                _close(value, list_credit[name], f"{sheet_name} {name}: {sec} = list row")
            _assert(h in incoming, f"{sheet_name} {name}: no link into the block")
            _assert(any(_display(c.value).startswith("⬅ Back to") for c in ws[h]), f"{sheet_name} {name}: no way back")
            _assert(ws.row_dimensions[h + 1].outlineLevel == 1, f"{sheet_name} {name}: block is collapsible")
            blocks_seen += 1
    _assert(runs > 0 and blocks_seen > 0 and inner > 0, "blocks checked")
    print(f"ok  blocks: {blocks_seen} user/tool blocks, {runs} log runs, {inner} in-block links; "
          "every back link round-trips")


def _freeze_col(ws) -> int:
    """How many columns the sheet freezes (0 = none)."""
    if not ws.freeze_panes:
        return 0
    letters = re.sub(r"\d", "", ws.freeze_panes)
    return max(0, column_index_from_string(letters) - 1)


def _freeze_row(ws) -> int:
    return int(re.sub(r"[A-Z]", "", ws.freeze_panes)) - 1 if ws.freeze_panes else 0


def WP_TABLES_HEADER(model):
    from utils.credit_report.workbook_periods import plan_explorer
    return plan_explorer(model).tables_header


def _audit_landings(wb):
    """No link lands on an empty cell; a cross-sheet landing keeps its row's
    name in view (column A frozen, or the cell within 120 characters of
    width of column A)."""
    checked = 0
    for ws, cell, target, label in _links(wb):
        t = _internal(cell.value)
        if not t:
            continue
        dest = wb[t[0]]
        landing = dest[t[1]]
        _assert(landing.value not in (None, ""), f"{ws.title}!{cell.coordinate} '{label}' lands on empty {t}")
        if t[0] == ws.title:
            continue
        offset = sum(dest.column_dimensions[W.get_column_letter(c)].width or 13
                     for c in range(1, landing.column))
        _assert(_freeze_col(dest) >= 1 or offset <= 120,
                f"{ws.title}!{cell.coordinate} -> {t}: lands {offset:.0f} wide from column A, which is not frozen")
        checked += 1
    return checked


def _audit_graph(wb):
    """0 self-links; 0 back↔back pairs; 0 forward↔forward loops (two cells
    not labelled "⬅" that link to each other's rows); at most 3 "⬅ Back"
    columns per sheet."""
    edges = []
    for ws, cell, _target, label in _links(wb):
        t = _internal(cell.value)
        if not t:
            continue
        _assert((ws.title, cell.coordinate) != t, f"{ws.title}!{cell.coordinate} links to itself")
        edges.append(((ws.title, cell.row), (t[0], int(re.sub(r"[A-Z]", "", t[1]))), label.startswith("⬅"),
                      f"{ws.title}!{cell.coordinate}"))
    by_kind = {}
    for src, dst, back, where in edges:
        by_kind.setdefault((back, src, dst), []).append(where)
    loops, back_pairs = [], []
    for (back, src, dst), where in by_kind.items():
        if src == dst or (back, dst, src) not in by_kind or src > dst:
            continue
        (back_pairs if back else loops).append((where[0], by_kind[(back, dst, src)][0]))
    _assert(not loops, f"forward↔forward loops: {loops[:10]}")
    _assert(not back_pairs, f"back↔back pairs: {back_pairs[:10]}")
    for ws in wb.worksheets:
        slots = {c.column for row in ws.iter_rows(max_row=min(ws.max_row, 400)) for c in row
                 if isinstance(c.value, str) and c.value.startswith("⬅ Back ") and c.value[7:].isdigit()}
        _assert(len(slots) <= 3, f"{ws.title}: {len(slots)} back columns")
    return len(edges)


def _audit_labels(model, wb):
    """Item links land on the item; back labels are never cut; one label
    means one target within a row."""
    items = ({u.label for u in model.users} | {t.name for t in model.tools} | {c.name for c in model.clients}
             | {d.name for d in model.departments})
    named = 0
    by_row = {}
    for ws, cell, target, label in _links(wb):
        _assert("…" not in label, f"{ws.title}!{cell.coordinate}: cut label '{label}'")
        key = (ws.title, cell.row, label)
        by_row.setdefault(key, set()).add(target)
        t = _internal(cell.value)
        name = N.strip_decoration(label)
        if not t or label.startswith(("⬅", "↔", "↓")) or name not in items or cell.row == 1:
            continue
        dest = wb[t[0]]
        row = int(re.sub(r"[A-Z]", "", t[1]))
        _assert(_display(dest.cell(row, 1).value, dest.cell(row, 1).number_format) == name,
                f"{ws.title}!{cell.coordinate} '{label}' lands on {t}, whose column A is "
                f"'{_display(dest.cell(row, 1).value)}'")
        named += 1
    # Two different items may share a name (the user "Unassigned" and the
    # department "Unassigned"): allowed when each link opens its own kind of
    # item - a user's block on By User, a department on By Department.
    item_sheet = {W.USER, W.DEPT, W.TOOL, W.CLIENT}

    def kinds(targets):
        sheets = [INTERNAL_RE.match(t).group("sheet") if INTERNAL_RE.match(t) else t for t in targets]
        return len(set(sheets)) == len(sheets) and set(sheets) <= item_sheet

    dupes = {k: v for k, v in by_row.items()
             if len(v) > 1 and not (N.strip_decoration(k[2]) in items and kinds(v))}
    _assert(not dupes, f"same label, different targets in one row: {list(dupes.items())[:5]}")
    return named


def _reachability(wb, view=25, max_clicks=4):
    """Follow links only, from Home: each landing shows its row + the next
    `view` rows, frozen rows always visible. Every non-empty row of every
    summary sheet (all but the Generation Log) must come into view."""
    sheets = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    links_by_row = {}
    for ws, cell, _target, _label in _links(wb):
        t = _internal(cell.value)
        if t:
            links_by_row.setdefault((ws.title, cell.row), []).append((t[0], int(re.sub(r"[A-Z]", "", t[1]))))
    # Breadth first: a landing at depth d shows its rows; a link among them is click d + 1.
    seen, frontier, landed = {}, [(W.HOME, 1)], set()
    for depth in range(max_clicks + 1):
        nxt = []
        for sheet, row in frontier:
            if (sheet, row) in landed:
                continue
            landed.add((sheet, row))
            ws = wb[sheet]
            visible = set(range(1, _freeze_row(ws) + 1)) | set(range(row, row + view + 1))
            seen.setdefault(sheet, set()).update(visible)
            for r in visible:
                nxt.extend(links_by_row.get((sheet, r), []))
        frontier = nxt
    missing = {}
    for name in sheets:
        if name == W.LOG:
            continue
        ws = wb[name]
        rows = {c.row for row in ws.iter_rows() for c in row if c.value not in (None, "")}
        gap = sorted(rows - seen.get(name, set()))
        if gap:
            missing[name] = gap
    return missing


def _validate_top_users(wb):
    """One rule: everyone counts (Unassigned too) by Charged credits, then
    generations; "Top named user" is filled exactly when Unassigned is on top."""
    checked = 0
    for name in (W.DEPT, W.TOOL, W.CLIENT):
        headers, rows = _table(wb[name])
        tu, nu = headers.index("Top user"), headers.index("Top named user")
        for r in rows:
            top, named = _display(r[tu]), _display(r[nu])
            # Filled only when Unassigned is on top ("—" there means nobody named used it).
            _assert(named == "—" or (top == UNASSIGNED and named != UNASSIGNED),
                    f"{name} {_display(r[0])}: top {top}, named {named}")
            checked += 1
    # By Tool: the top user is the first active row of the tool's Users table.
    ws = wb[W.TOOL]
    headers, rows = _table(ws)
    for r in rows:
        tool = _display(r[0])
        block = ws[_internal(r[0])[1]].row
        k = block
        while ws.cell(k, 1).value != "User":
            k += 1
        first = next(_display(ws.cell(j, 1).value) for j in range(k + 1, k + 200)
                     if (ws.cell(j, 3).value or 0) > 0 or (ws.cell(j, 4).value or 0) > 0)
        _assert(_display(r[headers.index("Top user")]) == first, f"{tool}: top user {first}")
    return checked


def test_workbook_end_to_end():
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        model = _build(path)
        _assert(model.log_stats.plan_mismatches == 0, "every log row landed where the blocks planned it")
        wb = load_workbook(path)
        _assert(wb.sheetnames == list(W.SHEET_ORDER) + [W.LISTS], f"sheet order {wb.sheetnames}")
        _assert(wb[W.LISTS].sheet_state == "hidden", "Lists sheet is hidden")
        _assert_no_banned_functions(wb)
        log_headers = [c.value for c in wb[W.LOG][W.HEADER_ROW]]
        _assert({"Date", "Month", "Week", "Quarter", "Charged"} <= set(log_headers), f"log date columns {log_headers}")
        log_ws = wb[W.LOG]
        _lh, data_rows = _log_rows(log_ws)
        first_data = data_rows[0][0]
        month_col = log_headers.index("Month") + 1
        first_month = log_ws.cell(first_data, month_col)
        _assert(isinstance(first_month.value, datetime) and first_month.value.day == 1
                and first_month.number_format == "mmm yyyy", "Month is a real date shown mmm yyyy")
        _assert(log_ws.cell(first_data, log_headers.index("Date") + 1).number_format == "dd mmm yyyy (ddd)",
                "dates shown dd mmm yyyy (ddd)")
        week = log_ws.cell(first_data, log_headers.index("Week") + 1).value
        _assert(isinstance(week, datetime) and week.weekday() == 0, "Week starts on Monday")
        charged_vals = {r[log_headers.index("Charged")] for _i, r in data_rows}
        _assert(charged_vals == {True, False}, f"Charged is TRUE/FALSE: {charged_vals}")
        _assert(log_ws.freeze_panes == "F5", "log back columns frozen with the header")
        for ws in wb.worksheets:
            if ws.title == W.LISTS:          # hidden dropdown data, not a page
                continue
            _assert(ws.freeze_panes, f"{ws.title} has frozen panes")
        _validate_log_layout(wb)
        for name in (W.DEPT, W.TOOL, W.CLIENT, W.DEPT_TOOL, W.USER_TOOL, W.USER_CLIENT, W.QUALITY, W.LOG):
            _assert(len(wb[name].tables) == 1, f"{name} has one Excel Table")
        user_tables = list(wb[W.USER].tables)
        _assert(user_tables[0] == "tblUsers" and len(user_tables) > 1
                and all(t.startswith("tblUserTDC") for t in user_tables[1:]),
                f"By User: the list, then one filterable Tool › Date › Client table per block: {user_tables}")
        _assert(len(wb[W.TREND].tables) == 9, f"Monthly Trend tables: {list(wb[W.TREND].tables)}")
        _assert(wb[W.DEPT]["B5"].number_format == W.CREDITS_FMT, "credits use thousands separators")

        # Charged totals reconcile everywhere, Unassigned included.
        home = wb[W.HOME]
        log = wb[W.LOG]
        totals = {
            "Home KPI": home["A11"].value,
            "By Department": _col_sum(wb[W.DEPT], "Credits"),
            "By User": _col_sum(wb[W.USER], "Credits"),
            "By Tool": _col_sum(wb[W.TOOL], "Credits"),
            "By Client": _col_sum(wb[W.CLIENT], "Credits"),
            "Dept × Tool": _col_sum(wb[W.DEPT_TOOL], "Total"),
            "User × Tool": _col_sum(wb[W.USER_TOOL], "Total"),
            "User × Client": _col_sum(wb[W.USER_CLIENT], "Credits"),
            "Trend by tool": _col_sum(wb[W.TREND], "Total", "tblTrendToolCredits"),
            "Trend by department": _col_sum(wb[W.TREND], "Total", "tblTrendDeptCredits"),
            "Generation Log (Charged)": _col_sum(log, "Credits", where=("Charge status", CHARGED)),
        }
        for name, value in totals.items():
            _close(value, CHARGED_TOTAL, f"{name} charged total")
        _close(home["B11"].value, PENDING_TOTAL, "Home pending KPI")
        _close(_col_sum(wb[W.DEPT], "Pending credits"), PENDING_TOTAL, "By Department pending")
        _close(_col_sum(log, "Credits", where=("Charge status", PENDING)), PENDING_TOTAL, "log pending")
        _assert(home["C11"].value == CHARGED_ROWS, "Home generations KPI counts charged rows")
        _assert(_col_sum(wb[W.TREND], "Total", "tblTrendToolGenerations") == CHARGED_ROWS, "trend generations")

        headers, rows = _table(log)
        status_i = headers.index("Charge status")
        rows = [r for r in rows if r[status_i] is not None]      # generations, not header / separator rows
        counts = {s: sum(1 for r in rows if r[status_i] == s) for s in (CHARGED, PENDING, FAILED)}
        _assert(counts == {CHARGED: CHARGED_ROWS, PENDING: PENDING_ROWS, FAILED: FAILED_ROWS}, f"log statuses {counts}")
        _assert(sum(counts.values()) == len(rows) == ALL_ROWS, "Charged + Pending + Failed = every log row")
        raw_i = headers.index("Raw status")
        _assert({"settled", "submitted", "reconciling", "generating_music", "draft", "streaming", "created",
                 W.NO_RAW_STATUS} <= {r[raw_i] for r in rows}, "raw status kept in the log")
        el1 = next(r for r in rows if r[headers.index(W.PROMPT_FULL)] == INJECTION)
        _assert(el1[raw_i] == "created" and el1[status_i] == CHARGED,
                "ElevenLabs TTS raw status read from metadata; still Charged")
        unassigned_rows = [r for r in rows if r[headers.index("User")] == UNASSIGNED]
        _assert(len(unassigned_rows) == 1 and unassigned_rows[0][headers.index("Department")] == UNASSIGNED,
                "Unassigned row in the log")

        # Duplicates, dropped columns, injection, truncation, IST.
        dup_i = headers.index(W.DUPLICATE_LABEL)
        _assert(sum(1 for r in rows if r[dup_i] == W.DUPLICATE_LABEL) == 2, "exactly one duplicate pair flagged")
        _assert("Task" not in headers and "Model / type" not in headers, f"empty Task/Model dropped: {headers}")
        for header in ("Prompt", W.PROMPT_FULL):
            prompt_col = W.get_column_letter(headers.index(header) + 1)
            prompts = [log[f"{prompt_col}{r}"] for r, _v in data_rows]
            injected = [c for c in prompts if c.value == INJECTION]
            _assert(injected and injected[0].data_type == "s", f"formula-looking prompt stays text in {header}")
            longest = max((c.value or "" for c in prompts), key=len)
            if header == "Prompt":
                _assert(longest == "x" * W.PROMPT_PREVIEW_CHARS + "…", "long prompt previewed")
            else:
                _assert(len(longest) <= 32_767 and longest.endswith(W.TRUNCATION_MARKER), "long prompt truncated")
        when = headers.index("Date / time (IST)")
        _assert(any(isinstance(r[when], datetime) and (r[when].hour, r[when].minute) == (1, 30) for r in rows),
                "IST times")

        # Links: none local, every internal one lands on matching text, and
        # every jump has a way back.
        internal, external = _validate_links(wb)
        _assert(internal > 80 and external == 3, f"links checked: {internal} internal, {external} output")
        forward, round_trips, shared = _validate_navigation(wb)
        print(f"    navigation: {forward} links, {round_trips} round trips, {shared} list returns")
        landings = _audit_landings(wb)
        missing = _reachability(wb)
        _assert(not missing, f"rows not reachable within 4 clicks: { {k: v[:12] for k, v in missing.items()} }")
        edges = _audit_graph(wb)
        named = _audit_labels(model, wb)
        print(f"    labels: {named} item links land on their item's row; no cut labels; one label, one target per row")
        print(f"    reachability: every summary row within 4 clicks of Home; {edges} links: no self-links, "
              "no back↔back pairs, no forward↔forward loops, ≤ 3 back columns per sheet")
        tops = _validate_top_users(wb)
        print(f"    audit: {landings} cross-sheet landings keep column A in view, no link lands on an empty cell, "
              f"{tops} top users follow one rule")

        _validate_blocks(wb)

        # By Department: a department name opens its block, whose People are
        # exactly its users and add up to its credits; Tools add up too.
        dh, drows = _table(wb[W.DEPT])
        uh, urows = _table(wb[W.USER])
        dept_of = {_display(r[0]): _display(r[uh.index("Department")]) for r in urows}
        dws = wb[W.DEPT]
        for (cell,) in dws.iter_rows(min_row=W.FIRST_DATA_ROW, max_row=W.FIRST_DATA_ROW + len(drows) - 1, max_col=1):
            dept = _display(cell.value)
            sheet, coord = _internal(cell.value)
            _assert(sheet == W.DEPT, f"{dept} opens its block on By Department, not {sheet}")
            block = dws[coord].row
            _assert(_display(dws.cell(block, 1).value) == dept, f"{dept} lands on its own block")
            sections = {}
            r = block + 1
            while r <= dws.max_row and not (dws.cell(r, 1).fill.fgColor.rgb in ("00BDD7EE", "FFBDD7EE")):
                title = dws.cell(r, 1).value
                if title in ("People (most credits first)", "Tools"):
                    rows, k = [], r + 2
                    while dws.cell(k, 1).value not in (None, ""):
                        rows.append(k)
                        k += 1
                    sections[title] = rows
                    r = k
                r += 1
            people = [_display(dws.cell(k, 1).value) for k in sections["People (most credits first)"]]
            _assert(sorted(people) == sorted(u for u, d in dept_of.items() if d == dept), f"{dept} people {people}")
            dept_credits = next(float(row[1]) for row in drows if _display(row[0]) == dept)
            _close(sum(float(dws.cell(k, 3).value or 0) for k in sections["People (most credits first)"]), dept_credits,
                   f"{dept}: people add up to the department")
            _close(sum(float(dws.cell(k, 2).value or 0) for k in sections["Tools"]), dept_credits,
                   f"{dept}: tools add up to the department")

        # Data Quality and Home.
        dq_headers, dq_rows = _table(wb[W.QUALITY])
        dq = {r[0]: r for r in dq_rows}
        _assert(dq[UNASSIGNED][1] == 1 and abs(dq[UNASSIGNED][2] - 999) < 0.01, "DQ Unassigned")
        _assert(dq[PENDING][1] == PENDING_ROWS and abs(dq[PENDING][2] - PENDING_TOTAL) < 0.01, "DQ Pending")
        _assert(dq[W.DUPLICATE_LABEL][1] == 2 and abs(dq[W.DUPLICATE_LABEL][2] - 20) < 0.01, "DQ duplicates")
        _assert(dq[W.ZERO_MISSED][1] == 1, "DQ: Kling's clamped row is charged with 0 credits")
        _assert(dq[W.NOT_CAPTURED][1] == 2 and "Suno 1" in dq[W.NOT_CAPTURED][3] and "Flow 1" in dq[W.NOT_CAPTURED][3],
                "DQ: Suno and Flow rows are 'cost not captured'")
        if datetime.now() > datetime(2026, 10, 6):
            _assert(dq[W.STALE_PENDING][1] == PENDING_ROWS, "DQ: every seeded pending row is over 24 h old")
        _assert(dq["Test & admin accounts excluded"][1] == 1, "DQ excluded (admin's one generation)")
        dq_link = _internal(wb[W.QUALITY][f"{W.QUALITY_LINK_COL}{W.FIRST_DATA_ROW + W.QUALITY_ISSUES.index(W.EXCLUDED)}"].value)
        _assert(dq_link == (W.HOME, "A5"), f"Data Quality 'Test & admin accounts' -> {dq_link}")
        home_back = _internal(wb[W.HOME]["C5"].value)
        _assert(_display(wb[W.HOME]["C5"].value).startswith("⬅") and home_back and home_back[0] == W.QUALITY
                and int(re.sub(r"[A-Z]", "", home_back[1])) == W.FIRST_DATA_ROW + W.QUALITY_ISSUES.index(W.EXCLUDED),
                f"Home row 5 has a back link to its Data Quality row: {home_back}")
        qa_headers, qa_rows = _table(home, "tblQuickAnswers")
        _assert(len(qa_rows) <= 11 and not any("top user on" in (r[0] or "").lower() for r in qa_rows),
                f"{len(qa_rows)} quick answers, no per-tool rows")
        home_text = " ".join(str(c.value) for row in home.iter_rows() for c in row if c.value)
        _assert("Excluded: Admin, RoleAdmin" in home_text, "exclusion setting printed on Home")
        _assert(f"Output links open on {DASHBOARD}" in home_text.replace("  ", " ") or DASHBOARD in home_text,
                "dashboard URL printed")
        _assert(len(home._charts) == 3, f"department, tool and trend charts: {len(home._charts)}")
        tool_headers, tool_rows = _table(wb[W.TOOL])
        cost = {r[0]: r[tool_headers.index("Cost")] for r in tool_rows}
        tool_names = {_display(k): v for k, v in cost.items()}
        _assert(tool_names["Suno"].startswith("Fixed price per song") and tool_names["Flow"].startswith("Not captured"),
                "Suno priced by an admin setting; Flow marked not captured")
        _assert(tool_names["Kling"] == "Recorded", "recorded tools say so")
        user_headers, _user_rows = _table(wb[W.USER])
        _assert("Credits per generation" in tool_headers and "Credits per generation" in user_headers,
                "credits per generation on By Tool and By User")
        _assert("% generations with a client" in user_headers, "client-tagging rate on By User")
    finally:
        os.remove(path)
    print("ok  workbook end to end")


# Excel 365 / 2019-only functions: the file must work in Excel 2016, LibreOffice and Google Sheets.
BANNED = re.compile(r"\b(FILTER|UNIQUE|SORT|SORTBY|LET|LAMBDA|XLOOKUP|XMATCH|MAXIFS|MINIFS|TEXTJOIN|IFS|SWITCH)\s*\(",
                    re.IGNORECASE)


def _assert_no_banned_functions(wb):
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and isinstance(cell.value, str):
                    n += 1
                    _assert(not BANNED.search(cell.value), f"{ws.title}!{cell.coordinate} uses a 365-only function")
    for dv_ws in wb.worksheets:
        for dv in dv_ws.data_validations.dataValidation:
            _assert(not BANNED.search(str(dv.formula1 or "")), "validation formula")
    _assert(n > 100, "formulas were scanned")


def _assert_month_consistency(wb, model):
    """Month Drill-down total = Monthly Trend row = SUM of the month's Charged log rows."""
    headers, rows = _table(wb[W.TREND], "tblTrendSummary")
    month = lambda v: datetime.strptime(_display(v), "%b %Y").strftime("%Y-%m")  # noqa: E731  (a link to Drill-down)
    trend = {month(r[0]): r[headers.index("Credits")] for r in rows}
    trend_gens = {month(r[0]): r[headers.index("Generations")] for r in rows}
    drill = {}
    for row in wb["Month Drill-down"].iter_rows(values_only=True):
        if row and isinstance(row[0], str) and len(row) > 4 and row[1] == "Credits":
            try:
                month = datetime.strptime(row[0], "%b %Y").strftime("%Y-%m")
            except ValueError:
                continue                     # a table header such as "Department | Credits | ..."
            drill[month] = (row[2], row[4])
    lh, lrows = _table(wb[W.LOG])
    log = {}
    for r in lrows:
        if r[lh.index("Charged")] is True:
            key = r[lh.index("Month")].strftime("%Y-%m")
            c, g = log.get(key, (0.0, 0))
            log[key] = (c + float(r[lh.index("Credits")] or 0), g + 1)
    _assert(set(trend) == set(drill) == set(model.months), f"months {set(trend)} {set(drill)}")
    for m in model.months:
        _close(trend[m], drill[m][0], f"{m}: Monthly Trend vs Drill-down")
        _close(trend[m], log.get(m, (0, 0))[0], f"{m}: Monthly Trend vs log")
        _assert(trend_gens[m] == drill[m][1] == log.get(m, (0, 0))[1], f"{m}: generations agree")
    print("ok  month totals agree: trend = drill-down = log, for", ", ".join(model.months))


SOFFICE_CANDIDATES = (r"C:\Program Files\LibreOffice\program\soffice.com", r"C:\Program Files\LibreOffice\program\soffice.exe",
                      "/usr/bin/soffice", "/usr/bin/libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice")


def _soffice():
    import shutil
    for path in SOFFICE_CANDIDATES:
        if os.path.exists(path):
            return path
    return shutil.which("soffice") or shutil.which("libreoffice")


def _recalc(src: str) -> "Workbook":
    """Open in LibreOffice headless, recalculate, save as xlsx, read the values."""
    import subprocess
    out = tempfile.mkdtemp()
    profile = tempfile.mkdtemp()
    subprocess.run([_soffice(), f"-env:UserInstallation=file:///{profile.replace(os.sep, '/')}", "--headless",
                    "--norestore", "--convert-to", "xlsx", "--outdir", out, src],
                   check=True, timeout=300, capture_output=True)
    return load_workbook(os.path.join(out, os.path.basename(src)), data_only=True)


def _with_inputs(path: str, **cells) -> str:
    wb = load_workbook(path)
    ws = wb[W.PERIOD]
    for ref, value in cells.items():
        ws[ref] = value
    fd, out = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(out)
    return out


def test_period_explorer_recalc(path, wb):
    """Recalculate with LibreOffice: defaults = Home total; one month = that
    month's Monthly Trend total; Department X = By Department row X."""
    if not _soffice():
        print("SKIP period explorer recalculation: LibreOffice not installed")
        return
    home_total = wb[W.HOME]["A11"].value
    vals = _recalc(path)
    pe = vals[W.PERIOD]
    _close(pe["B20"].value, home_total, "Period Explorer default total = Home total")
    _assert(pe["B21"].value == wb[W.HOME]["C11"].value, "default generations = Home generations")
    _assert(isinstance(pe["B27"].value, str) and pe["B27"].value not in ("", "—"), f"top user: {pe['B27'].value}")

    headers, rows = _table(wb[W.TREND], "tblTrendSummary")
    sept = next(r for r in rows if _display(r[0]) == "Sep 2026")
    one_month = _recalc(_with_inputs(path, B5=datetime(2026, 9, 1), B6=datetime(2026, 9, 30)))
    _close(one_month[W.PERIOD]["B20"].value, sept[headers.index("Credits")], "From/To = Sep -> Sep's Monthly Trend total")

    dh, drows = _table(wb[W.DEPT])
    design = next(r for r in drows if _display(r[0]) == "Design")
    dept = _recalc(_with_inputs(path, B7="Design"))
    _close(dept[W.PERIOD]["B20"].value, design[dh.index("Credits")], "Department = Design -> By Department row")

    pending = _recalc(_with_inputs(path, B11="Yes"))
    _assert(pending[W.PERIOD]["B20"].value > home_total, "Include Pending adds pending credits")

    # Force a tie: zero every log credit, then give two people's charged rows 5 each.
    tie_wb = load_workbook(path)
    log = tie_wb[W.LOG]
    heads = [c.value for c in log[W.HEADER_ROW]]
    ci, si, ui = heads.index("Credits") + 1, heads.index("Charged") + 1, heads.index("User") + 1
    chosen = {}
    for r in range(W.FIRST_DATA_ROW, log.max_row + 1):
        if log.cell(r, ci).value is None or not isinstance(log.cell(r, ui).value, str):
            continue
        log.cell(r, ci).value = 0
        user = log.cell(r, ui).value
        if log.cell(r, si).value is True and user not in chosen and user not in ("Unassigned",) and len(chosen) < 2:
            chosen[user] = r
    for r in chosen.values():
        log.cell(r, ci).value = 5
    fd, tie_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    tie_wb.save(tie_path)
    top_user = _recalc(tie_path)[W.PERIOD]["B27"].value
    _assert(top_user.startswith("Tie: ") and all(name in top_user for name in chosen), f"tie shown: {top_user!r}")
    print("ok  Period Explorer recalculated in LibreOffice: defaults, one month, department, pending")


def test_trend_filtered_and_empty_workbooks():
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        model = _build(path, start=SEP1)
        _assert(model.months == ["2026-09", "2026-10"], f"months {model.months}")
        wb = load_workbook(path)
        headers, rows = _table(wb[W.TREND], "tblTrendToolCredits")
        sept = next(r for r in rows if _display(r[0]) == "Sep 2026")
        _close(sept[headers.index("Kling")], 70, "September Kling credits")
        _validate_links(wb)
        _validate_navigation(wb)
        _assert_month_consistency(wb, model)
        _assert_no_banned_functions(wb)
        test_period_explorer_recalc(path, wb)
        _validate_blocks(wb)

        _build(path, department="Design", client="Acme")
        wb = load_workbook(path)
        text = " ".join(str(c.value) for row in wb[W.HOME].iter_rows() for c in row if c.value)
        _assert("Department: Design" in text and "Client: Acme" in text, "applied filters printed")
        _close(wb[W.HOME]["A11"].value, 150, "Design + Acme charged total")
        _validate_links(wb)
        _validate_navigation(wb, expect_links=False)

        _build(path, exclude_test_accounts=False)
        wb = load_workbook(path)
        _close(wb[W.HOME]["A11"].value, CHARGED_TOTAL + ADMIN_CREDITS, "exclusion off")
        text = " ".join(str(c.value) for row in wb[W.HOME].iter_rows() for c in row if c.value)
        _assert("Included (the export option was turned off)" in text, "exclusion off printed")

        model = _build(path, start="2020-01-01", end="2020-01-31")
        _assert(model.is_empty, "no data in range")
        wb = load_workbook(path)
        text = " ".join(str(c.value) for row in wb[W.HOME].iter_rows() for c in row if c.value)
        _assert("No credit activity was found" in text, "friendly empty message on Home")
        pe = " ".join(str(c.value) for row in wb[W.PERIOD].iter_rows() for c in row if c.value)
        _assert("nothing to explore" in pe, "Period Explorer explains there is no data")
        for name in W.SHEET_ORDER[1:]:
            if name in (W.QUALITY, W.PERIOD):
                continue
            _assert(wb[name]["A5"].value == W.EMPTY_MESSAGE or wb[name]["A4"].value == W.EMPTY_MESSAGE,
                    f"{name} empty message")
        _validate_links(wb)
        _validate_navigation(wb, expect_links=False)
    finally:
        os.remove(path)
    print("ok  trend, filtered and empty workbooks")


# --------------------------------------------------------------------------- #
# API
# --------------------------------------------------------------------------- #
def test_api():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    import routers.credit_report_router as R
    from database_config import get_operational_db
    from utils.permissions import get_current_user

    os.environ["PUBLIC_DASHBOARD_URL"] = DASHBOARD
    R.OperationalSessionLocal = SessionLocal
    app = FastAPI()
    app.include_router(R.router)
    who = {"id": IDS["Plain"]}

    def _db():
        db = SessionLocal()
        try:
            yield db
        finally:
            db.close()

    def _user():
        with SessionLocal() as db:
            u = db.query(User).filter(User.id == who["id"]).one()
            u.feature_grants, u.role_assignments  # load before the session closes
            db.expunge(u)
            return u

    app.dependency_overrides[get_operational_db] = _db
    app.dependency_overrides[get_current_user] = _user
    client = TestClient(app)

    _assert(client.get("/api/reports/credit/export.xlsx").status_code == 403, "ungranted user is refused")
    # Output links only need a signed-in account.
    r = client.get(f"/api/reports/credit/output/Kling/{IDS['kling_event']}")
    _assert(r.status_code == 200 and r.json()["url"] == "https://cdn.kling.example/k1.mp4", "output lookup")
    _assert(client.get("/api/reports/credit/output/Kling/999999").status_code == 404, "missing output 404")
    with SessionLocal() as db:
        suno_id = db.query(SunoGeneration.id).filter(SunoGeneration.provider_creation_id == "s1").scalar()
    r = client.get(f"/api/reports/credit/output/Suno/{suno_id}")
    _assert(r.status_code == 200 and r.json()["url"].startswith("https://"), "Suno falls back without R2")

    with SessionLocal() as db:
        db.add(UserFeatureAccess(user_id=IDS["Plain"], feature="credit_report"))
        db.commit()
    r = client.get("/api/reports/credit/export.xlsx", params={"start": OCT1, "end": OCT7})
    _assert(r.status_code == 200, f"granted user downloads: {r.status_code} {r.text[:200]}")
    _assert('filename="credit-report_2026-10-01_to_2026-10-07.xlsx"' in r.headers["content-disposition"],
            r.headers["content-disposition"])
    _assert(b"localhost" not in r.content, "no localhost anywhere in the file")
    r = client.get("/api/reports/credit/export.xlsx", params={"excludeTestAccounts": "false"})
    _assert(r.status_code == 200, "exclusion can be turned off")
    for params in ({"start": "07-10-2026"}, {"start": OCT7, "end": OCT1}, {"tool": "Nope"}):
        _assert(client.get("/api/reports/credit/export.xlsx", params=params).status_code == 400, f"rejects {params}")
    opts = client.get("/api/reports/credit/options").json()
    _assert(opts["dashboardUrl"] == DASHBOARD, "options report the link base")
    _assert(opts["users"][0]["id"] == UNASSIGNED_USER_ID, "Unassigned is a user option")
    _assert({a["name"] for a in opts["excludedAccounts"]} == {"Admin", "RoleAdmin"}, "excluded accounts listed")
    who["id"] = IDS["Admin"]
    _assert(client.get("/api/reports/credit/options").status_code == 200, "admins bypass the grant")
    print("ok  api")


def main() -> int:
    seed()
    test_status_mapping()
    test_group_queries()
    test_filters_and_exclusion()
    test_model()
    test_duplicate_detection()
    test_log_columns_dropped_when_empty()
    test_block_rules()
    test_workbook_end_to_end()
    test_trend_filtered_and_empty_workbooks()
    test_api()
    print("ALL CREDIT REPORT CHECKS PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
