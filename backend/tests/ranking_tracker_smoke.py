"""keyword_ranking sheets: parsing, normalisation, movement, importer, job reports, API.

Run from backend/:  python tests/ranking_tracker_smoke.py
"""
import io
import os
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace

BACKEND_DIR = Path(__file__).resolve().parents[1]
os.environ.setdefault("DATABASE_URL", "postgresql://placeholder:placeholder@localhost:5432/placeholder")
os.environ.setdefault("ARCHIVE_DATABASE_URL", os.environ["DATABASE_URL"])
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

import database_config  # noqa: E402
from database_config import get_operational_db  # noqa: E402
from models_new import (  # noqa: E402
    Base, ContentRequest, Keyword, RankingCheck, RankingRun, RequestEvent, SheetActivity, SheetUserGoogleEmail,
    TrackedSheet, TrackedSheetMember, User, UserFeatureAccess, UserRole,
)
import services.sheets as S  # noqa: E402
from services.sheets import ranking_queries as RQ  # noqa: E402
from services.sheets.ranking import (  # noqa: E402
    infer_trigger, normalize_ai, normalize_position, parse_header_date, parse_ranking_workbook,
)
from services.sheets.ranking_sync import sync_ranking_sheet  # noqa: E402
from utils.permissions import get_current_user  # noqa: E402

engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base.metadata.create_all(bind=engine, tables=[m.__table__ for m in (
    User, UserRole, UserFeatureAccess, TrackedSheet, TrackedSheetMember, SheetUserGoogleEmail, ContentRequest,
    Keyword, RankingRun, RankingCheck, RequestEvent, SheetActivity)])
database_config.OperationalSessionLocal = SessionLocal

SHEET_ID = "1hWTNhWjUGlB_unIsmJdzg3Sf7IwHLNRr9Yr38t4R93U"
SOURCE_ID = "1SourceSheetAbcdefghijklmnopqrstuvwxyz0123"
TAB = "RitzMediaWorld"
SECRET = "s3cret-for-tests-only-0123456789"


def _assert(condition, message):
    if not condition:
        raise AssertionError(message)


# --------------------------------------------------------------------------- #
def test_normalisation():
    cases = {
        30: (30, "ranked"), 30.0: (30, "ranked"), "7": (7, "ranked"), " 12 ": (12, "ranked"),
        "Not in Top 10        ": (None, "not_in_top_10"), "not in top 50": (None, "not_in_top_50"),
        "NOT IN TOP 100": (None, "not_in_top_100"), "CHECK FAILED": (None, "failed"), "Check failed": (None, "failed"),
        "": (None, "blank"), None: (None, "blank"), "N5": (None, "unparsed"), 0: (None, "unparsed"),
        3.5: (None, "unparsed"), "abc": (None, "unparsed"), True: (None, "unparsed"),
    }
    for raw, (pos, bucket) in cases.items():
        got = normalize_position(raw)
        _assert((got.position, got.bucket) == (pos, bucket), f"position {raw!r}: {got}")
    _assert(normalize_position("Not in Top 10").depth == 10, "depth kept")
    ai = {"Yes": "cited", "yes ": "cited", "No": "not_cited", "No AI Overview": "none", " no  ai overview ": "none",
          "Check failed": "failed", "N5": "unparsed", "": "blank", None: "blank"}
    for raw, state in ai.items():
        _assert(normalize_ai(raw) == state, f"ai {raw!r} -> {normalize_ai(raw)}")
    _assert(parse_header_date(datetime(2026, 8, 24))[0] == date(2026, 8, 24), "real date header")
    d, label, baseline = parse_header_date("21-Aug-2026 (Baseline)")
    _assert(d == date(2026, 8, 21) and baseline and label.startswith("21-Aug"), "baseline text header")
    for text, want in (("7-Oct-2026", date(2026, 10, 7)), ("2026-10-07", date(2026, 10, 7)),
                       ("7 October 2026", date(2026, 10, 7)), ("next week", None)):
        _assert(parse_header_date(text)[0] == want, f"date {text!r}")
    _assert(infer_trigger(date(2026, 10, 5), False) == "scheduled" and infer_trigger(date(2026, 10, 7), False) == "manual"
            and infer_trigger(date(2026, 8, 21), True) == "baseline", "trigger inference")
    print("ok  value normalisation and date headers")


def build_workbook(runs, rows, *, extra_tab=True, bad_header=False):
    """runs: [(header value, {row index: (pos, ai)})]; rows: [(sno, keyword, url)]."""
    wb = Workbook()
    ws = wb.active
    ws.title = TAB
    ws["A1"] = "RitzMediaWorld - Keyword Ranking Tracker"
    ws["A2"] = "Auto-synced weekly (Mondays, 9 AM) from the Keyword URL Source Sheet."
    ws.merge_cells("A1:E1")
    ws.cell(5, 1, "S.No")
    ws.cell(5, 2, "Keyword")
    ws.cell(5, 3, "Live URL")
    col = 4
    for header, values in runs:
        ws.cell(5, col, header)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        ws.cell(6, col, "Position")
        ws.cell(6, col + 1, "AI Overview")
        for i, (pos, ai) in values.items():
            ws.cell(7 + i, col, pos)
            ws.cell(7 + i, col + 1, ai)
        col += 2
    if bad_header:
        ws.cell(5, col, "someday")
        ws.cell(6, col, "Position")
        ws.cell(6, col + 1, "AI Overview")
    for i, (sno, keyword, url) in enumerate(rows):
        ws.cell(7 + i, 1, sno)
        ws.cell(7 + i, 2, keyword)
        ws.cell(7 + i, 3, url)
    if extra_tab:
        wb.create_sheet("Notes")["A1"] = "not a ranking tab"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


KW = [(1, "best creative agency", "https://ritzmediaworld.com/a"),
      (2, "creative agencies", "https://ritzmediaworld.com/b"),
      (3, "best creative agency", "https://ritzmediaworld.com/a"),     # duplicate on purpose
      (4, "seo company noida", "https://ritzmediaworld.com/c")]
R1 = ("21-Aug-2026 (Baseline)", {0: ("Not in Top 10        ", "No"), 1: (4, "Yes"), 2: (9, "No AI Overview"), 3: (2, "Yes")})
R2 = (datetime(2026, 8, 24), {0: (8, "No"), 1: (3, "Yes"), 2: (12.0, "N5"), 3: (9, "Yes")})
R3 = (datetime(2026, 8, 27), {0: (2, "Yes"), 1: ("CHECK FAILED", "Check failed"), 2: ("Not in Top 50", "No"), 3: ("Not in Top 50", "No")})


def test_parsing():
    tabs = parse_ranking_workbook(build_workbook([R1, R2, R3], KW, bad_header=True))
    t = tabs[TAB]
    _assert((t.header_row, t.sub_header_row, t.first_data_row) == (5, 6, 7), "rows found by header text")
    _assert(t.columns == {"sno": 1, "keyword": 2, "url": 3}, f"fixed columns {t.columns}")
    _assert([r.run_key for r in t.runs] == ["2026-08-21", "2026-08-24", "2026-08-27"], "date pairs")
    _assert(t.runs[0].is_baseline and t.runs[0].letter == "D" and t.runs[1].ai_col == 7, "pair columns")
    _assert(any("unreadable run date 'someday'" in p for p in t.problems), f"bad header flagged: {t.problems}")
    _assert(len(t.rows) == 4 and t.rows[0].cells["2026-08-21"][0].startswith("Not in Top 10"), "rows and cells")
    _assert(not tabs["Notes"].is_ranking, "other tabs ignored")
    # The sheet grows to the right: a 4th run is picked up with no config change.
    more = parse_ranking_workbook(build_workbook([R1, R2, R3, (datetime(2026, 8, 27), {0: (1, "Yes")})], KW))[TAB]
    _assert([r.run_key for r in more.runs][-1] == "2026-08-27#2", "second run on the same day gets its own key")
    # Header on another row still found.
    wb = Workbook()
    ws = wb.active
    ws.title = "X"
    ws.append(["Keyword", "Live URL", "1-Oct-2026", None])
    ws.append([None, None, "Position", "AI Overview"])
    ws.append(["kw", "u", 5, "Yes"])
    buf = io.BytesIO()
    wb.save(buf)
    x = parse_ranking_workbook(buf.getvalue())["X"]
    _assert(x.header_row == 1 and x.is_ranking and x.rows[0].cells["2026-10-01"] == (5, "Yes"), "header row 1 layout")
    # A wrong header-row setting (e.g. the title row) never stops the import.
    wrong = parse_ranking_workbook(build_workbook([R1, R2, R3], KW), {"headerRow": 1, "subHeaderRow": 2, "firstDataRow": 3})[TAB]
    _assert((wrong.header_row, wrong.sub_header_row, wrong.first_data_row, len(wrong.rows)) == (5, 6, 7, 4),
            "wrong row settings fall back to header-text detection")
    _assert(any("using row 5" in p for p in wrong.problems), "and the fallback is reported")
    print("ok  header and date-pair parsing")


def test_movement():
    c = lambda pos, bucket="ranked", depth=None: SimpleNamespace(position=pos, bucket=bucket, depth=depth)  # noqa: E731
    out10, out50 = c(None, "not_in_top_10", 10), c(None, "not_in_top_50", 50)
    cases = [
        (c(8), c(3), (5, "improved")), (c(3), c(9), (-6, "dropped")), (c(4), c(4), (0, "same")),
        (out10, c(5), (6, "improved")), (c(5), out10, (-6, "dropped")), (out50, out50, (0, "same")),
        (out10, c(30), (None, "not_comparable")), (c(30), out10, (None, "not_comparable")),
        (out10, out50, (None, "not_comparable")), (c(None, "failed"), c(3), (None, "unknown")),
        (None, c(3), (None, "new")),
    ]
    for prev, cur, (delta, direction) in cases:
        got = RQ.movement(prev, cur)
        _assert((got["delta"], got["direction"]) == (delta, direction), f"{prev} -> {cur}: {got}")
    _assert(RQ.url_mismatch("https://www.a.com/x/", "https://a.com/x") is False, "same page, different form")
    _assert(RQ.url_mismatch("https://a.com/x", "https://a.com/y") is True, "different page")
    runs = [SimpleNamespace(check_date=d) for d in (date(2026, 8, 21), date(2026, 8, 24), date(2026, 9, 7))]
    _assert(RQ.missed_mondays(runs, datetime(2026, 9, 14, 3)) == [date(2026, 8, 31)],
            "31 Aug missed; 14 Sep (09:00 IST + grace) not due yet")
    print("ok  movement, URL match, missed Mondays")


def _sheet(db):
    s = TrackedSheet(name="keyword_ranking_tracker", sheet_type="keyword_ranking", spreadsheet_id=SHEET_ID, url="u",
                     source_spreadsheet_id=SOURCE_ID, source_url="src",
                     tabs_json=[{"name": TAB, "tracked": True}, {"name": "Notes", "tracked": False}],
                     mapping_json={}, statuses_json={}, config_json={})
    db.add(s)
    db.commit()
    return s


def _events(db, kind):
    return db.query(RequestEvent).filter(RequestEvent.event_type == kind).order_by(RequestEvent.id).all()


def test_importer_lifecycle():
    t0 = datetime(2026, 8, 28, 6)
    with SessionLocal() as db:
        db.add(User(email="ravi@ritzmediaworld.com", name="Ravi", hashed_password="x"))
        db.commit()
        s = _sheet(db)
        st = sync_ranking_sheet(db, s, parse_ranking_workbook(build_workbook([R1, R2, R3], KW)), t0)
        _assert(st["keywords"] == 4 and st["newRuns"] == 3 and st["checks"] == 12, f"backfill {st}")
        _assert(len(_events(db, "IMPORTED")) == 4 and len(_events(db, "RUN_RECORDED")) == 3, "backfill events")
        runs = {r.run_key: r for r in db.query(RankingRun)}
        _assert((runs["2026-08-21"].trigger, runs["2026-08-24"].trigger, runs["2026-08-27"].trigger) ==
                ("baseline", "scheduled", "manual"), "triggers")
        _assert((runs["2026-08-27"].keywords_checked, runs["2026-08-27"].failed_count) == (4, 1), "run counts")
        kws = db.query(Keyword).order_by(Keyword.row_number).all()
        _assert([k.row_key for k in kws][:3] == ["best creative agency#1", "creative agencies#1", "best creative agency#2"],
                "duplicates kept as separate rows")
        chk = {(c.keyword_id, c.run_id): c for c in db.query(RankingCheck)}
        dup = chk[(kws[2].id, runs["2026-08-24"].id)]
        _assert(dup.position == 12 and dup.ai_state == "unparsed" and dup.raw_ai == "N5", "dirty value kept, flagged")

        # A person adds a keyword in the source sheet; the job copies it here.
        t1 = t0 + timedelta(days=3)
        db.add(SheetActivity(event_id="evt-src-000001", timestamp=t1 - timedelta(hours=5),
                             user_email="ravi@ritzmediaworld.com", spreadsheet_id=SOURCE_ID, sheet_name="Ritz",
                             row=9, column=2, num_rows=1, num_columns=1, change_type="EDIT", new_value="Logo Design Noida "))
        db.commit()
        kw2 = KW[:3] + [(4, "seo company in noida", "https://ritzmediaworld.com/c-new"), (5, "logo design noida", "")]
        r4 = (datetime(2026, 8, 31), {0: (1, "Yes"), 1: (2, "Yes"), 2: (5, "No"), 3: (7, "No"), 4: (11, "No")})
        r1_edited = (R1[0], {**R1[1], 1: (1, "Yes")})          # someone typed over a 10-day-old result
        r3_changed = (R3[0], {**R3[1], 0: (3, "Yes")})          # 4-day-old run changed too
        st = sync_ranking_sheet(db, s, parse_ranking_workbook(build_workbook([r1_edited, R2, r3_changed, r4], kw2)), t1)
        _assert((st["added"], st["edited"], st["newRuns"]) == (1, 1, 1), f"second poll {st}")
        added = _events(db, "KEYWORD_ADDED")[0]
        _assert(added.actor_email == "ravi@ritzmediaworld.com" and added.occurred_at == t1 - timedelta(hours=5),
                "adder named from the source sheet edit")
        edited = _events(db, "KEYWORD_EDITED")[0]
        _assert(edited.old_value == "seo company noida" and edited.new_value == "seo company in noida", "edit old -> new")
        _assert(edited.keyword_id == kws[3].id, "edited keyword keeps its id and history")
        url_ev = _events(db, "URL_CHANGED")
        _assert(len(url_ev) == 1 and url_ev[0].actor_type == "seo_api", "Live URL is the job's")
        new_run = db.query(RankingRun).filter_by(run_key="2026-08-31").one()
        page = db.query(RankingCheck).filter_by(keyword_id=kws[3].id, run_id=new_run.id).one().found_url
        _assert(page == "https://ritzmediaworld.com/c-new", "new run records the Live URL as its ranking page")
        old_pages = [c.found_url for c in db.query(RankingCheck).filter_by(keyword_id=kws[1].id)
                     if c.run_id in (runs["2026-08-21"].id, runs["2026-08-24"].id)]
        _assert(old_pages == [None, None], "backfill never credits today's Live URL to older runs")
        overrides = _events(db, "MANUAL_RESULT_OVERRIDE")
        _assert(len(overrides) == 2 and all(o.actor_type == "unknown" for o in overrides), "old results changed -> overrides")

        # A recent result rewritten by the job is not an override.
        t2 = t1 + timedelta(hours=2)
        r4b = (datetime(2026, 8, 31), {**r4[1], 4: (10, "No")})
        sync_ranking_sheet(db, s, parse_ranking_workbook(build_workbook([r1_edited, R2, r3_changed, r4b], kw2)), t2)
        _assert(len(_events(db, "RESULT_UPDATED")) == 1, "same-day change = job update")

        # A human edit captured on the result cell names the overrider.
        db.add(SheetActivity(event_id="evt-cell-000001", timestamp=t2, user_email="ravi@ritzmediaworld.com",
                             spreadsheet_id=SHEET_ID, sheet_name=TAB, row=8, column=6, num_rows=1, num_columns=1,
                             change_type="EDIT", new_value="2"))
        db.commit()
        r2_edit = (R2[0], {**R2[1], 1: (2, "Yes")})
        sync_ranking_sheet(db, s, parse_ranking_workbook(build_workbook([r1_edited, r2_edit, r3_changed, r4b], kw2)),
                           t2 + timedelta(minutes=2))
        named = _events(db, "MANUAL_RESULT_OVERRIDE")[-1]
        _assert(named.actor_type == "user" and named.actor_email == "ravi@ritzmediaworld.com", "override attributed")

        no_url = [a for a in RQ.alerts(db, s, now=t2) if a["type"] == "RANKED_WITHOUT_URL"]
        _assert([a["keyword"] for a in no_url] == ["logo design noida"], f"ranked but no Live URL: {no_url}")

        # Removal.
        sync_ranking_sheet(db, s, parse_ranking_workbook(build_workbook([r1_edited, r2_edit, r3_changed, r4b], kw2[:4])),
                           t2 + timedelta(minutes=5))
        _assert(len(_events(db, "KEYWORD_REMOVED")) == 1, "removed keyword")

        rows = RQ.keyword_rows(db, s, RQ.KeywordFilters())
        first = next(r for r in rows if r["keyword"] == "best creative agency")
        # previous run's value was overridden to 3; latest is 1
        _assert(first["change"] == {"delta": 2, "direction": "improved"} and first["best"] == 1, f"movement {first['change']}")
        _assert(first["isDuplicate"] and first["sparkline"][0]["v"] == 11, "duplicate flag; Not in Top 10 plotted as 11")
        alerts = RQ.alerts(db, s, now=datetime(2026, 9, 1, 12))
        kinds = {a["type"] for a in alerts}
        _assert({"DUPLICATE_KEYWORD", "UNPARSEABLE_VALUE", "MANUAL_OVERRIDE"} <= kinds, f"alerts {kinds}")
        people = RQ.contributors(db, s)
        ravi = next(p for p in people if p["email"] == "ravi@ritzmediaworld.com")
        unknown = next(p for p in people if p["email"] is None and not p.get("imported"))
        _assert(ravi["addedCount"] == 1 and ravi["name"] == "Ravi", "who added what")
        _assert((unknown["edits"], unknown["removed"]) == (1, 1), "unattributed edits counted under Unknown")
        _assert(people[-1]["name"] == "Before tracking" and people[-1]["addedCount"] == 4, "imported bucket")
        hist = "".join(RQ.history_csv(db, s, RQ.KeywordFilters(q="creative agencies")))
        _assert(hist.count("creative agencies") == 4, "history CSV: one line per run")
    print("ok  importer: backfill, add/edit/remove, URL, overrides, alerts, CSV")


def test_job_report_api():
    app = FastAPI()
    import routers.sheets_router as R
    app.include_router(R.router)
    with SessionLocal() as db:
        admin = User(email="boss@ritzmediaworld.com", name="Boss", hashed_password="x", is_admin=True, is_active=True)
        viewer = User(email="viewer@ritzmediaworld.com", name="Viewer", hashed_password="x", is_active=True)
        db.add_all([admin, viewer])
        db.flush()
        db.add(UserFeatureAccess(user_id=viewer.id, feature="sheet_activity"))
        db.commit()
        ids = {"admin": admin.id, "viewer": viewer.id}
    who = {"id": ids["admin"]}

    def _db():
        db = SessionLocal()
        try:
            yield db
        finally:
            db.close()

    def _user():
        with SessionLocal() as db:
            u = db.query(User).filter(User.id == who["id"]).one()
            u.feature_grants, u.role_assignments
            db.expunge(u)
            return u

    app.dependency_overrides[get_operational_db] = _db
    app.dependency_overrides[get_current_user] = _user
    client = TestClient(app)
    url = "/api/sheets/ranking-results"
    report = {"spreadsheetId": SHEET_ID, "tab": TAB, "runDate": "2026-08-31", "runId": "run-0831", "trigger": "scheduled",
              "api": "DataForSEO", "location": "Noida,Uttar Pradesh,India", "gl": "in", "cost": 0.012,
              "durationSeconds": 41.5, "results": [
                  {"keyword": "best creative agency", "position": 1, "positionRaw": "1", "aiOverview": "Yes",
                   "foundUrl": "https://ritzmediaworld.com/other-page", "raw": {"items": ["x" * 9000]}},
                  {"keyword": "brand new keyword", "positionRaw": "Not in Top 50", "aiOverview": "No AI Overview"},
                  {"keyword": "creative agencies", "error": "timeout"}]}
    os.environ.pop("SHEET_ACTIVITY_WEBHOOK_SECRET", None)
    _assert(client.post(url, json=report).status_code == 503, "unset secret refuses everything")
    os.environ["SHEET_ACTIVITY_WEBHOOK_SECRET"] = SECRET
    _assert(client.post(url, json={"bad": True}).status_code == 401, "no secret -> 401 before body validation")
    h = {"X-Sheet-Activity-Secret": SECRET}
    _assert(client.post(url, json={"bad": True}, headers=h).status_code == 422, "malformed report -> 422")
    _assert(client.post(url, json={**report, "spreadsheetId": "x" * 30}, headers=h).status_code == 404, "unknown sheet")
    _assert(client.post(url, json={**report, "tab": "Notes"}, headers=h).status_code == 400, "untracked tab")
    r = client.post(url, json=report, headers=h)
    _assert(r.status_code == 200, f"report accepted: {r.text[:200]}")
    body = r.json()
    _assert(body["runCreated"] is False and body["newKeywords"] == ["brand new keyword"], f"merged into importer run {body}")
    with SessionLocal() as db:
        run = db.get(RankingRun, body["runId"])
        _assert((run.api, run.cost, run.duration_seconds, run.external_run_id, run.source) ==
                ("DataForSEO", 0.012, 41.5, "run-0831", "job"), "run details from the job")
        _assert(run.failed_count >= 1, "error result counted as failed")
        kw = db.query(Keyword).filter_by(row_key="best creative agency#1").one()
        c = db.query(RankingCheck).filter_by(keyword_id=kw.id, run_id=run.id).one()
        _assert(c.found_url.endswith("other-page") and c.raw_response.get("truncated"), "found URL and trimmed raw")
        sheet_id = db.query(TrackedSheet).filter_by(spreadsheet_id=SHEET_ID).one().id
    again = client.post(url, json=report, headers=h).json()
    _assert(again["runId"] == body["runId"], "same runId reported twice -> same run")

    alerts = client.get(f"/api/sheets/{sheet_id}/ranking/alerts").json()["alerts"]
    _assert(any(a["type"] == "RANKING_PAGE_CHANGED" for a in alerts), "ranking page changed between runs")
    ov = client.get(f"/api/sheets/{sheet_id}/ranking/overview").json()
    _assert(ov["sites"][0]["tab"] == TAB and ov["sites"][0]["latest"]["checkDate"] == "2026-08-31", "overview")
    kws = client.get(f"/api/sheets/{sheet_id}/ranking/keywords", params={"movement": "improved"}).json()
    _assert(kws["total"] >= 1 and all(k["change"]["direction"] == "improved" for k in kws["items"]), "movement filter")
    detail = client.get(f"/api/sheets/{sheet_id}/ranking/keywords/{kws['items'][0]['id']}").json()
    _assert(detail["history"] and detail["events"], "keyword detail")
    _assert(client.get(f"/api/sheets/{sheet_id}/ranking/runs").json()["runs"][0]["api"] == "DataForSEO", "runs log")
    _assert(client.get(f"/api/sheets/{sheet_id}/ranking/history.csv").status_code == 200, "history CSV")
    _assert(client.get(f"/api/sheets/{sheet_id}/requests").status_code == 200, "shared endpoints still answer")
    who["id"] = ids["viewer"]
    _assert(client.get(f"/api/sheets/{sheet_id}/ranking/overview").status_code == 404, "unassigned viewer can't open it")
    print("ok  job report webhook and ranking API")


def main() -> int:
    test_normalisation()
    test_parsing()
    test_movement()
    test_importer_lifecycle()
    test_job_report_api()
    print("ALL RANKING CHECKS PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
