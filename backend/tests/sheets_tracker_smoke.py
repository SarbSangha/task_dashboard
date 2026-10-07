"""Sheets section: classification, snapshot diff, workbook parsing, access.

Run from backend/:  python tests/sheets_tracker_smoke.py

The webhook (secret, payload validation) is covered by
tests/sheet_activity_smoke.py: the Sheets section reuses that endpoint.
"""
import io
import os
import sys
from datetime import datetime, timedelta
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
os.environ.setdefault("DATABASE_URL", "postgresql://placeholder:placeholder@localhost:5432/placeholder")
os.environ.setdefault("ARCHIVE_DATABASE_URL", os.environ["DATABASE_URL"])
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import httpx  # noqa: E402
from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

import database_config  # noqa: E402
from database_config import get_operational_db  # noqa: E402
from models_new import (  # noqa: E402
    Base, ContentRequest, RequestEvent, SheetActivity, SheetUserGoogleEmail, TrackedSheet, TrackedSheetMember,
    User, UserFeatureAccess, UserRole,
)
import services.sheets as S  # noqa: E402
from services.sheets import mapping as M  # noqa: E402
from services.sheets import queries as Q  # noqa: E402
from services.sheets.public_export import RowData, SheetAccessError, TabData, fetch_workbook_bytes, parse_spreadsheet_id, parse_workbook  # noqa: E402
from services.sheets.tracker import classify_row_change, sync_sheet  # noqa: E402
from utils.permissions import get_current_user  # noqa: E402

engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base.metadata.create_all(bind=engine, tables=[m.__table__ for m in (
    User, UserRole, UserFeatureAccess, TrackedSheet, TrackedSheetMember, SheetUserGoogleEmail, ContentRequest,
    RequestEvent, SheetActivity)])
database_config.OperationalSessionLocal = SessionLocal

HEADERS = ["No.", "Topic", "Primary Keyword", "Structure Doc (Google Doc Link)", "Reference Link(s)",
           "Optimized Structure", "Status", "Final Doc Link"]
MAPPING = M.auto_mapping(HEADERS)
STATUSES = dict(M.DEFAULT_STATUSES)
SHEET_ID = "1xV60yuy8JUZIKcT17wuYXp7mNy6aC_5GSQ6pyb8LDf0"
TAB = "Ritz Media World"
T0 = datetime(2026, 10, 1, 6, 0)


def _assert(condition, message):
    if not condition:
        raise AssertionError(message)


def types(events):
    return [(e.event_type, e.actor_type) for e in events]


# --------------------------------------------------------------------------- #
def test_mapping_and_url():
    m = MAPPING
    _assert(m["No."]["role"] == "row_number" and m["Topic"] == {"role": "topic", "side": "input"}, "No. / Topic")
    _assert(m["Reference Link(s)"]["side"] == "input" and m["Primary Keyword"]["side"] == "input", "inputs")
    _assert(m["Optimized Structure"]["role"] == "structure" and m["Final Doc Link"]["role"] == "final_link", "outputs")
    _assert(M.pick_key_column(m, M.DEFAULT_CONFIG) == "No.", "No. is the row key without a Request ID column")
    _assert(M.pick_key_column(M.auto_mapping(HEADERS + ["Request ID"]), M.DEFAULT_CONFIG) == "Request ID",
            "Request ID wins when the Apps Script added it")
    for url in (f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit?gid=0#gid=0",
                f"https://docs.google.com/spreadsheets/u/1/d/{SHEET_ID}/edit", SHEET_ID):
        _assert(parse_spreadsheet_id(url) == SHEET_ID, f"id from {url}")
    try:
        parse_spreadsheet_id("https://example.com/x")
        raise AssertionError("bad URL accepted")
    except SheetAccessError:
        pass
    for bad in ({"Topic": {"role": "topic", "side": "input"}}, {"X": {"role": "nope", "side": "input"}}):
        try:
            M.validate_mapping(bad)
            raise AssertionError(f"accepted {bad}")
        except ValueError:
            pass
    print("ok  mapping by header name, URL parsing")


def test_classification():
    blank = {h: "" for h in HEADERS}
    row = {**blank, "No.": "4"}
    c = lambda old, new: types(classify_row_change(old, new, MAPPING, STATUSES))  # noqa: E731
    v1 = {**row, "Topic": "Sector 25 Noida", "Structure Doc (Google Doc Link)": "https://docs.google.com/document/d/a"}
    _assert(c(row, v1) == [("NEW_REQUEST", "user"), ("INPUT_EDIT", "user")], f"new request: {c(row, v1)}")
    v2 = {**v1, "Topic": "Sector 25 Noida rates"}
    _assert(c(v1, v2) == [("INPUT_EDIT", "user")], "topic edit")
    v3 = {**v2, "Optimized Structure": "# H1\n## H2", "Status": "Awaiting Approval"}
    out = classify_row_change(v2, v3, MAPPING, STATUSES)
    _assert(types(out) == [("STRUCTURE_GENERATED", "claude")] and out[0].new == "# H1\n## H2", "structure + status = one event")
    v4 = {**v3, "Optimized Structure": "# H1\n## H2 edited"}
    _assert(c(v3, v4) == [("STRUCTURE_EDITED_BY_USER", "user")], "user edits structure")
    v5 = {**v4, "Status": "approved "}
    _assert(c(v4, v5) == [("APPROVED", "user")], "approval, case and spaces ignored")
    v6 = {**v5, "Status": "In Progress"}
    _assert(c(v5, v6) == [("DRAFT_STARTED", "claude")], "draft started")
    v7 = {**v6, "Status": "Delivered", "Final Doc Link": "https://docs.google.com/open?id=x"}
    _assert(c(v6, v7) == [("DELIVERED", "claude")], "delivered with link")
    v7b = {**v6, "Status": "Delivered"}
    _assert(c(v6, v7b) == [("DELIVERED", "claude"), ("ERROR", "claude")], "delivered without link is an error")
    _assert(c(v6, {**v6, "Status": "Error: doc not shared"}) == [("ERROR", "claude")], "error status")
    _assert(c(v6, {**v6, "Status": "On hold"}) == [("OTHER_STATUS_CHANGE", "unknown")], "other status")
    _assert(c(v2, {**v2, "Optimized Structure": "draft"}) == [("STRUCTURE_GENERATED", "claude")],
            "structure appearing before the status still counts as Claude's")
    _assert(c(v7, v7) == [], "no change, no event")
    _assert(c(v7, {**v7, "No.": "99"}) == [], "ignored columns never create events")
    print("ok  event classification")


def _tab(rows, name=TAB):
    out = TabData(name=name, headers=list(HEADERS))
    for n, values in rows:
        out.rows.append(RowData(row_number=n, values={h: values.get(h, "") for h in HEADERS}))
    return {name: out}


def _sheet(db):
    s = TrackedSheet(name="Tracker", spreadsheet_id=SHEET_ID, url="u", tabs_json=[{"name": TAB, "tracked": True}],
                     mapping_json=MAPPING, statuses_json=STATUSES, config_json=dict(M.DEFAULT_CONFIG))
    db.add(s)
    db.commit()
    return s


def test_snapshot_diff_lifecycle():
    with SessionLocal() as db:
        db.add(User(email="ravi@ritzmediaworld.com", name="Ravi", hashed_password="x"))
        db.commit()
        s = _sheet(db)
        delivered_old = {"No.": "1", "Topic": "Old piece", "Status": "Delivered", "Final Doc Link": "https://d/1"}
        # First poll: existing rows are imported, empty pre-numbered rows ignored.
        st = sync_sheet(db, s, _tab([(2, delivered_old), (3, {"No.": "2"})]), T0)
        _assert(st["created"] == 1 and st["events"] == 1, f"import {st}")
        old = db.query(ContentRequest).one()
        _assert(old.is_imported and old.requested_at is None, "imported rows get no invented stage times")

        # A user starts a request; the Apps Script saw who.
        t1 = T0 + timedelta(minutes=10)
        db.add(SheetActivity(event_id="evt-topic-0001", timestamp=t1 - timedelta(minutes=1), user_email="ravi@ritzmediaworld.com",
                             spreadsheet_id=SHEET_ID, sheet_name=TAB, row=3, column=2, num_rows=1, num_columns=1,
                             change_type="EDIT", new_value="Sector 25"))
        db.commit()
        r2 = {"No.": "2", "Topic": "Sector 25", "Reference Link(s)": "https://a, https://b"}
        st = sync_sheet(db, s, _tab([(2, delivered_old), (3, r2)]), t1)
        req = db.query(ContentRequest).filter_by(row_key="2").one()
        evs = db.query(RequestEvent).filter_by(request_id=req.id).order_by(RequestEvent.id).all()
        _assert([(e.event_type, e.actor_type, e.actor_email) for e in evs] ==
                [("NEW_REQUEST", "user", "ravi@ritzmediaworld.com"), ("INPUT_EDIT", "user", None)], f"new {types(evs)}")
        _assert(evs[0].occurred_at == t1 - timedelta(minutes=1) and req.created_by_email == "ravi@ritzmediaworld.com",
                "exact time and creator from the Apps Script event")
        _assert(evs[1].metadata_json.get("approximateTime"), "unmatched change is timed at the poll")

        # Claude writes the structure (row moved: sorted to row 2).
        t2 = t1 + timedelta(hours=2)
        r3 = {**r2, "Optimized Structure": "# Sector 25\n## Rates", "Status": "Awaiting Approval"}
        sync_sheet(db, s, _tab([(2, r3), (3, delivered_old)]), t2)
        gen = db.query(RequestEvent).filter_by(request_id=req.id, event_type="STRUCTURE_GENERATED").one()
        _assert(gen.actor_type == "claude" and gen.metadata_json["inputs"]["Topic"] == "Sector 25", "inputs recorded")
        _assert(gen.metadata_json["durationSeconds"] == int((t2 - evs[0].occurred_at).total_seconds()), "duration")
        _assert(db.query(ContentRequest).count() == 2, "moved row keeps its request (keyed by No.)")

        # User edits the structure, approves; Claude drafts and delivers.
        t3, t4, t5, t6 = (t2 + timedelta(hours=h) for h in (1, 2, 3, 5))
        r4 = {**r3, "Optimized Structure": "# Sector 25\n## Rates\n## Metro"}
        sync_sheet(db, s, _tab([(2, r4), (3, delivered_old)]), t3)
        r5 = {**r4, "Status": "Approved"}
        sync_sheet(db, s, _tab([(2, r5), (3, delivered_old)]), t4)
        r6 = {**r5, "Status": "In Progress"}
        sync_sheet(db, s, _tab([(2, r6), (3, delivered_old)]), t5)
        r7 = {**r6, "Status": "Delivered", "Final Doc Link": "https://docs.google.com/open?id=z"}
        sync_sheet(db, s, _tab([(2, r7), (3, delivered_old)]), t6)
        db.refresh(req)
        seq = [e.event_type for e in db.query(RequestEvent).filter_by(request_id=req.id).order_by(RequestEvent.id)]
        _assert(seq == ["NEW_REQUEST", "INPUT_EDIT", "STRUCTURE_GENERATED", "STRUCTURE_EDITED_BY_USER", "APPROVED",
                        "DRAFT_STARTED", "DELIVERED"], f"timeline {seq}")
        _assert((req.awaiting_at, req.approved_at, req.in_progress_at, req.delivered_at) == (t2, t4, t5, t6), "stage times")
        _assert(req.status == "Delivered" and req.error is None, "final state")

        detail = Q.request_detail(db, s, req, now=t6)
        _assert(detail["claudeStructure"] == "# Sector 25\n## Rates", "Claude's version kept for the diff")
        _assert(detail["currentStructure"].endswith("## Metro"), "user's edited version")
        _assert([e["eventType"] for e in detail["inputEvents"]] == ["NEW_REQUEST", "INPUT_EDIT"], "input side")
        _assert(detail["request"]["createdBy"]["name"] == "Ravi", "creator mapped to a dashboard user")
        ov = Q.overview(db, s, now=t6)
        _assert(ov["avgHoursToAwaiting"] == round((t2 - evs[0].occurred_at).total_seconds() / 3600, 1), "avg to awaiting")
        _assert(ov["avgHoursApprovalToDelivered"] == 3.0 and ov["deliveredThisWeek"] == 1, "avg approval -> delivered")
        ravi = next(p for p in ov["perUser"] if p["email"] == "ravi@ritzmediaworld.com")
        _assert(ravi["created"] == 1 and ravi["name"] == "Ravi", "per-user stats")

        # Errors, stuck, removal.
        r8 = {"No.": "3", "Topic": "Waiting piece", "Optimized Structure": "# W", "Status": "Awaiting Approval"}
        r9 = {"No.": "4", "Topic": "Broken piece", "Status": "Delivered"}
        sync_sheet(db, s, _tab([(2, r7), (4, r8), (5, r9)]), t6)       # delivered_old row deleted
        later = t6 + timedelta(hours=49)
        ov = Q.overview(db, s, now=later)
        reasons = sorted(x["reason"] for x in ov["stuck"])
        _assert(reasons == ["Awaiting Approval for over 48 h", "Delivered with no Final Doc Link"], f"stuck {reasons}")
        _assert(ov["errors"] == 1, "Delivered with no link counted as an error")
        gone = db.query(ContentRequest).filter_by(row_key="1").one()
        _assert(gone.removed_at == t6 and ov["totalRequests"] == 3, "deleted row marked removed")
        stuck_list = Q.list_requests(db, s, Q.RequestFilters(stuck_only=True), 1, 50, now=later)
        _assert(stuck_list["total"] == 2, "stuck filter")
        _assert(Q.list_requests(db, s, Q.RequestFilters(q="sector"), 1, 50)["total"] == 1, "topic search")
        _assert(Q.list_requests(db, s, Q.RequestFilters(status="Delivered"), 1, 50)["total"] == 2, "status filter")
        _assert(Q.list_requests(db, s, Q.RequestFilters(user="unknown"), 1, 50)["total"] == 2, "unknown creator filter")
        csv_text = "".join(Q.requests_csv(db, s, Q.RequestFilters()))
        _assert("Sector 25" in csv_text and "Unknown user" in csv_text, "requests CSV")
        events_text = "".join(Q.events_csv(db, s, Q.RequestFilters(q="sector")))
        import csv as _csv
        parsed = list(_csv.reader(io.StringIO(events_text.lstrip("﻿"))))
        _assert(len(parsed) == 8 and parsed[3][7] == "STRUCTURE_GENERATED",
                "events CSV: header + 7 events, multi-line values intact")
    print("ok  snapshot diff lifecycle, stuck, removal, CSV")


def test_workbook_parsing_and_fetch():
    wb = Workbook()
    ws = wb.active
    ws.title = TAB
    ws.append(HEADERS)
    ws.append([1, "Example topic", "", "", "", "", "", ""])
    ws.append([2.0, "Real topic", "kw", "Open doc", "", "", "Delivered", "Final"])
    yellow = PatternFill("solid", fgColor="FFFFF2CC")
    for cell in ws[2]:
        cell.fill = yellow
    ws["D3"].hyperlink = "https://docs.google.com/document/d/abc"
    ws["H3"].hyperlink = "https://docs.google.com/open?id=final"
    buf = io.BytesIO()
    wb.save(buf)
    tabs = parse_workbook(buf.getvalue())
    rows = tabs[TAB].rows
    _assert(rows[0].is_example and not rows[1].is_example, "pale-yellow example row detected")
    _assert(rows[1].values["No."] == "2" and rows[1].values["Structure Doc (Google Doc Link)"] == "Open doc", "cell text")
    from services.sheets.tracker import row_values
    vals = row_values(rows[1], MAPPING)
    _assert(vals["Structure Doc (Google Doc Link)"] == "https://docs.google.com/document/d/abc", "link cell -> URL")
    _assert(vals["Final Doc Link"] == "https://docs.google.com/open?id=final" and vals["Topic"] == "Real topic", "links")

    def client(status, body, ctype="text/html"):
        return httpx.Client(transport=httpx.MockTransport(lambda req: httpx.Response(status, content=body,
                                                                                      headers={"content-type": ctype})))
    _assert(fetch_workbook_bytes("x" * 30, client=client(200, buf.getvalue())) == buf.getvalue(), "xlsx accepted")
    for status, body, needle in ((200, b"<html>Sign in</html>", "Anyone with the link"), (404, b"", "doesn't exist"),
                                 (429, b"", "rate-limiting"), (500, b"", "answered 500")):
        try:
            fetch_workbook_bytes("x" * 30, client=client(status, body))
            raise AssertionError(f"{status} accepted")
        except SheetAccessError as exc:
            _assert(needle in str(exc), f"{status}: {exc}")
    print("ok  workbook parsing, example rows, links, access errors")


def test_api_access():
    with SessionLocal() as db:
        ids = {}
        for name, admin in (("Boss", True), ("Member", False), ("Outsider", False), ("NoGrant", False)):
            u = User(email=f"{name.lower()}@ritzmediaworld.com", name=name, hashed_password="x", is_admin=admin, is_active=True)
            db.add(u)
            db.flush()
            ids[name] = u.id
        for name in ("Member", "Outsider"):
            db.add(UserFeatureAccess(user_id=ids[name], feature="sheet_activity"))
        db.add(SheetActivity(event_id="evt-scope-0001", timestamp=T0, spreadsheet_id="other-spreadsheet-id",
                             sheet_name="X", change_type="EDIT", user_email="a@b.com"))
        db.commit()

    app = FastAPI()
    import routers.sheet_activity_router as SA
    import routers.sheets_router as R
    app.include_router(R.router)
    app.include_router(SA.router)
    who = {"id": ids["Boss"]}

    def _db():
        db = SessionLocal()
        try:
            yield db
        finally:
            db.close()

    def _user():
        with SessionLocal() as db:
            u = db.query(User).filter(User.id == who["id"]).one()
            u.feature_grants, u.role_assignments
            db.expunge(u)
            return u

    app.dependency_overrides[get_operational_db] = _db
    app.dependency_overrides[get_current_user] = _user
    client = TestClient(app)

    site = Workbook()
    ws = site.active
    ws.title = "Site A"
    ws.append(HEADERS)
    ws.append(["1", "From API", "", "", "", "", "", ""])
    buf = io.BytesIO()
    site.save(buf)
    S.fetch_workbook_bytes = lambda sid, client=None: buf.getvalue()   # inspect and every poll read this
    insp = client.post("/api/sheets/inspect", json={"url": f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit"}).json()
    _assert(insp["tabs"][0]["suggestTracked"] and insp["mapping"]["Topic"]["role"] == "topic", "inspect")
    body = {"url": f"https://docs.google.com/spreadsheets/d/{SHEET_ID}2/edit", "name": "API sheet",
            "tabs": [{"name": "Site A", "tracked": True}], "mapping": insp["mapping"], "memberIds": [ids["Member"]]}
    r = client.post("/api/sheets", json=body)
    _assert(r.status_code == 201 and r.json()["firstSync"]["created"] == 1, f"create {r.text[:200]}")
    sheet_id = r.json()["sheet"]["id"]
    _assert(client.post("/api/sheets", json=body).status_code == 409, "duplicate spreadsheet refused")
    _assert(client.post("/api/sheets", json={**body, "url": body["url"] + "x", "tabs": [{"name": "Site A", "tracked": False}]}).status_code == 400,
            "needs a tracked tab")

    who["id"] = ids["Member"]
    _assert([s["id"] for s in client.get("/api/sheets").json()["sheets"]] == [sheet_id], "member sees assigned sheet")
    _assert(client.get(f"/api/sheets/{sheet_id}/requests").json()["total"] == 1, "member reads requests")
    _assert(client.post("/api/sheets/inspect", json={"url": body["url"]}).status_code == 403, "member can't add sheets")
    _assert(client.get("/api/sheet-activity").json()["total"] == 0, "member doesn't see other spreadsheets' raw edits")
    who["id"] = ids["Outsider"]
    _assert(client.get("/api/sheets").json()["sheets"] == [], "unassigned user sees no sheets")
    _assert(client.get(f"/api/sheets/{sheet_id}/overview").status_code == 404, "and can't open one by id")
    who["id"] = ids["NoGrant"]
    _assert(client.get("/api/sheets").status_code == 403, "no Section Access grant -> 403")
    who["id"] = ids["Boss"]
    with SessionLocal() as db:
        all_raw = db.query(SheetActivity).count()
    _assert(all_raw >= 2 and client.get("/api/sheet-activity").json()["total"] == all_raw, "admins see every raw edit")
    r = client.put(f"/api/sheets/people/{ids['Member']}/google-email", json={"googleEmail": "member.personal@gmail.com"})
    _assert(r.status_code == 200, "google email override")
    with SessionLocal() as db:
        _assert(Q.people_directory(db)["member.personal@gmail.com"]["name"] == "Member", "override used for attribution")
    _assert(client.patch(f"/api/sheets/{sheet_id}", json={"isActive": False}).json()["sheet"]["isActive"] is False, "pause")
    _assert(client.get(f"/api/sheets/{sheet_id}/requests.csv").status_code == 200, "csv")
    print("ok  API access and registry")


def main() -> int:
    test_mapping_and_url()
    test_classification()
    test_snapshot_diff_lifecycle()
    test_workbook_parsing_and_fetch()
    test_api_access()
    print("ALL SHEETS CHECKS PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
