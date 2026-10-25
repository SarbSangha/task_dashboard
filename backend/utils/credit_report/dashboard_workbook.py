"""Interactive dashboard workbook: the Usage & Spend page as one Excel file.

The Dashboard sheet copies the page - KPI cards, the filter bar (Tool,
Account, Department, Person, Client, Search), the three charts with their
measure toggles - and the tab sheets copy its tabs (Activity, People, Tools,
Clients, Daily, Monthly, Departments, Compare). Every number is a live
formula over the period's rows, so changing a filter cell recalculates the
cards, redraws the charts and re-ranks every tab, as the page does.

Plain formulas only (SUMIFS / COUNTIFS / INDEX / MATCH, no macros, no
dynamic arrays), so it works in any Excel, Google Sheets and LibreOffice.

The figures follow live_sheet.build_live_sheet and management.PeriodView:
  - a row's ₹ is fixed for the period (costing shares a fee over the whole
    period before any filter), so a filtered total is a SUMIFS over rows;
  - Total spent = row ₹ + unused plan fees + untracked subscriptions, the
    last two only while no filter is on (PeriodView.filtered);
  - Search narrows Activity and People only, as on the page.

Hidden sheets: Data (one row per generation, for Activity), Agg (charged
rows summed per period/day/tool/account/person/client, for every total),
Calc (all the formulas), Lists (dropdown values).
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Callable, Optional
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from models_new import User

from .costing import PAY_PER_USE
from .facts import (
    CHARGED,
    IST_MINUTES,
    NO_CLIENT,
    UNASSIGNED,
    UNASSIGNED_DEPARTMENT,
    UNASSIGNED_USER_ID,
    ReportFilters,
    load_day_groups,
    load_excluded_summary,
    load_groups,
    prepare,
    recent_log,
)
from .live_sheet import TOOL_OUTPUT_CATEGORY, classify_generation
from .management import build_management, period_label
from .model import build_model
from .workbook import INR_FMT, clean_text, link_formula

IST_DELTA = timedelta(minutes=IST_MINUTES)

MAX_ROWS = 50_000          # per period; more than this is too slow to recalculate as live formulas
ACTIVITY_ROWS = 1000       # the page shows 800
PROMPT_CHARS = 300
MATRIX_CELL_LIMIT = 4000   # person x client lookups above this are left out (too slow)

CUR, PREV = "Current", "Previous"
ALL_TOOLS, ALL_ACCOUNTS, ALL_DEPTS, EVERYONE, ALL_CLIENTS = (
    "All tools", "All accounts", "All departments", "Everyone", "All clients")
TOOL_MEASURES = ("By outputs", "By spend", "By credits")      # page TOOL_MODES order
DEPT_MEASURES = ("By spend", "By outputs", "By credits")      # page DEPT_MODES order
PALETTE = ("10B981", "8B5CF6", "38BDF8", "F59E0B", "EC4899", "14B8A6", "F97316", "A3E635")

DASH, ACTIVITY, PEOPLE, TOOLS, CLIENTS, DAILY, MONTHLY, DEPTS, COMPARE = (
    "Dashboard", "Activity", "People", "Tools", "Clients", "Daily", "Monthly", "Departments", "Compare")
TABS = (ACTIVITY, PEOPLE, TOOLS, CLIENTS, DAILY, MONTHLY, DEPTS, COMPARE)
DATA, AGG, CALC, LISTS = "Data", "Agg", "Calc", "Lists"

COUNT_FMT = "#,##0"
CREDITS_FMT = "#,##0.00"
PCT_FMT = "0%"
DATE_FMT = "dd mmm yyyy"
INR2_FMT = '"₹"#,##0.00'

# styles - the page's palette: ink text, muted greys, green as the brand colour
INK, MUTED, FAINT, LINE, PANEL = "111827", "6B7280", "9CA3AF", "E5E7EB", "F9FAFB"
GREEN, GREEN_DARK, GREEN_SOFT = "10B981", "047857", "D1FAE5"
RED, RED_SOFT, AMBER = "DC2626", "FEE2E2", "B45309"
HEAD_BG = "1F2937"
CARD_ACCENTS = (GREEN, "8B5CF6", "F59E0B", "EC4899")


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


F_TITLE = Font(size=22, bold=True, color=INK)
F_SUB = Font(size=10, color=MUTED)
F_LABEL = Font(size=10, color=MUTED)
F_CAPTION = Font(size=9, bold=True, color=MUTED)
F_BOLD = Font(bold=True, color=INK)
F_SECTION = Font(size=12, bold=True, color=INK)
F_KPI = Font(size=22, bold=True, color=INK)
F_FOOT = Font(size=9, color=MUTED)
F_CHANGE = Font(size=9, bold=True, color=MUTED)       # coloured red / green by a rule
F_LINK = Font(color="0563C1", underline="single")
F_TAB_ON = Font(bold=True, color="FFFFFF")
F_TAB_OFF = Font(bold=True, color="374151")
F_HEAD = Font(bold=True, color="374151", size=10)
F_THEAD = Font(bold=True, color="FFFFFF", size=10)
F_HELP = Font(size=9, italic=True, color=FAINT)
F_INPUT = Font(bold=True, color=INK)
FILL_INPUT = _fill("FFF7D6")
FILL_HEAD = _fill("F3F4F6")
FILL_THEAD = _fill(HEAD_BG)
FILL_PANEL = _fill(PANEL)
FILL_TAB_ON = _fill(GREEN)
FILL_TAB_OFF = _fill("F3F4F6")
CARD_FILLS = tuple(_fill(c) for c in ("ECFDF5", "F5F3FF", "FFFBEB", "FDF2F8"))
THIN = Side(style="thin", color="D1D5DB")
HAIR = Side(style="thin", color=LINE)
GAP = Side(style="thick", color="FFFFFF")              # white gap between cards / tabs
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
INPUT_BOX = Border(left=Side(style="thin", color="F59E0B"), right=Side(style="thin", color="F59E0B"),
                   top=Side(style="thin", color="F59E0B"), bottom=Side(style="thin", color="F59E0B"))
UNDERLINE = Border(bottom=Side(style="medium", color=GREEN))
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")
LEFT_MID = Alignment(horizontal="left", vertical="center")
RIGHT_MID = Alignment(horizontal="right", vertical="center")
NAV_ALIGN = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
HEAD_ALIGN = Alignment(vertical="center", wrap_text=True)

# Indian digit grouping (1,97,224), as the page shows numbers
INDIAN = "[>=10000000]##\\,##\\,##\\,##0;[>=100000]##\\,##\\,##0;#,##0"


def _indian(expr: str) -> str:
    return f'TEXT({expr},"{INDIAN}")'


class DashboardTooLarge(ValueError):
    """More rows than an interactive workbook can recalculate."""


# --------------------------------------------------------------------------- #
# Data
# --------------------------------------------------------------------------- #
@dataclass
class Row:
    period: str
    when: Optional[datetime]        # IST
    tool: str
    account: str
    account_id: Optional[int]
    user_id: int
    person: str
    dept: str
    client: str
    task: str
    out_type: str
    what: str
    prompt: str
    credits: float
    cost: float
    status: str
    basis: str
    link: str

    @property
    def day(self) -> Optional[date]:
        return self.when.date() if self.when else None


@dataclass
class DashboardData:
    start: date
    end: date
    prev_start: date
    prev_end: date
    generated_at: datetime
    rows: list                      # current period, newest first
    prev_rows: list
    plan_fee: dict                  # tool -> period fee of its tracked subscriptions
    unused: dict                    # tool -> unused fee (in Total spent when unfiltered)
    untracked: dict                 # tool -> fee of a subscription whose usage isn't captured
    prev_unused: float
    prev_untracked: float
    people: dict                    # person -> department, for every person in the rows
    person_by_id: dict              # user id -> person display name
    account_by_id: dict             # account id -> label

    @property
    def days(self) -> int:
        return (self.end - self.start).days + 1

    @property
    def label(self) -> str:
        return period_label(self.start, self.end)


def _people_meta(db: Session, filters: ReportFilters, model) -> dict:
    """user id -> (name, department), as the live sheet names them."""
    meta = {UNASSIGNED_USER_ID: (UNASSIGNED, UNASSIGNED_DEPARTMENT, "")}
    for u in model.users:
        meta[u.user_id] = (u.label, u.department, u.employee_id or "")
    for uid, uname, uemp, udept in db.query(User.id, User.name, User.employee_id, User.department).all():
        if int(uid) not in meta:
            meta[int(uid)] = (
                filters.display.person_name(uid, uname) or f"User {uid}",
                filters.display.department_name(udept) or UNASSIGNED_DEPARTMENT,
                (uemp or "").strip(),
            )
    return meta


def _load_rows(db: Session, filters: ReportFilters, period: str, meta: dict, base_url: str) -> tuple[list, set]:
    prepare(db, filters)
    raw = recent_log(db, filters, MAX_ROWS + 1)
    if len(raw) > MAX_ROWS:
        raise DashboardTooLarge(
            f"This period has more than {MAX_ROWS:,} generations - too many for the interactive dashboard. "
            "Pick a shorter period, or use the Full Excel workbook.")
    rows, user_ids = [], set()
    for r in raw:
        name, dept, _emp = meta.get(r.user_id, meta[UNASSIGNED_USER_ID])
        user_ids.add(r.user_id)
        out_type, what = classify_generation(r.tool, r.model, r.prompt)
        link = (f"{base_url}/open-output?ref={quote(r.tool)}:{quote(str(r.record_id))}"
                if r.has_output and base_url else "")
        rows.append(Row(
            period=period,
            when=(r.occurred_at + IST_DELTA) if r.occurred_at else None,
            tool=r.tool,
            account=filters.display.account_label(r.account, r.tool) or "—",
            account_id=r.account,
            user_id=r.user_id,
            person=name,
            dept=dept,
            client=r.client or NO_CLIENT,
            task=(r.task or "").strip(),
            out_type=out_type,
            what=what,
            prompt=(r.prompt or "").strip()[:PROMPT_CHARS],
            credits=round(float(r.credits or 0.0), 2),
            cost=round(float(r.rupees or 0.0), 2),
            status=r.charge_status,
            basis=r.basis or "",
            link=link,
        ))
    return rows, user_ids


def collect(db: Session, filters: ReportFilters, *, dashboard_url: str = "",
            generated_at: Optional[datetime] = None) -> DashboardData:
    """Every row of the period and the one before, unfiltered: the workbook
    applies the filters itself. `filters` carries only the dates and the
    test-account setting."""
    prepare(db, filters)
    groups = load_groups(db, filters)
    day_groups = load_day_groups(db, filters)
    mgmt = build_management(db, filters, groups, day_groups=day_groups)
    model = build_model(groups, load_excluded_summary(db, filters))
    meta = _people_meta(db, filters, model)

    base_url = (dashboard_url or "").rstrip("/")
    rows, ids = _load_rows(db, filters, CUR, meta, base_url)
    prev_filters = filters.shifted(mgmt.prev.start, mgmt.prev.end)
    prev_rows, prev_ids = _load_rows(db, prev_filters, PREV, meta, base_url)

    # The Person filter matches on the name, so two people sharing one get
    # their employee id added.
    names = defaultdict(set)
    for uid in ids | prev_ids:
        names[meta.get(uid, meta[UNASSIGNED_USER_ID])[0]].add(uid)
    person_by_id = {}
    for uid in ids | prev_ids:
        name, _dept, emp = meta.get(uid, meta[UNASSIGNED_USER_ID])
        person_by_id[uid] = f"{name} ({emp or uid})" if len(names[name]) > 1 else name
    for r in rows + prev_rows:
        r.person = person_by_id[r.user_id]
    people = {}
    for r in rows + prev_rows:
        people.setdefault(r.person, r.dept)

    cur = mgmt.cur
    plan_fee, unused, untracked = defaultdict(float), defaultdict(float), defaultdict(float)
    for s in cur.subscriptions:
        plan_fee[s.tool] += float(s.period_fee or 0.0)
        unused[s.tool] += float(s.not_used or 0.0)
    for s in cur.other_subscriptions:
        if not s.to_confirm:
            untracked[s.tool] += float(s.period_fee or 0.0)
    account_by_id = {r.account_id: r.account for r in rows + prev_rows if r.account_id is not None}
    return DashboardData(
        start=filters.start, end=filters.end, prev_start=mgmt.prev.start, prev_end=mgmt.prev.end,
        generated_at=generated_at or (datetime.utcnow() + IST_DELTA),
        rows=rows, prev_rows=prev_rows,
        plan_fee=dict(plan_fee), unused={t: v for t, v in unused.items() if v}, untracked=dict(untracked),
        prev_unused=float(mgmt.prev.not_used or 0.0), prev_untracked=float(mgmt.prev.other_total or 0.0),
        people=people, person_by_id=person_by_id, account_by_id=account_by_id,
    )


# --------------------------------------------------------------------------- #
# Writing helpers
# --------------------------------------------------------------------------- #
def _cell(ws, value=None, fmt=None, font=None, fill=None, align=None, border=None, text=False):
    if text and value is not None:
        value = clean_text(value)
    c = WriteOnlyCell(ws, value=value)
    if text and value is not None:
        c.data_type = "s"          # a name such as "=x" must stay text
    if fmt:
        c.number_format = fmt
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def _txt(ws, value):
    """Text for the big data sheets: a bare str is far cheaper than a cell
    object; only text starting with '=' needs one to stay text."""
    if value is None:
        return None
    value = clean_text(value)
    return _cell(ws, value, text=True) if value.startswith("=") else value


class Grid:
    """A small sheet's cells by (row, col), appended in row order at the end
    (write-only sheets only append)."""

    def __init__(self):
        self.cells: dict = {}

    def put(self, row: int, col: int, value=None, **style) -> None:
        self.cells[(row, col)] = (value, style)

    def flush(self, ws) -> None:
        if not self.cells:
            return
        max_row = max(r for r, _c in self.cells)
        max_col = max(c for _r, c in self.cells)
        for r in range(1, max_row + 1):
            line = []
            for c in range(1, max_col + 1):
                spec = self.cells.get((r, c))
                line.append(None if spec is None else _cell(ws, spec[0], **spec[1]))
            ws.append(line)


def _col(n: int) -> str:
    return get_column_letter(n)


def _letters(cell: str) -> str:
    return cell.rstrip("0123456789")


def _name(wb, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _abs(sheet: str, col: str, r1: int, r2: Optional[int] = None) -> str:
    return f"'{sheet}'!${col}${r1}" if r2 is None else f"'{sheet}'!${col}${r1}:${col}${r2}"


def _rank(key_name: str, idx_name: str, k: str, i: str) -> str:
    """1 = largest key; ties keep list order; a negative key is not ranked."""
    return f'=IF({k}<0,"",COUNTIF({key_name},">"&{k})+COUNTIFS({key_name},{k},{idx_name},"<"&{i})+1)'


def _sumifs(measure: str, *extra: str) -> str:
    """SUMIFS over the Agg sheet's current-period rows that pass the filters."""
    crit = ",".join((f'aPeriod,"{CUR}"', "aMatch,1") + extra)
    return f"SUMIFS({measure},{crit})"


def _sum_by(measure: str, field: str, key: str = "name") -> Callable:
    """Column spec: this row's `measure` total for its `key` cell on Agg column `field`."""
    return lambda i, r, c: "=" + _sumifs(measure, f"{field},{c(key)}")


def _filter_test(tool: str, acct: str, dept: str, person: str, client: str) -> str:
    return (f'OR(fTool="{ALL_TOOLS}",{tool}=fTool),OR(fAcct="{ALL_ACCOUNTS}",{acct}=fAcct),'
            f'OR(fDept="{ALL_DEPTS}",{dept}=fDept),OR(fPerson="{EVERYONE}",{person}=fPerson),'
            f'OR(fClient="{ALL_CLIENTS}",{client}=fClient)')


def _link(sheet: str, label: str) -> str:
    return link_formula(f"#'{sheet}'!A1", label)


FILTER_SUMMARY = ('="Filters: "&fTool&" · "&fAcct&" · "&fDept&" · "&fPerson&" · "&fClient'
                  '&IF(fSearch="",""," · Search: "&fSearch)&"   (change them on the Dashboard)"')


@dataclass
class Block:
    letters: dict
    first: int
    last: int

    def col(self, key: str) -> int:
        from openpyxl.utils import column_index_from_string
        return column_index_from_string(self.letters[key])


class Calc:
    """The hidden Calc sheet: KPI scalars in A:B, then blocks left to right.
    Every block's rows start at FIRST, so a matrix lines up row for row with
    the block it belongs to."""

    FIRST = 3

    def __init__(self, wb):
        self.wb = wb
        self.grid = Grid()
        self.next_col = 4
        self.scalar_row = self.FIRST
        self.grid.put(1, 1, "KPIs", font=F_BOLD)

    def scalar(self, name: str, label: str, value, fmt=None) -> None:
        self.grid.put(self.scalar_row, 1, label, text=True)
        self.grid.put(self.scalar_row, 2, value, fmt=fmt)
        _name(self.wb, name, _abs(CALC, "B", self.scalar_row))
        self.scalar_row += 1

    def block(self, prefix: str, title: str, n: int, columns: list) -> Block:
        """columns: (key, header, spec, fmt). spec is a list of static
        values, or fn(i, r, cell) -> formula, where cell(key) is that
        column's cell on row r. Each column gets the name prefix_key."""
        n = max(n, 1)
        first, last = self.FIRST, self.FIRST + n - 1
        start = self.next_col
        letters = {key: _col(start + j) for j, (key, *_rest) in enumerate(columns)}
        self.next_col = start + len(columns) + 1
        self.grid.put(1, start, title, font=F_BOLD)
        for j, (key, header, spec, fmt) in enumerate(columns):
            ci = start + j
            self.grid.put(2, ci, header, font=F_HEAD)
            for i in range(n):
                r = first + i
                if callable(spec):
                    self.grid.put(r, ci, spec(i, r, lambda k, rr=r: f"{letters[k]}{rr}"), fmt=fmt)
                else:
                    value = spec[i] if i < len(spec) else None
                    self.grid.put(r, ci, value, fmt=fmt, text=isinstance(value, str))
            _name(self.wb, f"{prefix}_{key}", _abs(CALC, letters[key], first, last))
        return Block(letters, first, last)

    def matrix(self, title: str, n: int, headers: list, fn: Callable) -> dict:
        """One column per header, rows aligned with an n-row block;
        fn(header_cell) -> formula. Returns header -> column letter."""
        n = max(n, 1)
        start = self.next_col
        self.grid.put(1, start, title, font=F_BOLD)
        letters = {}
        for j, h in enumerate(headers):
            ci = start + j
            letters[h] = _col(ci)
            self.grid.put(2, ci, h, font=F_HEAD, text=True)
            for i in range(n):
                self.grid.put(self.FIRST + i, ci, fn(f"{_col(ci)}$2"))
        self.next_col = start + max(len(headers), 1) + 1
        return letters


# --------------------------------------------------------------------------- #
# Hidden data sheets
# --------------------------------------------------------------------------- #
def _write_data(wb, ws, rows: list) -> None:
    """One row per current-period generation, newest first, for Activity.
    Row 2 is a sentinel (Seq 0), so MATCH(k-1, Seq, 1)+1 is the k-th match."""
    headers = ("Date", "Time", "Tool", "Account", "Person", "Department", "Client", "Output", "What",
               "Prompt", "Credits", "Cost", "Status", "Link", "Search", "Match", "Seq",
               "Link person", "Link tool", "Link client", "Link date")
    ws.append([_cell(ws, h, font=F_HEAD) for h in headers])
    ws.append([None] * 15 + [0, 0, None, None, None, None])
    r = 2
    for row in rows:
        r += 1
        hay = " ".join((row.person, row.dept, row.tool, row.out_type, row.what, row.prompt, row.client,
                        row.task)).lower()
        test = _filter_test(f"C{r}", f"D{r}", f"F{r}", f"E{r}", f"G{r}")
        ws.append([
            _cell(ws, row.when.date() if row.when else None, fmt=DATE_FMT),
            row.when.strftime("%H:%M") if row.when else "—",
            _txt(ws, row.tool),
            _txt(ws, row.account),
            _txt(ws, row.person),
            _txt(ws, row.dept),
            _txt(ws, row.client),
            _txt(ws, row.out_type),
            _txt(ws, row.what),
            _txt(ws, row.prompt or "—"),
            row.credits,
            row.cost,
            _txt(ws, row.status),
            _txt(ws, row.link) if row.link else None,
            _txt(ws, hay[:2000]),
            f'=IF(AND({test},OR(fSearch="",ISNUMBER(SEARCH(fSearch,O{r})))),1,0)',
            f"=Q{r - 1}+P{r}",
            # interlinks find an entity's first filtered row here (see _linked)
            f'=IF(P{r}=1,E{r},"")', f'=IF(P{r}=1,C{r},"")', f'=IF(P{r}=1,G{r},"")', f'=IF(P{r}=1,A{r},"")',
        ])
    for key, col in (("dDate", "A"), ("dTime", "B"), ("dTool", "C"), ("dAcct", "D"), ("dPerson", "E"),
                     ("dDept", "F"), ("dClient", "G"), ("dOut", "H"), ("dWhat", "I"), ("dPrompt", "J"),
                     ("dCred", "K"), ("dCost", "L"), ("dStatus", "M"), ("dLink", "N"), ("dSeq", "Q"),
                     ("dLnkPerson", "R"), ("dLnkTool", "S"), ("dLnkClient", "T"), ("dLnkDate", "U")):
        _name(wb, key, _abs(DATA, col, 2, r))


def _write_agg(wb, ws, rows: list) -> None:
    """Charged rows of both periods summed per (period, person, day, tool,
    account, client), sorted so each (period, person, day) is contiguous.
    First marks the first matching row of one, which makes people-per-day
    and active-days-per-person plain SUMIFS."""
    sums = defaultdict(lambda: [0, 0.0, 0.0])
    for r in rows:
        if r.status != CHARGED or r.day is None:
            continue
        s = sums[(r.period, r.person, r.day, r.tool, r.account, r.client, r.dept, r.basis)]
        s[0] += 1
        s[1] += r.credits
        s[2] += r.cost
    keys = sorted(sums, key=lambda k: (k[0], k[1].lower(), k[1], k[2], k[3], k[4], k[5]))
    headers = ("Period", "Date", "Month", "Tool", "Account", "Person", "Department", "Client", "Basis",
               "Outputs", "Credits", "Cost", "Match", "First")
    ws.append([_cell(ws, h, font=F_HEAD) for h in headers])
    r, group_start, prev_group = 1, 2, None
    for k in keys:
        r += 1
        period, person, day, tool, account, client, dept, basis = k
        if (period, person, day) != prev_group:
            group_start, prev_group = r, (period, person, day)
        outs, credits, cost = sums[k]
        first = f"=M{r}" if r == group_start else f"=IF(AND(M{r}=1,SUM(M{group_start}:M{r - 1})=0),1,0)"
        test = _filter_test(f"D{r}", f"E{r}", f"G{r}", f"F{r}", f"H{r}")
        ws.append([
            period,
            _cell(ws, day, fmt=DATE_FMT),
            day.strftime("%Y-%m"),
            _txt(ws, tool),
            _txt(ws, account),
            _txt(ws, person),
            _txt(ws, dept),
            _txt(ws, client),
            _txt(ws, basis),
            outs,
            round(credits, 2),
            round(cost, 2),
            f"=IF(AND({test}),1,0)",
            first,
        ])
    if r == 1:          # no rows: keep every range valid
        r = 2
        ws.append(["", None, "", "", "", "", "", "", "", 0, 0, 0, 0, 0])
    for key, col in (("aPeriod", "A"), ("aDate", "B"), ("aMonth", "C"), ("aTool", "D"), ("aAcct", "E"),
                     ("aPerson", "F"), ("aDept", "G"), ("aClient", "H"), ("aBasis", "I"), ("aOut", "J"),
                     ("aCred", "K"), ("aCost", "L"), ("aMatch", "M"), ("aFirst", "N")):
        _name(wb, key, _abs(AGG, col, 2, r))


# --------------------------------------------------------------------------- #
# Calc: every formula behind the dashboard
# --------------------------------------------------------------------------- #
@dataclass
class Layout:
    tools: list
    people: list            # (name, department) sorted by department: each department's people are contiguous
    clients: list
    depts: list
    accounts: list
    days: list
    months: list
    blocks: dict = field(default_factory=dict)


def _layout(data: DashboardData) -> Layout:
    rows = data.rows + data.prev_rows
    tools = _ranked_by_cost(data.rows, lambda r: r.tool)
    for t in _ranked_by_cost(data.prev_rows, lambda r: r.tool) + sorted(data.plan_fee) + sorted(data.untracked):
        if t not in tools:
            tools.append(t)
    clients = _ranked_by_cost(data.rows, lambda r: r.client)
    for c in _ranked_by_cost(data.prev_rows, lambda r: r.client):
        if c not in clients:
            clients.append(c)
    people = sorted(data.people.items(), key=lambda kv: (kv[1].lower(), kv[0].lower()))
    days = [data.start + timedelta(days=i) for i in range(data.days)]
    months = sorted({d.strftime("%Y-%m") for d in days})
    return Layout(
        tools=tools or ["—"], people=people or [("—", "—")], clients=clients or ["—"],
        depts=sorted({d for _p, d in people}, key=str.lower) or ["—"],
        accounts=sorted({r.account for r in rows}, key=str.lower), days=days, months=months,
    )


def _ranked_by_cost(rows: list, key) -> list:
    total = defaultdict(float)
    for r in rows:
        total[key(r)] += r.cost
    return sorted(total, key=lambda k: (-total[k], str(k).lower()))


def _build_calc(wb, data: DashboardData, lay: Layout) -> Calc:
    calc = Calc(wb)
    filtered = (f'IF(OR(fTool<>"{ALL_TOOLS}",fAcct<>"{ALL_ACCOUNTS}",fDept<>"{ALL_DEPTS}",'
                f'fPerson<>"{EVERYONE}",fClient<>"{ALL_CLIENTS}"),1,0)')
    ppu = f'aBasis,"{PAY_PER_USE}"'
    calc.scalar("kFlt", "Any filter on (plan fees then left out)", f"={filtered}")
    calc.scalar("kItem", "Usage spend (rows)", "=" + _sumifs("aCost"), INR_FMT)
    calc.scalar("kPPU", "Pay-per-use", "=" + _sumifs("aCost", ppu), INR_FMT)
    calc.scalar("kUnused", "Unused plan fees", round(sum(data.unused.values()), 2), INR_FMT)
    calc.scalar("kUntracked", "Untracked subscriptions", round(sum(data.untracked.values()), 2), INR_FMT)
    calc.scalar("kPlans", "Plans", "=kItem-kPPU+IF(kFlt=1,0,kUnused)", INR_FMT)
    calc.scalar("kTotal", "Total spent", "=kPPU+kPlans+IF(kFlt=1,0,kUntracked)", INR_FMT)
    calc.scalar("kPrevItem", "Previous usage", f'=SUMIFS(aCost,aPeriod,"{PREV}",aMatch,1)', INR_FMT)
    calc.scalar("kPrevUnused", "Previous unused plan fees", round(data.prev_unused, 2), INR_FMT)
    calc.scalar("kPrevUntracked", "Previous untracked", round(data.prev_untracked, 2), INR_FMT)
    calc.scalar("kPrevTotal", "Previous total", "=kPrevItem+IF(kFlt=1,0,kPrevUnused+kPrevUntracked)", INR_FMT)
    calc.scalar("kOutputs", "Outputs", "=" + _sumifs("aOut"), COUNT_FMT)
    calc.scalar("kCredits", "Credits", "=" + _sumifs("aCred"), CREDITS_FMT)
    calc.scalar("kPrevOutputs", "Previous outputs", f'=SUMIFS(aOut,aPeriod,"{PREV}",aMatch,1)', COUNT_FMT)
    calc.scalar("kPrevCredits", "Previous credits", f'=SUMIFS(aCred,aPeriod,"{PREV}",aMatch,1)', CREDITS_FMT)
    calc.scalar("kDays", "Days in period", data.days)
    calc.scalar("kMonths", "Months (days / 30, at least 1)", "=MAX(kDays/30,1)")
    calc.scalar("kPeople", "People", '=COUNTIFS(p_out,">0",p_named,1)')
    calc.scalar("kTools", "Tools", '=COUNTIF(t_out,">0")')
    # management.change_text: '▲ 12%', '▼ 5%', '▲ <1%', '— no change', 'new', '—'
    calc.scalar("kChange", "Change vs previous", (
        '=IF(kPrevTotal=0,IF(kTotal>0,"new","—"),'
        'IF(ROUND((kTotal-kPrevTotal)/kPrevTotal*100,0)=0,'
        'IF(ROUND(kTotal-kPrevTotal,2)=0,"— no change",IF(kTotal>kPrevTotal,"▲","▼")&" <1%"),'
        'IF(kTotal>kPrevTotal,"▲ ","▼ ")&ABS(ROUND((kTotal-kPrevTotal)/kPrevTotal*100,0))&"%"))'))
    calc.scalar("mToolIdx", "Usage-by-tool measure",
                f'=IF(mTool="{TOOL_MEASURES[1]}",2,IF(mTool="{TOOL_MEASURES[2]}",3,1))')
    calc.scalar("mDeptIdx", "Department measure",
                f'=IF(mDept="{DEPT_MEASURES[1]}",2,IF(mDept="{DEPT_MEASURES[2]}",3,1))')
    calc.scalar("kActCount", "Activity rows matching", "=MAX(dSeq)")

    people, tools, clients, depts = lay.people, lay.tools, lay.clients, lay.depts
    np_, nt, nc = len(people), len(tools), len(clients)

    def p_column(col: str) -> str:
        """A People-aligned matrix column, all its rows."""
        return f"${col}${Calc.FIRST}:${col}${Calc.FIRST + np_ - 1}"

    def cross(row_field: str, row_name: str, col_field: str) -> Callable:
        """Matrix cell: outputs for this row's entity (INDEX into its block) x the column header."""
        return lambda h: "=" + _sumifs("aOut", f"{row_field},INDEX({row_name},ROW()-2)", f"{col_field},{h}")

    # Matrices first: the blocks below read their columns.
    pt = calc.matrix("Person x tool (outputs)", np_, tools, cross("aPerson", "p_name", "aTool"))
    pm = calc.matrix("Person x month (outputs)", np_, lay.months, cross("aPerson", "p_name", "aMonth"))
    pc = (calc.matrix("Person x client (outputs)", np_, clients, cross("aPerson", "p_name", "aClient"))
          if np_ * nc <= MATRIX_CELL_LIMIT else {})
    ct = calc.matrix("Client x tool (outputs)", nc, tools, cross("aClient", "c_name", "aTool"))
    dt = calc.matrix("Department x tool (outputs)", len(depts), tools, cross("aDept", "d_name", "aTool"))
    dyt = calc.matrix("Day x tool (outputs)", len(lay.days), tools, cross("aDate", "dl_date", "aTool"))
    mt = calc.matrix("Month x tool (outputs)", len(lay.months), tools, cross("aMonth", "ml_month", "aTool"))

    def top_of(letters: dict, r: int, count_cell: str) -> str:
        """The header of the largest cell in row r of a matrix."""
        if not letters:
            return '="—"'
        cols = list(letters.values())
        head, line = f"${cols[0]}$2:${cols[-1]}$2", f"${cols[0]}{r}:${cols[-1]}{r}"
        return f'=IF({count_cell}<=0,"—",INDEX({head},MATCH(MAX({line}),{line},0)))'

    def tools_text(r: int) -> str:
        parts = [f'IF({c}{r}>0,{c}$2&" ("&{c}{r}&"), ","")' for c in pt.values()]
        return "=" + "&".join(parts)

    def column_count(letters: dict, names: list, i: int) -> str:
        col = letters.get(names[i]) if i < len(names) else None
        return f'=COUNTIF({p_column(col)},">0")' if col else '="—"'

    def column_top(letters: dict, names: list, i: int) -> str:
        col = letters.get(names[i]) if i < len(names) else None
        if not col:
            return '="—"'
        rng = p_column(col)
        return f'=IF(MAX({rng})<=0,"—",INDEX(p_name,MATCH(MAX({rng}),{rng},0)))'

    # ---- people ----
    lay.blocks["p"] = calc.block("p", "People", np_, [
        ("name", "Person", [p for p, _d in people], None),
        ("dept", "Department", [d for _p, d in people], None),
        ("named", "Named", [0 if p == UNASSIGNED else 1 for p, _d in people], None),
        ("out", "Outputs", _sum_by("aOut", "aPerson"), COUNT_FMT),
        ("cred", "Credits", _sum_by("aCred", "aPerson"), CREDITS_FMT),
        ("spend", "Spent", _sum_by("aCost", "aPerson"), INR_FMT),
        ("days", "Active days", lambda i, r, c: f'=SUMIFS(aFirst,aPeriod,"{CUR}",aPerson,{c("name")})', COUNT_FMT),
        ("toolsh", "Tools (helper)", lambda i, r, c: tools_text(r), None),
        ("tools", "Tools used",
         lambda i, r, c: f'=IF(LEN({c("toolsh")})=0,"—",LEFT({c("toolsh")},LEN({c("toolsh")})-2))', None),
        ("toptool", "Top tool", lambda i, r, c: top_of(pt, r, c("out")), None),
        ("topclient", "Top client", lambda i, r, c: top_of(pc, r, c("out")), None),
        ("perout", "Per output", lambda i, r, c: f'=IF({c("out")}>0,{c("spend")}/{c("out")},0)', INR2_FMT),
        ("share", "Share",
         lambda i, r, c: f'=IF(kItem>0,{c("spend")}/kItem,IF(kOutputs>0,{c("out")}/kOutputs,0))', PCT_FMT),
        ("hit", "Search hit", lambda i, r, c: (
            f'=IF(fSearch="",1,IF(ISNUMBER(SEARCH(fSearch,{c("name")}&" "&{c("dept")}&" "&{c("tools")})),1,0))'),
         None),
        ("idx", "#", list(range(1, np_ + 1)), None),
        ("key", "Rank key",
         lambda i, r, c: f'=IF(AND({c("out")}>0,{c("hit")}=1),ROUND({c("spend")},2),-1)', None),
        ("rank", "Rank", lambda i, r, c: _rank("p_key", "p_idx", c("key"), c("idx")), None),
    ])

    # ---- tools ----
    fixed_fee = [round(data.plan_fee.get(t, 0.0) + data.untracked.get(t, 0.0), 2) for t in tools]
    fixed_unused = [round(data.unused.get(t, 0.0) + data.untracked.get(t, 0.0), 2) for t in tools]
    categories = [TOOL_OUTPUT_CATEGORY.get(t, "Subscription" if t in data.untracked else "AI Output") for t in tools]
    lay.blocks["t"] = calc.block("t", "Tools", nt, [
        ("name", "Tool", list(tools), None),
        ("cat", "Category", categories, None),
        ("out", "Outputs", _sum_by("aOut", "aTool"), COUNT_FMT),
        ("cred", "Credits", _sum_by("aCred", "aTool"), CREDITS_FMT),
        ("usage", "Usage", _sum_by("aCost", "aTool"), INR_FMT),
        ("fee", "Plan fee", fixed_fee, INR_FMT),
        ("unusedc", "Unused (no filter)", fixed_unused, INR_FMT),
        ("unused", "Unused", lambda i, r, c: f'=IF(kFlt=1,0,{c("unusedc")})', INR_FMT),
        ("total", "Total", lambda i, r, c: f'={c("usage")}+{c("unused")}', INR_FMT),
        ("people", "People", lambda i, r, c: column_count(pt, tools, i), COUNT_FMT),
        ("top", "Top user", lambda i, r, c: column_top(pt, tools, i), None),
        ("perout", "Per output", lambda i, r, c: f'=IF({c("out")}>0,{c("usage")}/{c("out")},0)', INR2_FMT),
        ("perday", "Outputs / day", lambda i, r, c: f'={c("out")}/kDays', "0.0"),
        ("share", "Share", lambda i, r, c: f'=IF(kTotal>0,{c("total")}/kTotal,0)', PCT_FMT),
        ("measure", "Measure", lambda i, r, c: f'=CHOOSE(mToolIdx,{c("out")},{c("total")},{c("cred")})', None),
        ("idx", "#", list(range(1, nt + 1)), None),
        ("key", "Rank key",
         lambda i, r, c: f'=IF(OR({c("out")}>0,{c("total")}>0),ROUND({c("total")},2),-1)', None),
        ("rank", "Rank", lambda i, r, c: _rank("t_key", "t_idx", c("key"), c("idx")), None),
        ("mkey", "Measure key", lambda i, r, c: f'=IF({c("measure")}>0,ROUND({c("measure")},4),-1)', None),
        ("mrank", "Measure rank", lambda i, r, c: _rank("t_mkey", "t_idx", c("mkey"), c("idx")), None),
    ])

    # ---- clients ----
    lay.blocks["c"] = calc.block("c", "Clients", nc, [
        ("name", "Client / Project", list(clients), None),
        ("out", "Outputs", _sum_by("aOut", "aClient"), COUNT_FMT),
        ("cred", "Credits", _sum_by("aCred", "aClient"), CREDITS_FMT),
        ("spend", "Spent", _sum_by("aCost", "aClient"), INR_FMT),
        ("people", "People", lambda i, r, c: column_count(pc, clients, i), COUNT_FMT),
        ("top", "Top user", lambda i, r, c: column_top(pc, clients, i), None),
        ("toptool", "Top tool", lambda i, r, c: top_of(ct, r, c("out")), None),
        ("perout", "Per output", lambda i, r, c: f'=IF({c("out")}>0,{c("spend")}/{c("out")},0)', INR2_FMT),
        ("share", "Share", lambda i, r, c: f'=IF(kItem>0,{c("spend")}/kItem,0)', PCT_FMT),
        ("idx", "#", list(range(1, nc + 1)), None),
        ("key", "Rank key", lambda i, r, c: f'=IF({c("out")}>0,ROUND({c("spend")},2),-1)', None),
        ("rank", "Rank", lambda i, r, c: _rank("c_key", "c_idx", c("key"), c("idx")), None),
    ])

    # ---- departments: a department's people are rows a..b of the People block ----
    pb = lay.blocks["p"]
    span = {}
    for i, (_p, d) in enumerate(people):
        a = span.get(d, (pb.first + i, None))[0]
        span[d] = (a, pb.first + i)

    def dept_top(i, r, c):
        if i >= len(depts) or depts[i] not in span:
            return '="—"'
        a, b = span[depts[i]]
        s, n = pb.letters["spend"], pb.letters["name"]
        rng = f"${s}${a}:${s}${b}"
        return f'=IF({c("people")}=0,"—",INDEX(${n}${a}:${n}${b},MATCH(MAX({rng}),{rng},0)))'

    lay.blocks["d"] = calc.block("d", "Departments", len(depts), [
        ("name", "Department", list(depts), None),
        ("out", "Outputs", _sum_by("aOut", "aDept"), COUNT_FMT),
        ("cred", "Credits", _sum_by("aCred", "aDept"), CREDITS_FMT),
        ("spend", "Spent", _sum_by("aCost", "aDept"), INR_FMT),
        ("people", "People", lambda i, r, c: f'=COUNTIFS(p_dept,{c("name")},p_out,">0")', COUNT_FMT),
        ("top", "Top user", dept_top, None),
        ("toptool", "Top tool", lambda i, r, c: top_of(dt, r, c("out")), None),
        ("perperson", "Per person",
         lambda i, r, c: f'=IF({c("people")}>0,{c("spend")}/{c("people")},0)', INR_FMT),
        ("share", "Share", lambda i, r, c: f'=IF(kItem>0,{c("spend")}/kItem,0)', PCT_FMT),
        ("measure", "Measure", lambda i, r, c: f'=CHOOSE(mDeptIdx,{c("spend")},{c("out")},{c("cred")})', None),
        ("idx", "#", list(range(1, len(depts) + 1)), None),
        ("key", "Rank key", lambda i, r, c: f'=IF({c("out")}>0,ROUND({c("spend")},2),-1)', None),
        ("rank", "Rank", lambda i, r, c: _rank("d_key", "d_idx", c("key"), c("idx")), None),
        ("mkey", "Measure key", lambda i, r, c: f'=IF({c("measure")}>0,ROUND({c("measure")},4),-1)', None),
        ("mrank", "Measure rank", lambda i, r, c: _rank("d_mkey", "d_idx", c("mkey"), c("idx")), None),
    ])

    # ---- days (oldest first: the trend chart's order) ----
    lay.blocks["dl"] = calc.block("dl", "Daily", len(lay.days), [
        ("date", "Date", list(lay.days), DATE_FMT),
        ("label", "Label", [f"{d.day} {d.strftime('%b')}" for d in lay.days], None),
        ("month", "Month", [d.strftime("%Y-%m") for d in lay.days], None),
        ("out", "Outputs", _sum_by("aOut", "aDate", "date"), COUNT_FMT),
        ("people", "People", lambda i, r, c: f'=SUMIFS(aFirst,aPeriod,"{CUR}",aDate,{c("date")})', COUNT_FMT),
        ("cred", "Credits", _sum_by("aCred", "aDate", "date"), CREDITS_FMT),
        ("spend", "Spent", _sum_by("aCost", "aDate", "date"), INR_FMT),
        ("toptool", "Top tool", lambda i, r, c: top_of(dyt, r, c("out")), None),
        ("perout", "Per output", lambda i, r, c: f'=IF({c("out")}>0,{c("spend")}/{c("out")},0)', INR2_FMT),
        ("idx", "#", list(range(1, len(lay.days) + 1)), None),
        ("key", "Rank key", lambda i, r, c: f'=IF({c("out")}>0,{c("date")},-1)', None),
        ("rank", "Rank", lambda i, r, c: _rank("dl_key", "dl_idx", c("key"), c("idx")), None),
    ])

    # ---- months ----
    lay.blocks["ml"] = calc.block("ml", "Monthly", len(lay.months), [
        ("month", "Month", list(lay.months), None),
        ("label", "Label", [datetime.strptime(m, "%Y-%m").strftime("%b %Y") for m in lay.months], None),
        ("days", "Active days", lambda i, r, c: f'=COUNTIFS(dl_month,{c("month")},dl_out,">0")', COUNT_FMT),
        ("people", "People", lambda i, r, c: column_count(pm, lay.months, i), COUNT_FMT),
        ("out", "Outputs", _sum_by("aOut", "aMonth", "month"), COUNT_FMT),
        ("cred", "Credits", _sum_by("aCred", "aMonth", "month"), CREDITS_FMT),
        ("spend", "Spent", _sum_by("aCost", "aMonth", "month"), INR_FMT),
        ("toptool", "Top tool", lambda i, r, c: top_of(mt, r, c("out")), None),
        ("perday", "Per active day", lambda i, r, c: f'=IF({c("days")}>0,{c("spend")}/{c("days")},0)', INR_FMT),
        ("idx", "#", list(range(1, len(lay.months) + 1)), None),
        ("key", "Rank key", lambda i, r, c: f'=IF({c("out")}>0,{c("idx")},-1)', None),
        ("rank", "Rank", lambda i, r, c: _rank("ml_key", "ml_idx", c("key"), c("idx")), None),
    ])

    # ---- chart feeds: top 5 tools + Others, top 6 departments ----
    def donut_value(i, r, c):
        if i < 5:
            return f'=IF({c("idx")}="",0,INDEX(t_measure,{c("idx")}))'
        col = _letters(c("value"))
        return f'=MAX(0,SUMIF(t_measure,">0")-SUM({col}{r - 5}:{col}{r - 1}))'

    lay.blocks["dn"] = calc.block("dn", "Usage by tool (chart)", 6, [
        ("idx", "Row", lambda i, r, c: '=""' if i == 5 else f'=IFERROR(MATCH({i + 1},t_mrank,0),"")', None),
        ("name", "Tool", lambda i, r, c: (f'=IF({c("value")}>0,"Others","")' if i == 5
                                           else f'=IF({c("idx")}="","",INDEX(t_name,{c("idx")}))'), None),
        ("value", "Value", donut_value, None),
        ("pct", "Share", lambda i, r, c: f'=IF(SUM(dn_value)>0,{c("value")}/SUM(dn_value),0)', PCT_FMT),
    ])
    lay.blocks["dc"] = calc.block("dc", "Departments (chart)", 6, [
        ("idx", "Row", lambda i, r, c: f'=IFERROR(MATCH({i + 1},d_mrank,0),"")', None),
        ("name", "Department", lambda i, r, c: f'=IF({c("idx")}="","",INDEX(d_name,{c("idx")}))', None),
        ("value", "Value", lambda i, r, c: f'=IF({c("idx")}="",0,ROUND(INDEX(d_measure,{c("idx")}),0))', "#,##0"),
    ])
    return calc


# --------------------------------------------------------------------------- #
# Visible sheets
# --------------------------------------------------------------------------- #
TAB_HEAD = 5
HEADER_FILTER_NOTE = "Use ▼ on the headers to filter or sort; after changing Dashboard filters, click Data › Reapply."
CHANGE_RULES = (  # the change text: spend up is bad (red), down is good (green)
    ('LEFT({c},1)="▲"', RED), ('LEFT({c},1)="▼"', GREEN_DARK))


def _nav(grid: Grid, row: int, current: str) -> None:
    """A tab bar: the current sheet filled green, the others grey; white
    borders between them read as separate tabs."""
    grid.put(row, 2, "Go to", font=F_CAPTION, align=LEFT_MID)
    for j, name in enumerate((DASH,) + TABS):
        border = Border(left=GAP, right=GAP)
        if name == current:
            grid.put(row, 3 + j, name, font=F_TAB_ON, fill=FILL_TAB_ON, align=NAV_ALIGN, border=border, text=True)
        else:
            grid.put(row, 3 + j, _link(name, name), font=F_TAB_OFF, fill=FILL_TAB_OFF, align=NAV_ALIGN,
                     border=border)


def _sheet_setup(ws, *, tab_color: str = "9CA3AF", landscape: bool = True) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab_color
    if landscape:
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[3].height = 22


def _sheet_head(grid: Grid, title: str, note) -> None:
    grid.put(1, 2, title, font=F_TITLE, align=LEFT_MID, text=True)
    grid.put(2, 2, FILTER_SUMMARY, font=F_SUB)
    _nav(grid, 3, title)
    grid.put(4, 2, note, font=F_HELP, text=not str(note).startswith("="))


def _widths(ws, widths: dict) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _table_head(grid: Grid, ws, headers: list) -> None:
    """A dark header row; number columns' headers right-aligned over their figures."""
    ws.row_dimensions[TAB_HEAD].height = 26
    for j, (header, numeric) in enumerate(headers):
        align = Alignment(horizontal="right" if numeric else "left", vertical="center", wrap_text=True)
        grid.put(TAB_HEAD, 2 + j, header, font=F_THEAD, fill=FILL_THEAD, align=align, text=True)


def _table_rules(ws, last_col: int, n_rows: int) -> None:
    """Row styling that follows the data: a hairline under and zebra on rows
    that hold an entry, nothing on the unused rows below; top 3 ranks marked."""
    first, last = TAB_HEAD + 1, TAB_HEAD + max(n_rows, 1)
    body = f"B{first}:{_col(last_col)}{last}"
    ws.conditional_formatting.add(body, FormulaRule(
        formula=[f'AND($A{first}<>"",MOD(ROW(),2)=0)'], fill=FILL_PANEL, border=Border(bottom=HAIR)))
    ws.conditional_formatting.add(body, FormulaRule(formula=[f'$A{first}<>""'], border=Border(bottom=HAIR)))
    ws.conditional_formatting.add(f"B{first}:B{last}", FormulaRule(
        formula=[f"AND(ISNUMBER($B{first}),$B{first}<=3)"], fill=_fill(GREEN_SOFT),
        font=Font(bold=True, color=GREEN_DARK)))


def _column_rules(ws, col: str, style: Optional[str], n_rows: int) -> None:
    first, last = TAB_HEAD + 1, TAB_HEAD + max(n_rows, 1)
    rng = f"{col}{first}:{col}{last}"
    if style in ("bar", "share"):
        ws.conditional_formatting.add(rng, DataBarRule(
            start_type="num", start_value=0, end_type="max", color=GREEN if style == "bar" else "8B5CF6"))
    elif style == "scale":
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color="6EE7B7"))
    elif style == "status":
        for formula, color in ((f'{col}{first}="{CHARGED}"', GREEN_DARK), (f'{col}{first}="Pending"', AMBER),
                               (f'AND({col}{first}<>"",{col}{first}<>"{CHARGED}",{col}{first}<>"Pending")', RED)):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], font=Font(bold=True, color=color)))


def _header_filter(grid: Grid, ws, last_col: int, n_rows: int) -> None:
    """Excel's filter dropdowns on the header row. The range starts at the
    hidden key column A, so sorting from a dropdown carries each row's key
    (which entity it shows) with it."""
    grid.put(TAB_HEAD, 1, "Key", font=F_HEAD)
    ws.auto_filter.ref = f"A{TAB_HEAD}:{_col(last_col)}{TAB_HEAD + max(n_rows, 1)}"


def _spec(column: tuple) -> tuple:
    """(header, block column name, number format, width[, style[, link]]) -> 6-tuple.
    style: name (bold), bar / share (data bar), scale (colour scale), status.
    link: a key of the link targets (see _link_targets): the cell jumps there."""
    return column + (None,) * (6 - len(column))


# Interlinks: a name in one tab jumps to the same entity's row in another,
# as clicking it drills in on the page. The target row is looked up live
# (MATCH on the target tab's column), so it stays right however that tab is
# ranked, filtered or sorted; an entity the target doesn't list (filtered
# out, or older than the Activity window) stays plain text.
F_XLINK = Font(color="1D4ED8")
F_XLINK_BOLD = Font(color="1D4ED8", bold=True)
LINKS_NOTE = "Blue names are links: they open that person, tool, client, department or day in its own tab."

ACTIVITY_COLUMNS = (
    ("When", "dDate", DATE_FMT, 13, None, "date"), ("Time", "dTime", None, 7),
    ("Person", "dPerson", None, 22, "name", "person"), ("Tool", "dTool", None, 12, "name", "tool"),
    ("Account", "dAcct", None, 26), ("Output", "dOut", None, 15), ("What", "dWhat", None, 18),
    ("Prompt", "dPrompt", None, 48), ("Client / Project", "dClient", None, 20, None, "client"),
    ("Department", "dDept", None, 16, None, "dept"), ("Credits", "dCred", CREDITS_FMT, 10),
    ("Cost (₹)", "dCost", INR2_FMT, 11, "bar"), ("Status", "dStatus", None, 12, "status"),
)


def _link_targets(lay: Layout, n_activity: int) -> dict:
    """key -> where a link lands. Targets are found WITHOUT reading the
    target tab's own cells (those hold links back, which would be a circular
    reference):
      ("ranked", sheet, calc names, rows): the entity's index in its Calc
          block, looked up in the tab's hidden key column A - which depends
          on Calc only, and moves with its row when the tab is sorted;
      ("activity", Data link column, rows): the entity's first filtered row
          in Data, whose Seq is its Activity row - shown only if within the
          Activity window."""
    return {
        "person": ("ranked", PEOPLE, "p_name", len(lay.people)),
        "tool": ("ranked", TOOLS, "t_name", len(lay.tools)),
        "client": ("ranked", CLIENTS, "c_name", len(lay.clients)),
        "dept": ("ranked", DEPTS, "d_name", len(lay.depts)),
        "date": ("ranked", DAILY, "dl_date", len(lay.days)),
        "act_person": ("activity", "dLnkPerson", n_activity),
        "act_tool": ("activity", "dLnkTool", n_activity),
        "act_client": ("activity", "dLnkClient", n_activity),
        "act_date": ("activity", "dLnkDate", n_activity),
    }


def _linked(value: str, target: tuple, label: Optional[str] = None) -> str:
    """value, as a link to its row in the target table (plain if not there)."""
    shown = label or value
    if target[0] == "ranked":
        _kind, sheet, names, n = target
        keys = f"'{sheet}'!$A${TAB_HEAD + 1}:$A${TAB_HEAD + max(n, 1)}"
        row = f"MATCH(MATCH({value},{names},0),{keys},0)"
    else:
        _kind, helper, n = target
        sheet = ACTIVITY
        seq = f"INDEX(dSeq,MATCH({value},{helper},0))"
        row = f"IF({seq}<={n},{seq},NA())"
    return f"IFERROR(HYPERLINK(\"#'{sheet}'!C\"&({TAB_HEAD}+{row}),{shown}),{shown})"


def _swatch(expr: str) -> str:
    """A legend entry: a colour square before the name (the cell font gives the colour)."""
    return f'"■  "&{expr}'


def _body_cell(grid: Grid, r: int, c: int, spec: tuple, targets: dict) -> None:
    """link is a target key, or (target key, column to look up) when the cell
    shows one thing but opens another (a department opens its top person)."""
    _h, name_ref, fmt, _w, style, link = spec
    value = f"INDEX({name_ref},$A{r})"
    if link:
        key, lookup = link if isinstance(link, tuple) else (link, name_ref)
        font = F_XLINK_BOLD if style == "name" else F_XLINK
        formula = _linked(f"INDEX({lookup},$A{r})", targets[key], label=value)
        grid.put(r, c, f'=IF($A{r}="","",{formula})', fmt=fmt, font=font)
    else:
        grid.put(r, c, f'=IF($A{r}="","",{value})', fmt=fmt, font=F_BOLD if style == "name" else None)


def _write_ranked_tab(ws, title: str, note: str, rank_name: str, n: int, columns: list, targets: dict) -> None:
    """A tab listing a block's rows in rank order: row k shows the entity
    ranked k under the current filters, so the table re-sorts as they change.
    columns: (header, block column name, number format, width[, style[, link]])."""
    columns = [_spec(c) for c in columns]
    grid = Grid()
    _sheet_setup(ws)
    _sheet_head(grid, title, f"{note} {LINKS_NOTE} {HEADER_FILTER_NOTE}")
    _table_head(grid, ws, [("#", True)] + [(h, fmt is not None) for h, _r, fmt, *_x in columns])
    for k in range(1, max(n, 1) + 1):
        r = TAB_HEAD + k
        grid.put(r, 1, f'=IFERROR(MATCH({k},{rank_name},0),"")')
        grid.put(r, 2, f'=IF($A{r}="","",{k})', font=F_LABEL, align=CENTER)
        for j, spec in enumerate(columns):
            _body_cell(grid, r, 3 + j, spec, targets)
    last_col = 2 + len(columns)
    _header_filter(grid, ws, last_col, n)
    _table_rules(ws, last_col, n)
    for j, spec in enumerate(columns):
        _column_rules(ws, _col(3 + j), spec[4], n)
    widths = {"A": 4, "B": 5}
    widths.update({_col(3 + j): spec[3] for j, spec in enumerate(columns)})
    _widths(ws, widths)
    ws.column_dimensions["A"].hidden = True
    ws.freeze_panes = f"C{TAB_HEAD + 1}"
    grid.flush(ws)


def _write_activity(ws, n_rows: int, targets: dict) -> None:
    grid = Grid()
    _sheet_setup(ws)
    _sheet_head(grid, ACTIVITY, (f'="Showing "&MIN(kActCount,{n_rows})&" of "&kActCount&" matching generations, '
                                 f'newest first"&IF(kActCount>{n_rows},". The rest are in the Full Excel workbook.","")'
                                 f'&". {LINKS_NOTE} {HEADER_FILTER_NOTE}"'))
    columns = [_spec(c) for c in ACTIVITY_COLUMNS]
    _table_head(grid, ws, [("S.No", True)] + [(h, fmt is not None and h != "When") for h, _r, fmt, *_x in columns]
                + [("Link", False)])
    link_col = 3 + len(columns)
    for k in range(1, max(n_rows, 1) + 1):
        r = TAB_HEAD + k
        grid.put(r, 1, f'=IF({k}>kActCount,"",MATCH({k - 1},dSeq,1)+1)')
        grid.put(r, 2, f'=IF($A{r}="","",{k})', font=F_LABEL, align=CENTER)
        for j, spec in enumerate(columns):
            _body_cell(grid, r, 3 + j, spec, targets)
        grid.put(r, link_col, (f'=IF($A{r}="","",IF(INDEX(dLink,$A{r})="","",'
                               f'HYPERLINK(INDEX(dLink,$A{r}),"Open ↗")))'), font=F_LINK)
    _header_filter(grid, ws, link_col, n_rows)
    _table_rules(ws, link_col, n_rows)
    for j, spec in enumerate(columns):
        _column_rules(ws, _col(3 + j), spec[4], n_rows)
    widths = {"A": 4, "B": 6}
    widths.update({_col(3 + j): spec[3] for j, spec in enumerate(columns)})
    widths[_col(link_col)] = 9
    _widths(ws, widths)
    ws.column_dimensions["A"].hidden = True
    ws.freeze_panes = f"C{TAB_HEAD + 1}"
    grid.flush(ws)


TOOL_COMPARE = (("Outputs", "t_out", COUNT_FMT, True), ("Credits used", "t_cred", CREDITS_FMT, True),
                ("Total spend", "t_total", INR_FMT, True), ("Usage spend", "t_usage", INR_FMT, True),
                ("Plan fee", "t_fee", INR_FMT, False), ("People using it", "t_people", COUNT_FMT, True),
                ("Cost per output", "t_perout", INR2_FMT, False), ("Outputs / day", "t_perday", "0.0", True),
                ("Share of spend", "t_share", PCT_FMT, True))
PEOPLE_COMPARE = (("Department", "p_dept", None, False), ("Outputs", "p_out", COUNT_FMT, True),
                  ("Credits", "p_cred", CREDITS_FMT, True), ("Spent", "p_spend", INR_FMT, True),
                  ("Active days", "p_days", COUNT_FMT, True), ("Cost per output", "p_perout", INR2_FMT, False),
                  ("Share of spend", "p_share", PCT_FMT, True), ("Top tool", "p_toptool", None, False))
PICKS = 6


def _write_compare(ws, lay: Layout, list_ranges: dict, data: DashboardData) -> None:
    grid = Grid()
    _sheet_setup(ws)
    _sheet_head(grid, COMPARE, "Pick up to 6 tools or people in the yellow cells. Every figure follows the "
                               "Dashboard filters; the highest of each row is highlighted.")
    tool_order = lay.tools
    person_order = [p for p in _ranked_by_cost(data.rows, lambda r: r.person) if p != UNASSIGNED]
    last = _col(2 + PICKS)

    def section(top: int, title: str, label: str, picks: list, metrics: tuple, name_ref: str) -> int:
        for c in range(2, 3 + PICKS):
            grid.put(top, c, title if c == 2 else None, font=F_SECTION, border=UNDERLINE, text=c == 2)
        ws.row_dimensions[top].height = 22
        pick_row = top + 1
        ws.row_dimensions[pick_row].height = 22
        grid.put(pick_row, 2, label, font=F_CAPTION, fill=FILL_HEAD, align=LEFT_MID, text=True)
        for j in range(PICKS):
            grid.put(pick_row, 3 + j, picks[j] if j < len(picks) else None, fill=FILL_INPUT, border=INPUT_BOX,
                     font=F_INPUT, align=CENTER, text=True)
        for m, (metric, ref, fmt, best_high) in enumerate(metrics):
            r = top + 2 + m
            grid.put(r, 2, metric, font=F_LABEL, fill=FILL_HEAD, border=Border(bottom=HAIR), text=True)
            for j in range(PICKS):
                p = f"{_col(3 + j)}${pick_row}"
                grid.put(r, 3 + j, f'=IF({p}="","",IFERROR(INDEX({ref},MATCH({p},{name_ref},0)),""))',
                         fmt=fmt, border=Border(bottom=HAIR), align=RIGHT_MID if fmt else CENTER)
            if best_high:
                ws.conditional_formatting.add(f"C{r}:{last}{r}", FormulaRule(
                    formula=[f"AND(ISNUMBER(C{r}),C{r}=MAX($C{r}:${last}{r}),COUNT($C{r}:${last}{r})>1)"],
                    fill=_fill(GREEN_SOFT), font=Font(bold=True, color=GREEN_DARK)))
        return pick_row

    tool_row = section(6, "Tool comparison", "Tools →", tool_order[:3], TOOL_COMPARE, "t_name")
    person_row = section(8 + len(TOOL_COMPARE), "People comparison", "People →", person_order[:3],
                         PEOPLE_COMPARE, "p_name")

    top = 10 + len(TOOL_COMPARE) + len(PEOPLE_COMPARE)
    for c in range(2, 6):
        grid.put(top, c, "This period vs the previous one" if c == 2 else None, font=F_SECTION, border=UNDERLINE,
                 text=c == 2)
    grid.put(top + 1, 2, f"{period_label(data.start, data.end)} vs {period_label(data.prev_start, data.prev_end)}",
             font=F_HELP, text=True)
    _head = [("Metric", False), ("This period", True), ("Previous", True), ("Change", True)]
    for j, (h, numeric) in enumerate(_head):
        grid.put(top + 2, 2 + j, h, font=F_THEAD, fill=FILL_THEAD,
                 align=RIGHT_MID if numeric else LEFT_MID, text=True)
    for m, (metric, cur, prev, fmt, spend) in enumerate((("Total spent", "kTotal", "kPrevTotal", INR_FMT, True),
                                                         ("Usage spend", "kItem", "kPrevItem", INR_FMT, True),
                                                         ("Outputs", "kOutputs", "kPrevOutputs", COUNT_FMT, False),
                                                         ("Credits", "kCredits", "kPrevCredits", CREDITS_FMT, False))):
        r = top + 3 + m
        grid.put(r, 2, metric, font=F_BOLD, border=Border(bottom=HAIR), text=True)
        grid.put(r, 3, f"={cur}", fmt=fmt, border=Border(bottom=HAIR))
        grid.put(r, 4, f"={prev}", fmt=fmt, border=Border(bottom=HAIR), font=F_LABEL)
        grid.put(r, 5, f'=IF({prev}=0,"",{cur}/{prev}-1)', fmt='"▲ "0%;"▼ "0%;0%', border=Border(bottom=HAIR),
                 font=F_BOLD)
        if spend:     # more spend is red, less is green
            ws.conditional_formatting.add(f"E{r}", CellIsRule(operator="greaterThan", formula=["0"],
                                                              font=Font(bold=True, color=RED)))
            ws.conditional_formatting.add(f"E{r}", CellIsRule(operator="lessThan", formula=["0"],
                                                              font=Font(bold=True, color=GREEN_DARK)))
    _widths(ws, {"A": 2, "B": 20, **{_col(3 + j): 17 for j in range(PICKS)}})
    for row, key in ((tool_row, "tools_only"), (person_row, "people_only")):
        dv = DataValidation(type="list", formula1=list_ranges[key], allow_blank=True, showErrorMessage=True,
                            errorTitle="Pick from the list", error="Choose a value from the dropdown.")
        dv.add(f"C{row}:{last}{row}")
        ws.data_validations.append(dv)
    grid.flush(ws)


def _write_lists(ws, lay: Layout) -> dict:
    """Dropdown values, one column each. Returns key -> range for validation."""
    columns = {
        "tool": [ALL_TOOLS] + list(lay.tools),
        "acct": [ALL_ACCOUNTS] + list(lay.accounts),
        "dept": [ALL_DEPTS] + list(lay.depts),
        "person": [EVERYONE] + sorted({p for p, _d in lay.people}, key=str.lower),
        "client": [ALL_CLIENTS] + list(lay.clients),
        "tool_m": list(TOOL_MEASURES),
        "dept_m": list(DEPT_MEASURES),
        "tools_only": list(lay.tools),
        "people_only": sorted({p for p, _d in lay.people}, key=str.lower),
    }
    ranges = {}
    height = max(len(v) for v in columns.values())
    for j, (key, values) in enumerate(columns.items()):
        ranges[key] = _abs(LISTS, _col(1 + j), 1, max(len(values), 1))
    for i in range(height):
        ws.append([_cell(ws, v[i], text=True) if i < len(v) else None for v in columns.values()])
    return ranges


# Dashboard rows
R_CARD = 5            # accent strip; label, value, two footer lines below it
R_FILTER = 12         # filter inputs (labels on the row above, help below)
R_CHART_HEAD = 15     # chart titles; their totals and measure toggles on the next row
R_CHART = 17          # charts are anchored here
R_LEGEND = 34         # legend tables under the charts


def _write_dashboard(ws, wb, data: DashboardData, lay: Layout, calc_ws, list_ranges: dict, initial: dict,
                     targets: dict) -> None:
    grid = Grid()
    _sheet_setup(ws, tab_color=GREEN)
    grid.put(1, 2, "Usage & Spend", font=F_TITLE, align=LEFT_MID, text=True)
    grid.put(2, 2, "Who generated what, with which tool, for which client, and what it cost", font=F_SUB, text=True)
    grid.put(1, 10, f"{period_label(data.start, data.end)} · {data.days} days", font=F_SECTION, align=LEFT_MID,
             text=True)
    grid.put(2, 10, f"Data as of {data.generated_at.strftime('%d %b %Y, %H:%M')} IST", font=F_SUB, text=True)
    _nav(grid, 3, DASH)

    # KPI cards: an accent strip, label, value, two footer lines. Values are
    # text, so a big number spills across its card instead of "######".
    cards = (
        ("TOTAL SPENT", f'="₹"&{_indian("kTotal")}',
         '=IF(OR(kChange="—",kChange=""),"",kChange&" vs previous period")',
         f'="Usage ₹"&{_indian("kPPU")}&" · Plans ₹"&{_indian("kPlans")}'),
        ("AVERAGE SPEND", f'="₹"&{_indian("kTotal/kDays")}&" / day"',
         f'="₹"&{_indian("kTotal/kMonths")}&" / month · "&kDays&" days"', None),
        ("OUTPUTS GENERATED", f"={_indian('kOutputs')}",
         '=TEXT(kOutputs/kDays,"#,##0.0")&" / day · ₹"&TEXT(IF(kOutputs>0,kItem/kOutputs,0),"#,##0.00")&" each"',
         f'={_indian("kCredits")}&" credits"'),
        ("PEOPLE & TOOLS", '=kPeople&" people · "&kTools&" tools"',
         '="₹"&TEXT(IF(kPeople>0,kItem/kPeople,0),"#,##0.00")&" per person (avg)"', None),
    )
    ws.row_dimensions[R_CARD].height = 5
    ws.row_dimensions[R_CARD + 1].height = 20
    ws.row_dimensions[R_CARD + 2].height = 34
    for k, (label, value, foot1, foot2) in enumerate(cards):
        c0 = 2 + 3 * k
        fill, accent = CARD_FILLS[k], _fill(CARD_ACCENTS[k])
        for r in range(R_CARD, R_CARD + 5):
            for c in range(c0, c0 + 3):
                grid.put(r, c, None, fill=accent if r == R_CARD else fill,
                         border=Border(left=GAP) if c == c0 else None)
        grid.put(R_CARD + 1, c0, label, font=Font(size=9, bold=True, color=CARD_ACCENTS[k]), fill=fill,
                 border=Border(left=GAP), align=LEFT_MID, text=True)
        grid.put(R_CARD + 2, c0, value, font=F_KPI, fill=fill, border=Border(left=GAP), align=LEFT_MID)
        grid.put(R_CARD + 3, c0, foot1, font=F_CHANGE if k == 0 else F_FOOT, fill=fill, border=Border(left=GAP))
        if foot2:
            grid.put(R_CARD + 4, c0, foot2, font=F_FOOT, fill=fill, border=Border(left=GAP))
    for formula, color in CHANGE_RULES:
        ws.conditional_formatting.add(f"B{R_CARD + 3}", FormulaRule(
            formula=[formula.format(c=f"B{R_CARD + 3}")], font=Font(size=9, bold=True, color=color)))

    # Filter panel: labels, yellow dropdowns, help line, on a light panel
    filters = (("TOOL", "fTool", "tool", initial.get("tool") or ALL_TOOLS),
               ("ACCOUNT", "fAcct", "acct", initial.get("account") or ALL_ACCOUNTS),
               ("DEPARTMENT", "fDept", "dept", initial.get("department") or ALL_DEPTS),
               ("PERSON", "fPerson", "person", initial.get("person") or EVERYONE),
               ("CLIENT / PROJECT", "fClient", "client", initial.get("client") or ALL_CLIENTS),
               ("SEARCH", "fSearch", None, initial.get("search") or None))
    for r in range(R_FILTER - 1, R_FILTER + 2):
        for c in range(2, 14):
            grid.put(r, c, None, fill=FILL_PANEL)
    ws.row_dimensions[R_FILTER].height = 22
    for k, (label, name, list_key, value) in enumerate(filters):
        c = 2 + 2 * k
        grid.put(R_FILTER - 1, c, label, font=F_CAPTION, fill=FILL_PANEL, text=True)
        grid.put(R_FILTER, c, value, font=F_INPUT, fill=FILL_INPUT, align=LEFT_MID,
                 border=Border(left=INPUT_BOX.left, top=INPUT_BOX.top, bottom=INPUT_BOX.bottom),
                 text=value is not None)
        # the box's right half: a ▾ hint on dropdowns, empty on Search
        grid.put(R_FILTER, c + 1, "▾" if list_key else None, font=Font(color=AMBER), fill=FILL_INPUT,
                 align=RIGHT_MID, border=Border(right=INPUT_BOX.right, top=INPUT_BOX.top, bottom=INPUT_BOX.bottom),
                 text=bool(list_key))
        _name(wb, name, _abs(DASH, _col(c), R_FILTER))
        if list_key:
            dv = DataValidation(type="list", formula1=list_ranges[list_key], allow_blank=False,
                                showErrorMessage=True, errorTitle="Pick from the list",
                                error="Choose a value from the dropdown.")
            dv.add(f"{_col(c)}{R_FILTER}")
            ws.data_validations.append(dv)
    grid.put(R_FILTER + 1, 2, ("Click a yellow box and pick from its dropdown; pick “All …” (or clear Search) to reset. "
                               "The cards, charts and every tab follow these filters. Search narrows Activity and "
                               "People."), font=F_HELP, fill=FILL_PANEL, text=True)

    # Chart headers, underlined, with totals and measure toggles
    for c in range(2, 14):
        grid.put(R_CHART_HEAD, c, None, border=UNDERLINE)
    grid.put(R_CHART_HEAD, 2, "Spend trend (daily)", font=F_SECTION, border=UNDERLINE, text=True)
    grid.put(R_CHART_HEAD, 6, "Usage by tool", font=F_SECTION, border=UNDERLINE, text=True)
    grid.put(R_CHART_HEAD, 10, '="Department-wise "&CHOOSE(mDeptIdx,"spend","outputs","credits")', font=F_SECTION,
             border=UNDERLINE)
    vr = R_CHART_HEAD + 1
    ws.row_dimensions[vr].height = 22
    grid.put(vr, 2, f'="₹"&{_indian("kTotal")}', font=Font(size=14, bold=True, color=INK), align=LEFT_MID)
    grid.put(vr, 3, '=IF(OR(kChange="—",kChange=""),"",kChange)', font=F_CHANGE, align=RIGHT_MID)
    for formula, color in CHANGE_RULES:
        ws.conditional_formatting.add(f"C{vr}", FormulaRule(formula=[formula.format(c=f"C{vr}")],
                                                            font=Font(size=9, bold=True, color=color)))
    grid.put(vr, 6, (f'=IF(mToolIdx=2,"₹","")&{_indian("SUM(dn_value)")}&" "&'
                     'CHOOSE(mToolIdx,"outputs","spent","credits")'), font=Font(size=12, bold=True), align=LEFT_MID)
    grid.put(vr, 10, f'=IF(mDeptIdx=1,"₹","")&{_indian("CHOOSE(mDeptIdx,kItem,kOutputs,kCredits)")}',
             font=Font(size=12, bold=True), align=LEFT_MID)
    for col, name, value, key in ((9, "mTool", TOOL_MEASURES[0], "tool_m"), (13, "mDept", DEPT_MEASURES[0], "dept_m")):
        grid.put(vr, col - 1, "Show", font=F_CAPTION, align=RIGHT_MID, text=True)
        grid.put(vr, col, value, font=F_INPUT, fill=FILL_INPUT, border=INPUT_BOX, align=CENTER, text=True)
        _name(wb, name, _abs(DASH, _col(col), vr))
        dv = DataValidation(type="list", formula1=list_ranges[key], allow_blank=False)
        dv.add(f"{_col(col)}{vr}")
        ws.data_validations.append(dv)

    # Legend tables under the charts: colour swatch, name, share bar, value
    lg = R_LEGEND
    for col, text, numeric in ((2, "Busiest day", False), (3, None, True), (6, "Tool", False), (8, "Share", True),
                               (9, "Value", True), (10, "Department", False), (13, "Value", True)):
        grid.put(lg, col, text, font=F_THEAD, fill=FILL_THEAD, align=RIGHT_MID if numeric else LEFT_MID,
                 text=text is not None)
    for c in (4, 7, 11, 12):
        grid.put(lg, c, None, fill=FILL_THEAD)
    for i in range(6):
        r = lg + 1 + i
        line = Border(bottom=HAIR)
        tool = f"INDEX(dn_name,{i + 1})"         # the swatch shows the slice colour; the name links to Tools
        grid.put(r, 6, f'=IF({tool}="","",{_linked(tool, targets["tool"], label=_swatch(tool))})',
                 font=Font(color=PALETTE[i], bold=True), border=line)
        grid.put(r, 7, None, border=line)
        grid.put(r, 8, f'=IF(INDEX(dn_name,{i + 1})="","",INDEX(dn_pct,{i + 1}))', fmt=PCT_FMT, border=line)
        grid.put(r, 9, f'=IF(INDEX(dn_name,{i + 1})="","",INDEX(dn_value,{i + 1}))', fmt="#,##0", border=line)
        dept = f"INDEX(dc_name,{i + 1})"
        grid.put(r, 10, f'=IF({dept}="","",{_linked(dept, targets["dept"], label=_swatch(dept))})',
                 font=Font(color=PALETTE[i], bold=True), border=line)
        for c in (11, 12):
            grid.put(r, c, None, border=line)
        grid.put(r, 13, f'=IF(INDEX(dc_name,{i + 1})="","",INDEX(dc_value,{i + 1}))', fmt="#,##0", border=line)
    ws.conditional_formatting.add(f"H{lg + 1}:H{lg + 6}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1, color="8B5CF6"))
    ws.conditional_formatting.add(f"M{lg + 1}:M{lg + 6}", DataBarRule(
        start_type="num", start_value=0, end_type="max", color=GREEN))
    # Under the trend: busiest day (its header is the table header row), then
    # two captioned stats.
    busiest = "INDEX(dl_date,MATCH(MAX(dl_spend),dl_spend,0))"
    busiest_link = _linked(busiest, targets["date"], label=f'TEXT({busiest},"d mmm")')
    grid.put(lg + 1, 2, f'=IF(MAX(dl_spend)<=0,"—",{busiest_link})', font=F_XLINK_BOLD)
    grid.put(lg + 1, 3, "=MAX(dl_spend)", fmt=INR_FMT, font=F_BOLD, align=RIGHT_MID)
    grid.put(lg + 3, 2, "DAILY AVERAGE", font=F_CAPTION, text=True)
    grid.put(lg + 4, 2, "per calendar day", font=F_FOOT, text=True)
    grid.put(lg + 4, 3, "=kItem/kDays", fmt=INR_FMT, font=F_BOLD, align=RIGHT_MID)
    grid.put(lg + 5, 2, "ACTIVE DAYS", font=F_CAPTION, text=True)
    grid.put(lg + 6, 2, "days with outputs", font=F_FOOT, text=True)
    grid.put(lg + 6, 3, '=COUNTIF(dl_out,">0")', fmt=COUNT_FMT, font=F_BOLD, align=RIGHT_MID)

    _widths(ws, {"A": 2, **{_col(c): 13 for c in range(2, 14)}, "N": 2})
    grid.flush(ws)
    _add_charts(ws, calc_ws, lay)


def _no_border(chart) -> None:
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))


def _add_charts(ws, calc_ws, lay: Layout) -> None:
    dl, dn, dc = lay.blocks["dl"], lay.blocks["dn"], lay.blocks["dc"]
    anchor_row = R_CHART
    grid_line = GraphicalProperties(ln=LineProperties(solidFill=LINE))

    trend = AreaChart()
    trend.legend = None
    trend.height, trend.width = 8.2, 10.2
    trend.add_data(Reference(calc_ws, min_col=dl.col("spend"), min_row=dl.first - 1, max_row=dl.last),
                   titles_from_data=True)
    trend.set_categories(Reference(calc_ws, min_col=dl.col("label"), min_row=dl.first, max_row=dl.last))
    series = trend.series[0]
    series.graphicalProperties.solidFill = "A7F3D0"
    series.graphicalProperties.line.solidFill = PALETTE[0]
    series.graphicalProperties.line.width = 22000
    trend.y_axis.numFmt = '"₹"#,##0'
    trend.y_axis.majorGridlines.spPr = grid_line
    trend.x_axis.tickLblSkip = max(1, len(lay.days) // 6)
    trend.x_axis.delete = False
    trend.y_axis.delete = False
    _no_border(trend)
    ws.add_chart(trend, f"B{anchor_row}")

    donut = DoughnutChart(holeSize=62)
    donut.height, donut.width = 8.2, 10.2
    donut.add_data(Reference(calc_ws, min_col=dn.col("value"), min_row=dn.first - 1, max_row=dn.last),
                   titles_from_data=True)
    donut.set_categories(Reference(calc_ws, min_col=dn.col("name"), min_row=dn.first, max_row=dn.last))
    for i in range(6):
        point = DataPoint(idx=i)
        point.graphicalProperties.solidFill = PALETTE[i]
        point.graphicalProperties.line.solidFill = "FFFFFF"
        donut.series[0].dPt.append(point)
    donut.legend = None          # the table under the chart is its legend, as on the page
    _no_border(donut)
    ws.add_chart(donut, f"F{anchor_row}")

    bars = BarChart()
    bars.type = "col"
    bars.legend = None
    bars.gapWidth = 60
    bars.height, bars.width = 8.2, 10.2
    bars.add_data(Reference(calc_ws, min_col=dc.col("value"), min_row=dc.first - 1, max_row=dc.last),
                  titles_from_data=True)
    bars.set_categories(Reference(calc_ws, min_col=dc.col("name"), min_row=dc.first, max_row=dc.last))
    for i in range(6):
        point = DataPoint(idx=i)
        point.graphicalProperties.solidFill = PALETTE[i]
        bars.series[0].dPt.append(point)
    # Value only (Excel otherwise adds the series and category names), shown
    # as the rounded value the feed cells hold.
    bars.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False,
                                    showPercent=False)
    bars.y_axis.delete = True
    bars.x_axis.delete = False
    bars.y_axis.majorGridlines = None
    _no_border(bars)
    ws.add_chart(bars, f"J{anchor_row}")


# --------------------------------------------------------------------------- #
# Entry points
# --------------------------------------------------------------------------- #
def resolve_initial(data: DashboardData, *, tool: Optional[str] = None, account_ids=(),
                    department: Optional[str] = None, user_id: Optional[int] = None,
                    client: Optional[str] = None, search: Optional[str] = None) -> dict:
    """The page's current filters as the workbook's starting filter values.
    The Account dropdown is single-choice: several logins ticked start it
    at All accounts."""
    tools = {r.tool for r in data.rows + data.prev_rows} | set(data.plan_fee) | set(data.untracked)
    labels = {data.account_by_id[i] for i in account_ids if i in data.account_by_id}
    clients = {r.client for r in data.rows + data.prev_rows}
    return {
        "tool": tool if tool in tools else None,
        "account": labels.pop() if len(labels) == 1 else None,
        "department": department if department in set(data.people.values()) else None,
        "person": data.person_by_id.get(user_id) if user_id is not None else None,
        "client": client if client in clients else None,
        "search": (search or "").strip() or None,
    }


def write_dashboard_workbook(data: DashboardData, path: str, *, initial: Optional[dict] = None) -> None:
    initial = initial or {}
    lay = _layout(data)
    wb = Workbook(write_only=True)
    sheets = {name: wb.create_sheet(name) for name in (DASH,) + TABS + (DATA, AGG, CALC, LISTS)}

    _write_data(wb, sheets[DATA], data.rows)
    _write_agg(wb, sheets[AGG], data.rows + data.prev_rows)
    calc = _build_calc(wb, data, lay)
    calc.grid.flush(sheets[CALC])
    list_ranges = _write_lists(sheets[LISTS], lay)

    n_activity = min(len(data.rows), ACTIVITY_ROWS)
    targets = _link_targets(lay, n_activity)
    _write_dashboard(sheets[DASH], wb, data, lay, sheets[CALC], list_ranges, initial, targets)
    _write_activity(sheets[ACTIVITY], n_activity, targets)
    _write_ranked_tab(sheets[PEOPLE], PEOPLE, "Everyone with at least one output, by spend. Search also narrows this list.",
                      "p_rank", len(lay.people), [
                          ("Person", "p_name", None, 24, "name", "act_person"), ("Department", "p_dept", None, 18, None, "dept"),
                          ("Tools used", "p_tools", None, 34), ("Top client", "p_topclient", None, 20, None, "client"),
                          ("Outputs", "p_out", COUNT_FMT, 10), ("Active days", "p_days", COUNT_FMT, 11),
                          ("Credits", "p_cred", CREDITS_FMT, 11), ("Spent", "p_spend", INR_FMT, 12, "bar"),
                          ("Per output", "p_perout", INR2_FMT, 11), ("Share", "p_share", PCT_FMT, 9, "share")], targets)
    _write_ranked_tab(sheets[TOOLS], TOOLS, "Tools by total spend. Plan fees not used by anyone count only "
                                            "while no filter is on.", "t_rank", len(lay.tools), [
                          ("Tool", "t_name", None, 14, "name", "act_tool"), ("Category", "t_cat", None, 18),
                          ("People", "t_people", COUNT_FMT, 8), ("Top user", "t_top", None, 22, None, "person"),
                          ("Outputs", "t_out", COUNT_FMT, 10), ("Credits", "t_cred", CREDITS_FMT, 11),
                          ("Usage", "t_usage", INR_FMT, 12), ("Plan fee", "t_fee", INR_FMT, 12),
                          ("Total", "t_total", INR_FMT, 12, "bar"), ("Per output", "t_perout", INR2_FMT, 11),
                          ("Share", "t_share", PCT_FMT, 9, "share")], targets)
    _write_ranked_tab(sheets[CLIENTS], CLIENTS, "Clients and projects by spend.", "c_rank", len(lay.clients), [
                          ("Client / Project", "c_name", None, 28, "name", "act_client"), ("People", "c_people", COUNT_FMT, 8),
                          ("Top user", "c_top", None, 22, None, "person"), ("Top tool", "c_toptool", None, 14, None, "tool"),
                          ("Outputs", "c_out", COUNT_FMT, 10), ("Credits", "c_cred", CREDITS_FMT, 11),
                          ("Spent", "c_spend", INR_FMT, 12, "bar"), ("Per output", "c_perout", INR2_FMT, 11),
                          ("Share", "c_share", PCT_FMT, 9, "share")], targets)
    _write_ranked_tab(sheets[DAILY], DAILY, "Days with activity, newest first.", "dl_rank", len(lay.days), [
                          ("Date", "dl_date", DATE_FMT, 14, "name", "act_date"), ("People", "dl_people", COUNT_FMT, 8),
                          ("Outputs", "dl_out", COUNT_FMT, 10), ("Top tool", "dl_toptool", None, 14, None, "tool"),
                          ("Credits", "dl_cred", CREDITS_FMT, 11), ("Spent", "dl_spend", INR_FMT, 12, "scale"),
                          ("Per output", "dl_perout", INR2_FMT, 11)], targets)
    _write_ranked_tab(sheets[MONTHLY], MONTHLY, "Months with activity, newest first.", "ml_rank", len(lay.months), [
                          ("Month", "ml_label", None, 12, "name"), ("Active days", "ml_days", COUNT_FMT, 11),
                          ("People", "ml_people", COUNT_FMT, 8), ("Outputs", "ml_out", COUNT_FMT, 10),
                          ("Top tool", "ml_toptool", None, 14, None, "tool"), ("Credits", "ml_cred", CREDITS_FMT, 11),
                          ("Spent", "ml_spend", INR_FMT, 12, "bar"), ("Per active day", "ml_perday", INR_FMT, 13)], targets)
    _write_ranked_tab(sheets[DEPTS], DEPTS, "Departments by spend.", "d_rank", len(lay.depts), [
                          ("Department", "d_name", None, 22, "name", ("person", "d_top")), ("People", "d_people", COUNT_FMT, 8),
                          ("Top user", "d_top", None, 22, None, "person"), ("Top tool", "d_toptool", None, 14, None, "tool"),
                          ("Outputs", "d_out", COUNT_FMT, 10), ("Credits", "d_cred", CREDITS_FMT, 11),
                          ("Spent", "d_spend", INR_FMT, 12, "bar"), ("Per person", "d_perperson", INR_FMT, 12),
                          ("Share", "d_share", PCT_FMT, 9, "share")], targets)
    _write_compare(sheets[COMPARE], lay, list_ranges, data)

    for name in (DATA, AGG, CALC, LISTS):
        sheets[name].sheet_state = "hidden"
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def generate_dashboard_workbook(db: Session, filters: ReportFilters, path: str, *, dashboard_url: str = "",
                                selection: Optional[dict] = None) -> None:
    """`filters`: the period and the test-account setting only. `selection`:
    the page's current filters (resolve_initial's keywords), which become the
    workbook's starting filter values."""
    data = collect(db, filters, dashboard_url=dashboard_url)
    write_dashboard_workbook(data, path, initial=resolve_initial(data, **(selection or {})))
