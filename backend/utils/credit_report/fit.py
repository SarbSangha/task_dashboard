"""
Credit Consumption Report - column widths that fit what is written.

Run on each buffered sheet after navigation.resolve() and before flush, when
every cell (links and back cells included) is known. A text cell may spill
into empty cells to its right, so it only needs widening when the next
non-empty cell would cut it off; a number never spills. For each column:

    width = max(planned width, widest text that would be cut otherwise)

capped at MAX_WIDTH. Text that would need more than the cap (a sentence
with a filled cell beside it) wraps in its cell instead; Excel sizes the row.
Header cells already wrap and are left alone.
"""

from __future__ import annotations

from datetime import date, datetime

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from . import navigation as N

MAX_WIDTH = 45
DEFAULT_WIDTH = 8.43                # Excel's width for a column with none set
WRAP_WIDTH = 30                     # a column holding wrapped sentences is at least this wide


def shown(cell) -> str:
    """The text a cell displays, numbers formatted roughly as Excel shows them."""
    if cell is None or cell.value in (None, ""):
        return ""
    v = cell.value
    if isinstance(v, str):
        if v.startswith("=") and not v.startswith("=HYPERLINK("):
            return ""                                   # a formula's result is not known here
        if v.startswith("=HYPERLINK("):
            m = N.LINK_RE.match(v)                      # the whole label, arrows and "(list)" included
            return m.group("label").replace('""', '"') if m else ""
        return v
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (datetime, date)):
        fmt = cell.number_format or ""
        return "Sep 2026" if fmt == "mmm yyyy" else ("29 Sep 2026 (Tue)" if "ddd" in fmt else "29 Sep 2026")
    if isinstance(v, (int, float)):
        fmt = cell.number_format or ""
        if "%" in fmt:
            return f"{v:.1%}"
        if "₹" in fmt:
            return f"₹{v:,.0f}" + ("x" if "0.00" in fmt else "")
        return f"{v:,.2f}" if "0.00" in fmt else f"{v:,.0f}"
    return str(v)


def text_width(text: str, font) -> float:
    """Width in Excel units (one '0' at 11 pt), with a little padding."""
    size = float(getattr(font, "sz", None) or 11)
    bold = bool(getattr(font, "b", False))
    # A smaller font never shrinks the estimate: Excel versions render small text differently.
    return len(text) * max(size / 11.0, 1.0) * (1.08 if bold else 1.0) * 1.05 + 1.5


def _wraps(cell) -> bool:
    al = getattr(cell, "alignment", None)
    return bool(al and al.wrap_text)


def fit_columns(buf, skip_cols=()) -> None:
    """Widen (or wrap) so that no text in the BufferedSheet is cut off."""
    dims = buf.column_dimensions
    widths = {}

    def width(col: int) -> float:
        if col not in widths:
            w = dims[get_column_letter(col)].width if get_column_letter(col) in dims else None
            widths[col] = float(w) if w else DEFAULT_WIDTH
        return widths[col]

    to_wrap = []
    for cells in buf.rows:
        last = len(cells)
        for i, cell in enumerate(cells, start=1):
            if i in skip_cols:
                continue
            text = shown(cell)
            if not text or _wraps(cell):
                continue
            need = text_width(text, cell.font)
            # Room to the right: the empty cells a text value spills into.
            spill = 0.0
            if isinstance(cell.value, str):
                j = i + 1
                while j <= last and (cells[j - 1] is None or cells[j - 1].value in (None, "")):
                    spill += width(j)
                    j += 1
                if j > last:
                    continue                            # nothing to its right: it shows in full
            if need <= width(i) + spill:
                continue
            required = need - spill
            if required <= MAX_WIDTH:
                widths[i] = max(width(i), required)
            elif isinstance(cell.value, str):
                widths[i] = max(width(i), WRAP_WIDTH)
                to_wrap.append(cell)
            else:
                widths[i] = max(width(i), MAX_WIDTH)
    for col, w in widths.items():
        letter = get_column_letter(col)
        if letter not in dims or (dims[letter].width or DEFAULT_WIDTH) < w:
            dims[letter].width = round(w, 1)
    for cell in to_wrap:
        old = cell.alignment
        cell.alignment = Alignment(wrap_text=True, vertical=(old.vertical if old and old.vertical else "top"),
                                   horizontal=old.horizontal if old else None)


__all__ = ["fit_columns", "text_width", "shown", "MAX_WIDTH"]
