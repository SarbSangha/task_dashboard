"""
Credit Consumption Report - detail blocks under the By User and By Tool lists.

Each sheet keeps its ranked list (the Excel Table) on top; below it, one
block per user / per tool, in list order, grouped with Excel's row outline
so each can be collapsed with its +/- button (expanded by default).

User block:  header band (name, figures, "⬅ Back to user list", "All
             generations (n) →"), department + most used tool + mostly works
             on client, then Tools used, Clients, Date-wise and
             Tool › Date › Client (an Excel Table, filterable by Client).
Tool block:  header band, then Users, Clients, Date-wise and
             Date › User › Client.
Department block (By Department): header band, then People and Tools.

Each block's second or third row is a jump row ("↓ Tools used", ...) to its
sections; each section title has "⬅ Block top" back to it.

Links inside a block: a Date-wise row opens the first row of its date (or
week) in Tool › Date › Client / Date › User › Client, a Clients row the first
row of its client; those rows carry "⬅ Date-wise" / "⬅ Clients" cells that
return to the exact row. Every Tool › Date › Client and Date › User › Client
row opens its own (user, tool, day, client) run in the Generation Log.

Every row position is planned (plan_user_blocks / plan_tool_blocks) before
anything is written, because the Generation Log - written first - links
back to them; the writers assert they land where the plan said.
"""

from __future__ import annotations

from datetime import timedelta

from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.properties import Outline

from . import navigation as N
from . import workbook as W
from .blocks import BlockData, date_rows, day_client_key, pick_most_used, usage_key

DAY_FMT = "dd mmm yyyy (ddd)"
FILL_BAND = PatternFill("solid", fgColor="BDD7EE")
F_BAND = Font(bold=True, size=13, color=W.NAVY)
F_SUB = Font(bold=True, color=W.NAVY)
BAND_COLS = 14                     # header band spans A..N

# Link cells (fixed columns, so the log's back links can point at them).
USER_ALL_GENS_COL = "N"            # block header: "All generations (n) →"
USER_TOOLS_LINK_COL = "G"          # Tools used: "Generations (n) →"
USER_TDC_LINK_COL = "F"            # Tool › Date › Client: "Generations (n) →"
TOOL_USERS_LINK_COL = "I"          # Users: "Generations (n) →"
TOOL_DUC_LINK_COL = "F"            # Date › User › Client: "Generations (n) →"
BACK_DATEWISE = "⬅ Date-wise"
BACK_CLIENTS = "⬅ Clients"

TOOLS_HEADERS = ["Tool", "Generations", "Credits", "% of user's credits", "Top client on this tool", "Last used",
                 "Log rows", "In By Tool"]
SECTION_BACK_COL = 12              # "⬅ Block top" on each section title row (column L)
PEOPLE_HEADERS = ["User", "Employee ID", "Credits", "% of department", "Generations", "Pending credits",
                  "Most used tool"]
DEPT_TOOLS_HEADERS = ["Tool", "Credits", "Generations", "Users"]
CLIENTS_HEADERS = ["Client", "Generations", "Credits", "% of user's credits", "Tools used"]
TDC_HEADERS = ["Tool", "Date", "Client", "Generations", "Credits", "Log rows", BACK_DATEWISE, BACK_CLIENTS]
TOOL_USERS_HEADERS = ["User", "Department", "Generations", "Credits", "% of tool", "Top client", "Clients",
                      "Last used", "Log rows"]
TOOL_CLIENTS_HEADERS = ["Client", "Generations", "Credits", "Users"]
TOOL_DAYS_HEADERS = ["Date", "Credits", "Generations", "Users", "Clients"]
DUC_HEADERS = ["Date", "User", "Client", "Generations", "Credits", "Log rows", BACK_DATEWISE]


def gens_label(agg, prefix: str = "Generations") -> str:
    """'Generations (48) →', naming pending / failed rows inside the range."""
    extra = "".join(f" · {n:,} {what}" for n, what in ((agg.pending, "pending"), (agg.failed, "failed")) if n)
    return f"{prefix} ({agg.rows:,}{extra}) →"


def _span(first, last) -> str:
    if not first:
        return "no dates"
    return f"{first:%d %b %Y}" if first == last else f"{first:%d %b %Y} – {last:%d %b %Y}"


def _section(n_rows: int) -> int:
    return 3 + max(1, n_rows)       # title, header, rows (at least one), blank


def _client_order(data):
    return data.log.client_order


def _period(day, weekly: bool):
    """The Date-wise row a day belongs to: the day, or the Monday of its week."""
    return day - timedelta(days=day.weekday()) if weekly else day


def tdc_rows(u, data) -> list:
    """[(tool, day, client, Agg)] - tool most used first, day newest first, client in By Client order."""
    order = _client_order(data)
    out = []
    for tool in u.tool_order():
        keys = [(d, c) for (t, d), clients in u.tdc.items() if t == tool for c in clients]
        for day, client in sorted(keys, key=lambda k: day_client_key(k[0], k[1], order)):
            out.append((tool, day, client, u.tdc[(tool, day)][client]))
    return out


def duc_rows(t, data, users_order) -> list:
    """[(day, user id, client, Agg)] - day newest first, user as in the Users table, client in By Client order."""
    order, upos = _client_order(data), {uid: i for i, uid in enumerate(users_order)}
    out = []
    for day in sorted(t.duc, reverse=True):
        for (uid, client) in sorted(t.duc[day], key=lambda k: (upos.get(k[0], len(upos)), order.get(k[1], len(order)))):
            out.append((day, uid, client, t.duc[day][(uid, client)]))
    return out


def _tool_users(d):
    return sorted(d.users, key=lambda k: usage_key(str(k), d.users[k]))


# --------------------------------------------------------------------------- #
# Plans
# --------------------------------------------------------------------------- #
def plan_user_blocks(model, data: BlockData, lay, first_row: int) -> None:
    row = first_row
    for u in model.users:
        d = data.users.get(u.user_id)
        if not d:
            continue
        uid = u.user_id
        lay.user_block_row[uid] = row
        r = row + 3                                     # header, info, jump row
        sections = lay.block_sections[(W.USER, uid)] = []
        sections.append(("Tools used", r))
        for i, tool in enumerate(d.tool_order()):
            lay.user_tools_row[(uid, tool)] = r + 2 + i
        r += _section(len(d.tools))
        sections.append(("Clients", r))
        for i, client in enumerate(_clients_sorted(d.clients)):
            lay.user_clients_row[(uid, client)] = r + 2 + i
        r += _section(len(d.clients))
        sections.append(("Date-wise", r))
        days, weekly = date_rows(d.days)
        lay.user_weekly[uid] = weekly
        for i, dr in enumerate(days):
            lay.user_date_row[(uid, dr.day)] = r + 2 + i
        r += _section(len(days))
        sections.append(("Tool › Date › Client", r))
        tdc = tdc_rows(d, data)
        for i, (tool, day, client, _a) in enumerate(tdc):
            at = r + 2 + i
            lay.user_tdc_row[(uid, tool, day, client)] = at
            lay.user_tdc_date_row.setdefault((uid, _period(day, weekly)), at)
            lay.user_tdc_client_row.setdefault((uid, client), at)
        r += _section(len(tdc))
        lay.user_block_end[uid] = r - 1
        row = r + 1                                     # a blank row between blocks


def plan_tool_blocks(tool_order: list, data: BlockData, lay, first_row: int) -> None:
    row = first_row
    for tool in tool_order:
        d = data.tools.get(tool)
        if not d:
            continue
        lay.tool_block_row[tool] = row
        r = row + 2                                     # header, jump row
        sections = lay.block_sections[(W.TOOL, tool)] = [("Users", r)]
        users = _tool_users(d)
        for i, uid in enumerate(users):
            lay.tool_users_row[(tool, uid)] = r + 2 + i
        r += _section(len(users))
        sections.append(("Clients", r))
        r += _section(len(d.clients))
        sections.append(("Date-wise", r))
        days, weekly = date_rows(d.days)
        lay.tool_weekly[tool] = weekly
        for i, dr in enumerate(days):
            lay.tool_date_row[(tool, dr.day)] = r + 2 + i
        r += _section(len(days))
        sections.append(("Date › User › Client", r))
        duc = duc_rows(d, data, users)
        for i, (day, uid, client, _a) in enumerate(duc):
            at = r + 2 + i
            lay.tool_duc_row[(tool, day, uid, client)] = at
            lay.tool_duc_date_row.setdefault((tool, _period(day, weekly)), at)
        r += _section(len(duc))
        lay.tool_block_end[tool] = r - 1
        row = r + 1


def _clients_sorted(clients: dict) -> list:
    return sorted(clients, key=lambda c: (-round(clients[c].credits, 4), -clients[c].gens, c.lower()))


# --------------------------------------------------------------------------- #
# Writers
# --------------------------------------------------------------------------- #
def _here(ws) -> int:
    return len(ws.rows) + 1


def _expect(ws, row: int, what: str) -> None:
    if _here(ws) != row:
        raise AssertionError(f"{ws.title}: {what} planned for row {row}, writing row {_here(ws)}")


def _outline(ws, first: int, last: int) -> None:
    for r in range(first, last + 1):
        ws.row_dimensions[r].outlineLevel = 1
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False)
    ws.sheet_format.outlineLevelRow = 1


def _band(w, cells: dict) -> list:
    """A header band row: {column index: cell}, filled A..N."""
    out = []
    for i in range(1, max(BAND_COLS, max(cells)) + 1):
        c = cells.get(i) or w.cell()
        c.fill = FILL_BAND
        out.append(c)
    return out


def _jump_row(ws, w, sheet: str, lay, key) -> list:
    """'↓ Tools used', '↓ Clients', ... - down to the block's sections."""
    return [w.nav(W.internal_target(sheet, row), f"↓ {name}", font=Font(color="0563C1", underline="single", size=9))
            for name, row in lay.block_sections[key]]


def _section_back(w, sheet: str, lay, key, index: int):
    """'⬅ Block top' on a section title: back to its jump cell."""
    jump_row = lay.block_sections[key][0][1] - 1
    return _back(w, sheet, jump_row, W.get_column_letter(index + 1), "⬅ Block top")


def _table(ws, w, title: str, headers: list, rows: list, table_name: str = None, back=None) -> None:
    title_row = [w.text(title, font=F_SUB)]
    if back is not None:
        title_row += [w.cell() for _ in range(SECTION_BACK_COL - 2)] + [back]
    ws.append(title_row)
    header_row = _here(ws)
    ws.append(w.header(headers))
    for r in rows:
        ws.append(r)
    if not rows:
        ws.append([w.text("—", font=W.F_ITALIC)])
    elif table_name:
        W._add_table(ws, headers, header_row + len(rows), table_name, header_row=header_row)
    ws.append([])


def _back(w, sheet, row, col, label):
    return w.link(W.internal_target(sheet, row, col), label, font=N.F_BACK, fill=N.FILL_BACK)


def _day_label(day, weekly: bool) -> str:
    return f"Week of {day:%d %b %Y}" if weekly else f"{day:%d %b %Y (%a)}"


def _top_user_text(model, top_uid, named_uid) -> str:
    if top_uid is None:
        return "—"
    text = model.user(top_uid).label
    if top_uid == W.UNASSIGNED_USER_ID and named_uid is not None:
        text += f" · Top named user: {model.user(named_uid).label}"
    return text


def write_user_blocks(sheets, model, data: BlockData, lay) -> None:
    ws, w = sheets[W.USER]
    plan = data.log
    for index, u in enumerate(model.users, start=1):
        d = data.users.get(u.user_id)
        if not d:
            continue
        uid, h = u.user_id, lay.user_block_row[u.user_id]
        weekly = lay.user_weekly[uid]
        _expect(ws, h, f"{u.label}'s block")
        t = d.total
        info = (f"{t.credits:,.2f} credits · {t.gens:,} gen{'' if t.gens == 1 else 's'} · "
                f"{len(d.days):,} active day{'' if len(d.days) == 1 else 's'} · {_span(t.first, t.last)}")
        ws.append(_band(w, {
            1: w.text(u.label, font=F_BAND),
            3: w.text(info, font=F_SUB),
            12: w.link(W.internal_target(W.USER, lay.user_row[uid]), "⬅ Back to user list"),
            14: w.link(W.internal_target(W.LOG, plan.user_header[uid], "A"), gens_label(t, "All generations"))
            if uid in plan.user_header else w.text("—"),
        }))
        ws.append([w.link(W.internal_target(W.DEPT, lay.dept_row[u.department]), u.department), w.cell(),
                   w.text(f"Most used tool: {d.most_used_tool().text()}   ·   "
                          f"Mostly works on client: {d.top_client().text()}", font=W.F_ITALIC)])
        bkey = (W.USER, uid)
        ws.append(_jump_row(ws, w, W.USER, lay, bkey))

        tools = []
        for tool in d.tool_order():
            a = d.tools[tool]
            tools.append([
                W.tool_link(w, lay, tool),
                w.count(a.gens), w.credits(a.credits), w.pct(a.credits / t.credits if t.credits else 0.0),
                w.text(d.tool_top_client(tool).name or "—"),
                w.cell(W.as_datetime(a.last), fmt=DAY_FMT) if a.last else w.text("—"),
                w.link(W.internal_target(W.LOG, plan.pair_row[(uid, tool)], "A"), gens_label(a))
                if (uid, tool) in plan.pair_row else w.text("—"),
                # Sideways: this user's row in the tool's block (its "⬅" cell returns here).
                w.link(W.internal_target(W.TOOL, lay.tool_users_row[(tool, uid)]), W.sideways(tool, W.TOOL))
                if (tool, uid) in lay.tool_users_row else w.text("—"),
            ])
        _expect(ws, h + 3, "Tools used")
        _table(ws, w, "Tools used (most used first)", TOOLS_HEADERS, tools, back=_section_back(w, W.USER, lay, bkey, 0))

        clients = []
        for client in _clients_sorted(d.clients):
            a = d.clients[client]
            first = lay.user_tdc_client_row.get((uid, client))
            # Down to the client's first Tool › Date › Client row; its "⬅ Clients" cell links back.
            name = (w.nav(W.internal_target(W.USER, first, "H"), f"↓ {client}") if first else w.text(client))
            clients.append([name, w.count(a.gens), w.credits(a.credits),
                            w.pct(a.credits / t.credits if t.credits else 0.0),
                            w.text(", ".join(sorted(d.client_tools.get(client, ()), key=str.lower)))])
        _table(ws, w, "Clients", CLIENTS_HEADERS, clients, back=_section_back(w, W.USER, lay, bkey, 1))

        rows, _weekly = date_rows(d.days)
        order = d.tool_order()
        days = []
        for r in rows:
            first = lay.user_tdc_date_row.get((uid, r.day))
            label = _day_label(r.day, weekly)
            days.append([w.nav(W.internal_target(W.USER, first, "G"), label) if first else w.text(label)]
                        + [w.credits(r.credits.get(tool, 0.0)) for tool in order]
                        + [w.credits(r.total, font=W.F_BOLD), w.text(", ".join(sorted(r.clients, key=str.lower)))])
        _table(ws, w, "Date-wise, by week (more than 60 active dates)" if weekly else "Date-wise (newest first)",
               ["Week" if weekly else "Date"] + order + ["Total", "Clients that week" if weekly else "Clients that day"],
               days, back=_section_back(w, W.USER, lay, bkey, 2))

        tdc, seen_period, seen_client = [], set(), set()
        for tool, day, client, a in tdc_rows(d, data):
            key = (uid, tool, day, client)
            period = _period(day, weekly)
            back_date = back_client = w.cell()
            if period not in seen_period:
                seen_period.add(period)
                back_date = _back(w, W.USER, lay.user_date_row[(uid, period)], "A", BACK_DATEWISE)
            if client not in seen_client:
                seen_client.add(client)
                back_client = _back(w, W.USER, lay.user_clients_row[(uid, client)], "A", BACK_CLIENTS)
            tdc.append([w.text(tool), w.cell(W.as_datetime(day), fmt=DAY_FMT), w.text(client), w.count(a.gens),
                        w.credits(a.credits),
                        w.link(W.internal_target(W.LOG, plan.quad_row[key], "C"), gens_label(plan.quad[key]))
                        if key in plan.quad_row else w.text("—"),
                        back_date, back_client])
        _table(ws, w, "Tool › Date › Client (filter Client to see one client)", TDC_HEADERS, tdc,
               table_name=f"tblUserTDC{index}", back=_section_back(w, W.USER, lay, bkey, 3))
        _expect(ws, lay.user_block_end[uid] + 1, f"end of {u.label}'s block")
        ws.append([])
        _outline(ws, h + 1, lay.user_block_end[uid])


def write_tool_blocks(sheets, model, data: BlockData, lay, tool_order: list) -> None:
    ws, w = sheets[W.TOOL]
    plan = data.log
    for tool in tool_order:
        d = data.tools.get(tool)
        if not d:
            continue
        h = lay.tool_block_row[tool]
        weekly = lay.tool_weekly[tool]
        _expect(ws, h, f"{tool}'s block")
        t = d.total
        users = _tool_users(d)
        active = [uid for uid in users if d.users[uid].gens or d.users[uid].credits]
        named = [uid for uid in active if uid != W.UNASSIGNED_USER_ID]
        info = (f"{t.credits:,.2f} credits · {t.gens:,} gen{'' if t.gens == 1 else 's'} · {len(d.users):,} "
                f"user{'' if len(d.users) == 1 else 's'} · Top user: "
                f"{_top_user_text(model, active[0] if active else None, named[0] if named else None)} · Top client: "
                f"{d.top_client().name or '—'} · {_span(t.first, t.last)}")
        ws.append(_band(w, {1: w.text(tool, font=F_BAND), 3: w.text(info, font=F_SUB),
                            12: w.link(W.internal_target(W.TOOL, lay.tool_row[tool]), "⬅ Back to tool list")}))
        bkey = (W.TOOL, tool)
        ws.append(_jump_row(ws, w, W.TOOL, lay, bkey))

        rows = []
        for uid in users:
            a, u = d.users[uid], model.user(uid)
            clients = d.user_clients.get(uid, {})
            rows.append([
                W.user_link(w, lay, uid, u.label),
                w.text(u.department), w.count(a.gens), w.credits(a.credits),
                w.pct(a.credits / t.credits if t.credits else 0.0),
                w.text(pick_most_used(clients, a, skip=W.NO_CLIENT).name or "—"),
                w.text(", ".join(sorted(clients, key=str.lower))),
                w.cell(W.as_datetime(a.last), fmt=DAY_FMT) if a.last else w.text("—"),
                w.link(W.internal_target(W.LOG, plan.pair_row[(uid, tool)], "B"), gens_label(a))
                if (uid, tool) in plan.pair_row else w.text("—"),
            ])
        _table(ws, w, "Users (most used first)", TOOL_USERS_HEADERS, rows, back=_section_back(w, W.TOOL, lay, bkey, 0))

        clients = []
        for client in _clients_sorted(d.clients):
            a = d.clients[client]
            clients.append([W.client_link(w, lay, client), w.count(a.gens), w.credits(a.credits),
                            w.text(", ".join(sorted((model.user(x).label for x in d.client_users[client]),
                                                    key=str.lower)))])
        _table(ws, w, "Clients", TOOL_CLIENTS_HEADERS, clients, back=_section_back(w, W.TOOL, lay, bkey, 1))

        drows, _weekly = date_rows(d.days)
        days = []
        for r in drows:
            first = lay.tool_duc_date_row.get((tool, r.day))
            label = _day_label(r.day, weekly)
            days.append([w.nav(W.internal_target(W.TOOL, first, "G"), label) if first else w.text(label),
                         w.credits(r.total), w.count(r.gens),
                         w.text(", ".join(sorted((model.user(x).label for x in r.people), key=str.lower))),
                         w.text(", ".join(sorted(r.clients, key=str.lower)))])
        _table(ws, w, "Date-wise, by week (more than 60 active dates)" if weekly else "Date-wise (newest first)",
               ["Week" if weekly else "Date"] + TOOL_DAYS_HEADERS[1:4]
               + ["Clients that week" if weekly else "Clients that day"], days,
               back=_section_back(w, W.TOOL, lay, bkey, 2))

        duc, seen_period = [], set()
        for day, uid, client, a in duc_rows(d, data, users):
            key = (uid, tool, day, client)
            period = _period(day, weekly)
            back_date = w.cell()
            if period not in seen_period:
                seen_period.add(period)
                back_date = _back(w, W.TOOL, lay.tool_date_row[(tool, period)], "A", BACK_DATEWISE)
            duc.append([w.cell(W.as_datetime(day), fmt=DAY_FMT), w.text(model.user(uid).label), w.text(client),
                        w.count(a.gens), w.credits(a.credits),
                        w.link(W.internal_target(W.LOG, plan.quad_row[key], "D"), gens_label(plan.quad[key]))
                        if key in plan.quad_row else w.text("—"),
                        back_date])
        _table(ws, w, "Date › User › Client", DUC_HEADERS, duc, back=_section_back(w, W.TOOL, lay, bkey, 3))
        _expect(ws, lay.tool_block_end[tool] + 1, f"end of {tool}'s block")
        ws.append([])
        _outline(ws, h + 1, lay.tool_block_end[tool])


# --------------------------------------------------------------------------- #
# Department blocks (By Department)
# --------------------------------------------------------------------------- #
def _dept_tools(model, data: BlockData, department: str) -> dict:
    """{tool: Agg + users} for the department's people."""
    from .blocks import Agg

    out = {}
    for u in W.dept_members(model, department):
        d = data.users.get(u.user_id)
        if not d:
            continue
        for tool, a in d.tools.items():
            agg = out.setdefault(tool, (Agg(), set()))
            agg[0].rows += a.rows
            agg[0].gens += a.gens
            agg[0].credits += a.credits
            if a.gens or a.credits:
                agg[1].add(u.user_id)
    return out


def plan_dept_blocks(model, data: BlockData, lay, first_row: int) -> None:
    row = first_row
    for dept in model.departments:
        lay.dept_block_row[dept.name] = row
        r = row + 2                                     # header, jump row
        people = W.dept_members(model, dept.name)
        lay.block_sections[(W.DEPT, dept.name)] = [("People", r), ("Tools", r + _section(len(people)))]
        r += _section(len(people)) + _section(len(_dept_tools(model, data, dept.name)))
        lay.dept_block_end[dept.name] = r - 1
        row = r + 1


def write_dept_blocks(sheets, model, data: BlockData, lay) -> None:
    ws, w = sheets[W.DEPT]
    for dept in model.departments:
        h = lay.dept_block_row[dept.name]
        _expect(ws, h, f"{dept.name}'s block")
        people = W.dept_members(model, dept.name)
        t = dept.totals
        info = (f"{len(people)} {'person' if len(people) == 1 else 'people'} · {t.credits:,.2f} credits · "
                f"{t.generations:,} gen{'' if t.generations == 1 else 's'}")
        ws.append(_band(w, {1: w.text(dept.name, font=F_BAND), 3: w.text(info, font=F_SUB),
                            12: w.link(W.internal_target(W.DEPT, lay.dept_row[dept.name]), "⬅ Back to department list")}))
        bkey = (W.DEPT, dept.name)
        ws.append(_jump_row(ws, w, W.DEPT, lay, bkey))
        rows = []
        for u in people:
            ut = u.totals
            top = u.top_tool.first if u.top_tool else None
            rows.append([W.user_link(w, lay, u.user_id, u.label), w.text(u.employee_id), w.credits(ut.credits),
                         w.pct(ut.credits / t.credits if t.credits else 0.0), w.count(ut.generations),
                         w.credits(ut.pending_credits), W.tool_link(w, lay, top) if top else w.text("—")])
        _table(ws, w, "People (most credits first)", PEOPLE_HEADERS, rows, back=_section_back(w, W.DEPT, lay, bkey, 0))
        tools = _dept_tools(model, data, dept.name)
        trows = [[W.tool_link(w, lay, tool), w.credits(a.credits), w.count(a.gens), w.count(len(users))]
                 for tool, (a, users) in sorted(tools.items(), key=lambda kv: usage_key(kv[0], kv[1][0]))]
        _table(ws, w, "Tools", DEPT_TOOLS_HEADERS, trows, back=_section_back(w, W.DEPT, lay, bkey, 1))
        _expect(ws, lay.dept_block_end[dept.name] + 1, f"end of {dept.name}'s block")
        ws.append([])
        _outline(ws, h + 1, lay.dept_block_end[dept.name])
