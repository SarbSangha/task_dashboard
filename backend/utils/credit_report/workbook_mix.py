"""
Credit Consumption Report - the "Tool Mix" sheet (full version, after By Client).

Answers, by items generated (never credits - see mix.py):
1. In a month, which tool did a person use most, and what % went to each other tool?
2. Per day, which tool did that person use most / least, and which generations were they?
3. For a client, which tools were used (most -> least, %)?

Layout:
    rows 1-3   breadcrumbs, title, subtitle
    row 4      "Jump to:" ↓ People · ↓ Clients (and ↓ steps down long lists)
    People     one row per person, most items first ("Unknown user" last):
               the name or "Open block →" goes down to the person's block
    Clients    one compact row per client: its tool split, "Open block →" to
               its By Client block (the block itself is not repeated here)
    blocks     one per person: a band, a jump row, then per month (newest
               first) a sentence, "<month> tool split" and "<month> day by day"

Every "Generations (n) →" in a tool split opens that person's (user, tool,
month) run in the Generation Log; every non-zero tool cell of a day row its
(user, tool, day) run. Both runs are contiguous because the log is ordered
user -> tool -> day (newest first); their first rows carry "⬅ Mix month" /
"⬅ Mix day" cells back here (see workbook._write_log).

Every row position is planned (plan_mix) before the log is written.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.styles import Font

from . import mix as M
from . import workbook as W
from . import workbook_blocks as WB
from .blocks import BlockData, items_key

WIDTHS = [32, 13, 13, 22, 26, 14, 14, 20, 14, 14, 14, 14, 14, 14]
JUMP_ROW = 4
PEOPLE_HEADER_ROW = 5
FIRST_ROW = 6
JUMP_STEP = 20                      # list / day rows per "↓" step (a landing shows about 25)
MIX_SPLIT_LINK_COL = "H"            # tool split: "Generations (n) →" (the log's "⬅ Mix month" returns here)
DAY_TOOL_FIRST_COL = 6              # day by day: one column per tool from F

PEOPLE_HEADERS = ["Person", "Department", "Items", W.RUPEES, "Main tool (%)", "Tools used (whole period)", "Months",
                  "Open"]
CLIENT_HEADERS = ["Client", "Items", W.RUPEES, "People", "Tools used (most → least, %)", "Most / least", "Open"]
SPLIT_HEADERS = ["Tool", "Items", "%", W.RUPEES, "% of ₹", "Days used", "Main client", "Log rows"]
DAY_HEADERS = ["Date", "Items", W.RUPEES, "Most used tool (n, %)", "Least used"]


@dataclass
class MixLayout:
    people: list = field(default_factory=list)          # user ids, most items first, Unknown user last
    clients: list = field(default_factory=list)         # client names, By Client order
    index_row: dict = field(default_factory=dict)       # user id -> People row
    client_row: dict = field(default_factory=dict)      # client -> Clients row
    clients_title_row: int = 0
    clients_header_row: int = 0
    blocks_title_row: int = 0
    block_row: dict = field(default_factory=dict)       # user id -> block header row
    block_end: dict = field(default_factory=dict)
    month_row: dict = field(default_factory=dict)       # (user id, month) -> sentence row
    split_row: dict = field(default_factory=dict)       # (user id, month, tool) -> tool split row
    day_row: dict = field(default_factory=dict)         # (user id, day) -> day by day row
    day_tools: dict = field(default_factory=dict)       # (user id, month) -> [tools], the day table's tool columns
    day_cell: dict = field(default_factory=dict)        # (user id, tool, day) -> (row, column letter)
    jumps: list = field(default_factory=list)           # [(label, row)] on row 4
    example: str = ""                                   # a sample answer for Home's quick answer

    @property
    def header_rows(self) -> list:
        return [r for r in (PEOPLE_HEADER_ROW, self.clients_header_row) if r]


def people_order(model, data: BlockData) -> list:
    users = [u.user_id for u in model.users if u.user_id in data.users and data.users[u.user_id].total.rows]
    return sorted(users, key=lambda uid: (uid == W.UNASSIGNED_USER_ID,) + items_key(model.user(uid).label,
                                                                                     data.users[uid].total))


def _months(d) -> list:
    return sorted(d.months(), reverse=True)


def _month_tools(d, month: str) -> list:
    """Tools with Charged items that month, most items first."""
    return [s.key for s in WB.tool_split({t: a for (m, t), a in d.month_tools.items() if m == month})]


def _month_days(d, month: str) -> list:
    return sorted((day for day in d.days if f"{day:%Y-%m}" == month), reverse=True)


def person_label(model, uid) -> str:
    u = model.user(uid)
    return f"{u.label} ({u.department})" if u.department else u.label


def plan_mix(model, data: BlockData, lay) -> MixLayout:
    m = MixLayout()
    m.people = people_order(model, data)
    m.clients = [c.name for c in model.clients if c.name in data.clients]
    for i, uid in enumerate(m.people):
        m.index_row[uid] = FIRST_ROW + i
    row = FIRST_ROW + max(1, len(m.people)) + 1                 # list, blank
    m.clients_title_row = row
    m.clients_header_row = row + 1
    for i, client in enumerate(m.clients):
        m.client_row[client] = row + 2 + i
    row += 2 + max(1, len(m.clients)) + 1                       # title, header, rows, blank
    m.blocks_title_row = row
    m.jumps = [("People", PEOPLE_HEADER_ROW), ("Clients", m.clients_title_row)]
    m.jumps += [(f"{model.user(m.people[i]).label}", FIRST_ROW + i) for i in range(JUMP_STEP, len(m.people), JUMP_STEP)]
    m.jumps += [(m.clients[i], m.client_row[m.clients[i]]) for i in range(JUMP_STEP, len(m.clients), JUMP_STEP)]
    if m.people:
        m.jumps.append(("Blocks", m.blocks_title_row))
    row += 2                                                    # title, blank
    for uid in m.people:
        d = data.users[uid]
        m.block_row[uid] = row
        r = row + 2                                             # band, jump row
        sections = lay.block_sections[(W.TOOL_MIX, uid)] = []
        for month in _months(d):
            name = WB.month_label(month)
            m.month_row[(uid, month)] = r
            r += 1                                              # sentence
            sections.append((f"{name} tool split", r))
            for i, tool in enumerate(_month_tools(d, month)):
                m.split_row[(uid, month, tool)] = r + 2 + i
            r += WB._section(len(_month_tools(d, month)))
            sections.append((f"{name} day by day", r))
            tools = _month_tools(d, month)
            m.day_tools[(uid, month)] = tools
            days = _month_days(d, month)
            for i, day in enumerate(days):
                at = r + 2 + i
                m.day_row[(uid, day)] = at
                for j, tool in enumerate(tools):
                    if d.days[day].items.get(tool):
                        m.day_cell[(uid, tool, day)] = (at, W.get_column_letter(DAY_TOOL_FIRST_COL + j))
            for i in range(JUMP_STEP, len(days), JUMP_STEP):
                sections.append((f"{name} from {days[i]:%d %b}", r + 2 + i))
            r += WB._section(len(days))
        m.block_end[uid] = r - 1
        row = r + 1
    if m.people:
        uid = m.people[0]
        d = data.users[uid]
        month = _months(d)[0] if _months(d) else None
        if month:
            shares = WB.tool_split({t: a for (mm, t), a in d.month_tools.items() if mm == month})
            m.example = f"{model.user(uid).label}, {WB.month_label(month)}: {M.chain(shares, limit=4) or 'nothing'}"
    return m


def month_sentence(model, uid, d, month: str) -> str:
    """'Oct 2026 — Kushi Saini (Internal Brands): 399 items, ₹X. Used Kling most (95%), then Freepik (3%) and
    HeyGen (2%). Least used: HeyGen.'"""
    total = d.month_total.get(month)
    shares = WB.tool_split({t: a for (m, t), a in d.month_tools.items() if m == month})
    items = total.gens if total else 0
    rupees = total.rupees if total else 0.0
    return (f"{WB.month_label(month)} — {person_label(model, uid)}: {M.items_text(items)}, {M.rupees_text(rupees)}. "
            f"{M.used_most_then(shares)}")


def busiest_note(d, days: list) -> str:
    """'Busiest day: 07 Oct (52 items, mostly Kling)'; ties together; '' with no items."""
    counts = {day: d.days[day].gens for day in days if d.days[day].gens}
    if not counts:
        return ""
    best = max(counts.values())
    top = sorted((day for day, n in counts.items() if n == best), reverse=True)

    def mostly(day):
        r = d.days[day]
        return M.names(M.most(WB.tool_split(r.items, r.item_rupees)))

    if len(top) == 1:
        return f"Busiest day: {top[0]:%d %b} ({M.items_text(best)}, mostly {mostly(top[0])})"
    return f"Busiest days: {M.names(M.Share(x, f'{x:%d %b}', 0, 0, 0) for x in top)} ({M.items_text(best)} each)"


def day_cell_label(items: int, tool: str, run) -> str:
    """'25 Kling →', or '1 Suno (+31 pending) →' when the run it opens also holds
    pending / failed rows - the label names every row the link opens."""
    extra = " · ".join(f"+{n:,} {what}" for n, what in ((run.pending, "pending"), (run.failed, "failed")) if n)
    return f"{items:,} {tool}" + (f" ({extra})" if extra else "") + " →"


def block_top(w, block_row: int, index: int):
    """'⬅ Block top' on a section title: back to its jump cell (the row under the band)."""
    return WB._back(w, W.TOOL_MIX, block_row + 1, W.get_column_letter(index + 1), "⬅ Block top")


JUMPS_BACK_COL = 10                 # "⬅ Jumps" on each row a row-4 jump lands on (right of the lists)


def _jump_backs(ws, w, m: MixLayout, lay) -> None:
    """Every "↓" jump's landing row gets a cell back to that jump: "⬅ Jumps" for
    row 4, "⬅ Block top" for a step down a long day table."""
    for i, (_label, row) in enumerate(m.jumps):
        ws.put(row, JUMPS_BACK_COL, WB._back(w, W.TOOL_MIX, JUMP_ROW, W.get_column_letter(2 + i), "⬅ Jumps"))
    for uid in m.people:
        for i, (name, row) in enumerate(lay.block_sections.get((W.TOOL_MIX, uid), ())):
            if " from " in name:
                ws.put(row, len(ws.rows[row - 1]) + 1, block_top(w, m.block_row[uid], i))


def write_mix(sheets, model, data: BlockData, lay, period_label: str) -> None:
    ws, w = sheets[W.TOOL_MIX]
    m: MixLayout = lay.mix
    plan = data.log
    W._preamble(ws, w, W.TOOL_MIX, f"{period_label} · Most used = most items generated (Charged; credits of "
                                   "different tools are never compared), ₹ beside every count. Each person: the "
                                   "tool split per month and day by day; each tool count opens those generations "
                                   "in the log.")
    if not m.people and not m.clients:
        W._empty_row(ws, w, len(PEOPLE_HEADERS))
        return
    jump_font = Font(color="0563C1", underline="single", size=9)
    ws.append([w.text("Jump to:", font=W.F_ITALIC)]
              + [w.nav(W.internal_target(W.TOOL_MIX, row), f"↓ {label}", font=jump_font) for label, row in m.jumps])  # 4
    ws.append(w.header(PEOPLE_HEADERS))                                                                          # 5
    for uid in m.people:
        d = data.users[uid]
        u = model.user(uid)
        shares = WB.tool_split(d.tools)
        block = m.block_row[uid]
        ws.append([w.nav(W.internal_target(W.TOOL_MIX, block), u.label), w.text(u.department),
                   w.count(d.total.gens), w.rupees(d.total.rupees), w.text(M.main_of(shares)),
                   w.text(M.chain(shares) or "—"), w.count(len(d.months())),
                   w.nav(W.internal_target(W.TOOL_MIX, block), "Open block →")])
    if not m.people:
        W._empty_row(ws, w, len(PEOPLE_HEADERS))
    else:
        W._add_table(ws, PEOPLE_HEADERS, FIRST_ROW + len(m.people) - 1, "tblMixPeople", header_row=PEOPLE_HEADER_ROW)
    ws.append([])

    WB._expect(ws, m.clients_title_row, "Tool Mix clients")
    ws.append([w.text("Clients: which tools each client got (most → least, by items)", font=W.F_SECTION)])
    ws.append(w.header(CLIENT_HEADERS))
    for client in m.clients:
        c = data.clients[client]
        shares = WB.tool_split(c.tools)
        block = lay.client_block_row.get(client)
        people = WB._named_people({uid: a.gens for uid, a in c.users.items()})
        ws.append([w.text(client, font=Font(italic=True) if client == W.NO_CLIENT else None),
                   w.count(c.total.gens), w.rupees(c.total.rupees), w.count(len(people)),
                   w.text(M.chain(shares) or "—"), w.text(M.most_least(shares) or "—"),
                   w.link(W.internal_target(W.CLIENT, block), "Open block →") if block else w.text("—")])
    if not m.clients:
        W._empty_row(ws, w, len(CLIENT_HEADERS))
    else:
        W._add_table(ws, CLIENT_HEADERS, m.clients_header_row + len(m.clients), "tblMixClients",
                     header_row=m.clients_header_row)
    ws.append([])

    if not m.people:
        _jump_backs(ws, w, m, lay)
        return
    WB._expect(ws, m.blocks_title_row, "Tool Mix blocks")
    ws.append([w.text("Tool mix by person and month", font=W.F_SECTION)])
    ws.append([])
    for index, uid in enumerate(m.people, start=1):
        d = data.users[uid]
        h = m.block_row[uid]
        WB._expect(ws, h, f"{model.user(uid).label}'s tool mix block")
        t = d.total
        months = _months(d)
        dept = model.user(uid).department
        info = (f"{dept + ' · ' if dept else ''}{M.items_text(t.gens)} · {M.rupees_text(t.rupees)} · "
                f"{len(months)} month{'' if len(months) == 1 else 's'} · {WB._span(t.first, t.last)}")
        ws.append(WB._band(w, {1: w.text(model.user(uid).label, font=WB.F_BAND), 3: w.text(info, font=WB.F_SUB),
                               12: w.link(W.internal_target(W.TOOL_MIX, m.index_row[uid]), "⬅ Back to people list")}))
        bkey = (W.TOOL_MIX, uid)
        ws.append(WB._jump_row(ws, w, W.TOOL_MIX, lay, bkey))
        names = [n for n, _r in lay.block_sections[bkey]]

        for month in months:
            name = WB.month_label(month)
            WB._expect(ws, m.month_row[(uid, month)], f"{model.user(uid).label} {month}")
            ws.append([w.text(month_sentence(model, uid, d, month), font=Font(bold=True, color=W.NAVY))])

            shares = WB.tool_split({tl: a for (mm, tl), a in d.month_tools.items() if mm == month})
            rupee_pct = M.whole_percents({s.key: s.rupees for s in shares})
            rows = []
            for s in shares:
                a = d.month_tools[(month, s.key)]
                clients = d.month_tool_clients.get((month, s.key), {})
                main = WB.tool_split({c: x for c, x in clients.items() if c != W.NO_CLIENT} or clients)
                run = (uid, s.key, month)
                rows.append([W.tool_link(w, lay, s.key), w.count(s.items), w.text(f"{s.pct_text}"),
                             w.rupees(s.rupees), w.text(M.pct_text(rupee_pct[s.key], s.rupees) if sum(rupee_pct.values()) else "—"),
                             w.count(len(d.month_tool_days.get((month, s.key), ()))),
                             w.text(M.names(M.most(main)) or "—"),
                             w.link(W.internal_target(W.LOG, plan.month_run_row[run], lay.log_cols[W.LOG_BACK_MIX_MONTH]),
                                    WB.gens_label(plan.month_run[run])) if run in plan.month_run_row else w.text("—")])
            WB._table(ws, w, f"{name} tool split", SPLIT_HEADERS, rows,
                      back=block_top(w, h, names.index(f"{name} tool split")))

            tools = m.day_tools[(uid, month)]
            days = _month_days(d, month)
            note = busiest_note(d, days)
            rows = []
            for day in days:
                r = d.days[day]
                split = WB.tool_split(r.items, r.item_rupees)
                cells = [w.cell(W.as_datetime(day), fmt=WB.DAY_FMT), w.count(r.gens), w.rupees(r.rupees),
                         w.text(WB.count_pct(split)), w.text(WB.least_text(split))]
                for tool in tools:
                    n = r.items.get(tool, 0)
                    run = (uid, tool, day)
                    if n and run in plan.day_run_row:
                        cells.append(w.link(W.internal_target(W.LOG, plan.day_run_row[run],
                                                              lay.log_cols[W.LOG_BACK_MIX_DAY]),
                                            day_cell_label(n, tool, plan.day_run[run])))
                    else:
                        cells.append(w.count(n, font=Font(color="A6A6A6")))
                cells.append(w.text(", ".join(sorted(r.clients, key=str.lower))))
                rows.append(cells)
            WB._table(ws, w, f"{name} day by day" + (f" — {note}" if note else ""),
                      W.unique_headers(DAY_HEADERS + tools + ["Clients that day"]), rows,
                      table_name=f"tblMix{index}d{month.replace('-', '')}",
                      back=block_top(w, h, names.index(f"{name} day by day")))
        WB._expect(ws, m.block_end[uid] + 1, f"end of {model.user(uid).label}'s tool mix block")
        ws.append([])
        WB._outline(ws, h + 1, m.block_end[uid])
    _jump_backs(ws, w, m, lay)


__all__ = ["MixLayout", "plan_mix", "write_mix", "month_sentence", "busiest_note", "MIX_SPLIT_LINK_COL", "WIDTHS"]
